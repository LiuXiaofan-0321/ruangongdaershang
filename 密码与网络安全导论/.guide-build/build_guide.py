from pathlib import Path
import re, json, html, shutil
from pptx import Presentation
from lxml import html as LH

ROOT=Path(__file__).resolve().parent.parent
B=ROOT/'.guide-build'
TARGET=ROOT/'密码学_第二三讲_深度学习指南.html'
BACKUP=B/'original.html'
if not BACKUP.exists(): shutil.copy2(TARGET,BACKUP)
original=BACKUP.read_text(encoding='utf-8')
DATA=json.loads(re.search(r'const DATA=(\{[\s\S]*?\});\s*</script>',original).group(1))
deck=Presentation(str(ROOT/'第四讲——对称加密.pptx'))
def rawtable(page,index=0):
    ts=[s.table for s in deck.slides[page-1].shapes if s.has_table]
    return [[c.text for c in row.cells] for row in ts[index].rows]
def table(caption,heads,rows):
    th=''.join('<th scope="col">'+str(x)+'</th>' for x in heads)
    body=''.join('<tr>'+''.join('<td>'+str(x)+'</td>' for x in row)+'</tr>' for row in rows)
    return f'<div class="table-wrap"><table><caption>{caption}</caption><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table></div>'
def source_table(page,index,title,head=True):
    rows=rawtable(page,index)
    if not head and sum(bool(x) for x in rows[0])==1: rows=rows[1:]
    return table(title,rows[0] if head else [str(i+1) for i in range(len(rows[0]))],rows[1:] if head else rows)
def think(q,a,rubric):
    return f'<details class="think"><summary>先自己思考 · {q}</summary><div class="answer">{a}<p class="rubric">评分采分点：{rubric}</p></div></details>'
def call(kind,title,txt):return f'<div class="callout {kind}"><strong>{title}</strong><p>{txt}</p></div>'
def diagram(src,desc):
    return '<div class="diagram"><pre class="mermaid">'+html.escape(src)+'</pre><div class="diagram-render" role="img" aria-label="'+html.escape(desc,quote=True)+'"></div><p class="diagram-fallback">'+desc+'</p></div>'
modules=[]; quizzes=[]; bodies=[]
current=''; ci=0
def cloze(answer):
    global ci
    ci+=1
    return f'<button type="button" class="cloze" data-concept="{current}-c{ci}" aria-expanded="false"><span>{answer}</span></button>'
def start(id):
    global current,ci
    current=id;ci=0
def add(id,title,lecture,pages,summary,body,questions):
    n=19+len(modules)
    modules.append(dict(id=id,title=title,lecture=lecture,pages=pages,summary=summary))
    bodies.append(f'<section class="chapter added-chapter" id="{id}" data-title="{title}"><header class="chapter-head"><div class="chapter-number">学习模块 {n:02d} / 34</div><span class="source-tag">{lecture} {pages}｜符号铺垫、推导补全与反例为教学补充</span><h2>{title}</h2><p class="chapter-intro">{summary}</p></header><div class="chapter-body">{body}</div><div class="quiz-zone" data-quiz-module="{id}"></div><footer class="chapter-footer"><a class="btn" href="#welcome">学习路线</a><a class="btn" href="#coverage">核对原课件页码</a></footer></section>')
    for j,q in enumerate(questions,1):
        question,opts,explanation=q
        quizzes.append(dict(id=f'{id}-q{j}',module=id,type='multi' if sum(o[1] for o in opts)>1 else 'single',question=question,options=[dict(text=t,correct=c,reason=r) for t,c,r in opts],explanation=explanation,source=f'{lecture} {pages}；原例改编或依重难点原创'))
def opt(text,correct,reason):return (text,correct,reason)

AS=[int(c,16) for row in rawtable(55,0)[1:] for c in row[1:]]
SS=[int(c,16) for row in rawtable(47)[1:] for c in row[1:]]
def rol(x,n,w=32):return ((x<<n)|(x>>(w-n)))&((1<<w)-1)
def aeskeys(key):
    w=[list(key[i:i+4]) for i in range(0,16,4)];rc=[1,2,4,8,16,32,64,128,27,54]
    for i in range(4,44):
        t=w[i-1].copy()
        if i%4==0:t=[AS[x] for x in t[1:]+t[:1]];t[0]^=rc[i//4-1]
        w.append([a^b for a,b in zip(w[i-4],t)])
    return [bytes(sum(w[i:i+4],[])) for i in range(0,44,4)]
def tau(x):return sum(SS[(x>>(24-8*i))&255]<<(24-8*i) for i in range(4))
def lin(x,shifts):
    y=x
    for n in shifts:y^=rol(x,n)
    return y
FK=[0xA3B1BAC6,0x56AA3350,0x677D9197,0xB27022DC]
CK=[int.from_bytes(bytes(((4*i+j)*7)%256 for j in range(4)),'big') for i in range(32)]
MK=[0x01234567,0x89abcdef,0xfedcba98,0x76543210]
ks=[a^b for a,b in zip(MK,FK)];rk=[];smrows=[];X=MK.copy()
for i in range(32):
    ka=ks[i+1]^ks[i+2]^ks[i+3]^CK[i];kb=tau(ka);kt=lin(kb,[13,23]);knew=ks[i]^kt
    ks.append(knew);rk.append(knew)
    a=X[i+1]^X[i+2]^X[i+3]^rk[i];s=tau(a);t=lin(s,[2,10,18,24]);X.append(X[i]^t)
    smrows.append([i,f'{CK[i]:08X}',f'{rk[i]:08X}',f'{a:08X}',f'{s:08X}',f'{t:08X}',f'{X[i+4]:08X}'])
assert ''.join(f'{x:08X}' for x in X[-1:-5:-1])=='681EDF34D206965E86B3E94F536E4246'
assert aeskeys(bytes(range(16)))[1].hex()=='d6aa74fdd2af72fadaa678f1d6ab76fe'
PS=[14,4,13,1,2,15,11,8,3,10,6,12,5,9,0,7]
def sub16(x,s=PS):return sum(s[(x>>(12-4*i))&15]<<(12-4*i) for i in range(4))
def perm16(x):
    y=0
    for src,dst in enumerate([1,5,9,13,2,6,10,14,3,7,11,15,4,8,12,16],1):y|=((x>>(16-src))&1)<<(16-dst)
    return y
spkeys=[0x3A94,0xA94D,0x94D6,0x4D63,0xD63F];sp=0x26B7;sprows=[]
for i in range(4):
    u=sp^spkeys[i];v=sub16(u);sp=perm16(v) if i<3 else v^spkeys[4]
    sprows.append([i+1,f'{spkeys[i]:04X}',f'{u:04X}',f'{v:04X}',f'{sp:04X}' if i<3 else '末轮无P，白化后 '+f'{sp:04X}'])
assert sp==0xBCD6

start('t01')
body=r'''
<h3>1. 先区分消息、密钥和算法</h3>
<p>第二、三讲已经讲过“可逆加密”和“不可逆的摘要”。现在把可逆加密做成能处理大量数据的工具。发送者和接收者事先持有同一份秘密 K，发送者把明文 M 交给加密算法得到密文 C，接收者用同一 K 解密。共享的是密钥；加解密的操作方向仍可能不同。算法可以公开，密钥需要保密。</p>
<p>例如你与朋友事先约定一把密钥，之后在公开网络上发送密文。窃听者可以知道算法、拿到密文，甚至知道一些明密文对；安全要求仍是其在规定资源内难以恢复其他明文。密钥怎样安全交到朋友手里，是另一个需要解决的问题。对称算法通常速度快、适合批量数据，密钥分发与管理却是关键难点。结构简单不是所有对称算法的必然性质。</p>
<h3>2. 三个算法与一个必须成立的等式</h3>
<p>沿用第四讲定义2.3，明文空间 ℳ 是允许输入的消息集合，密钥空间 𝒦 是合法密钥集合，密文空间 𝒞 是算法输出集合。方案 SE=(Gen,Enc,Dec)：Gen(param*) 产生 K；Enc(M,K) 产生 C；Dec(C,K) 恢复 M。param* 表示一个或多个参数，例如安全参数；实际生成密钥还需要合格的随机性。</p>
<div class="formula">\[\operatorname{Dec}(\operatorname{Enc}(M,K),K)=M\quad(M\in\mathcal M,\ K\text{ 为合法生成的密钥}).\]</div>
<p>这是<strong>正确性</strong>：合法接收者能恢复原文。它没有证明保密性。例如 Enc(M,K)=M 也满足该等式，却直接公开明文。带随机数或 IV 的实际方案可写 Enc(K,M;IV)，并把解密所需的公开 IV 随密文携带；它不应被误认为额外秘密密钥。</p>
<h3>3. 分组密码、流密码与工作模式</h3>
<p>bit 是一个0或1，8 bit=1 byte。分组密码每次接收固定 n 位数据，在密钥 K 控制下输出同样 n 位数据。n 是<strong>分组长度</strong>，k 是<strong>密钥长度</strong>，二者没有必须相等的规定。AES 的 n=128，k 可为128、192或256；DES 的 n=64，有效 k=56。流密码先由短种子密钥和 IV 生成密钥流，再与消息逐位异或；实现也可以一次生成一个字，而非每次只计算一位。</p>
<p>一份1 MB文件远大于一个分组。工作模式规定如何反复调用分组密码、怎样安排分组之间的依赖。AES-CBC 中，AES 决定“一个128位块怎么加密”，CBC 决定“多个块怎样链接”。只说“使用AES”，还没有交代一套完整的文件加密方案。</p>
'''+diagram('flowchart LR\n A["主密钥 K"] --> B["密钥编排"]\n B --> C["轮密钥 K1 到 Kr"]\n M["一个 n 位明文块"] --> R["多轮迭代"]\n C --> R\n R --> O["一个 n 位密文块"]','多轮迭代属于分组算法内部；工作模式在算法外部组织多次完整调用。')+r'''
<h3>4. 为什么固定密钥下必须是双射</h3>
<p>{0,1}<sup>n</sup> 表示所有 n 位串，共有 2<sup>n</sup> 个。固定 K 后，E<sub>K</sub> 是这个有限集合上的排列。如果两个不同明文 M、M′ 被加密成相同 C，解密者只见 C 就无法判断是哪一个，因此 E<sub>K</sub> 必须单射；输入输出集合大小相等，所以同时满射，也就是双射。</p>
<div class="formula">\[E_K:\{0,1\}^n\to\{0,1\}^n,\quad D_K=E_K^{-1},\quad D_K(E_K(M))=M.\]</div>
<p>所有 n 位串上的排列共有 (2<sup>n</sup>)! 种；k 位密钥最多选择其中 2<sup>k</sup> 种。轮密钥编排是确定性展开，不会凭空增加独立的秘密信息。</p>
'''+call('intuition','核心直觉 · 先辨认层次','一轮是分组密码内部的一步；一次 E<sub>K</sub> 是执行完整分组算法；一种工作模式会多次调用完整 E<sub>K</sub>。AES-128 的10轮与“文件有10块”没有对应关系。')+f'<p>分组长度描述{cloze("一次处理多少数据")}；密钥长度描述{cloze("秘密参数有多少比特")}。满足正确性{cloze("不等于满足安全性")}。</p>'+think('为什么可逆变换不能把128位明文无损压缩成64位密文？','<p>128位明文有2<sup>128</sup>种，64位密文只有2<sup>64</sup>种。如果每种明文都合法，且固定密钥下仅靠64位密文恢复，鸽巢原理保证至少两个明文得到相同密文，违背唯一解密。限制消息集合或另带额外信息会改变问题条件。</p>','写出两种集合大小2分；说明必有碰撞2分；联系唯一解密1分。')
add('t01','对称加密：接口、分组与设计语言','第四讲','P1–10','理解共享密钥、Gen/Enc/Dec、正确性与双射，分清算法内部的轮和处理长消息的模式。',body,[
('AES-256 中的“256”表示什么？',[opt('分组长度256位。',False,'AES的分组长度始终128位。'),opt('主密钥长度256位。',True,'数字后缀标识密钥长度；AES-256有14轮。'),opt('执行256轮。',False,'轮数不是后缀值。'),opt('最多加密256字节。',False,'消息量限制需要结合模式等条件判断。')],'先把分组长度、密钥长度、轮数分开。'),
('关于 Gen、Enc、Dec，哪些说法成立？',[opt('正确性保证合法先加密再解密恢复原文。',True,'这是定义等式的含义。'),opt('正确性本身就证明保密。',False,'恒等变换就是反例。'),opt('Gen需要按方案要求生成合法密钥。',True,'合法性及密钥随机性影响安全。'),opt('IV必须与密钥一起保密。',False,'许多模式允许IV公开，约束是不可预测或唯一等。')],'接口说明怎样运算，安全还需另行分析。'),
('AES-CBC 中的 CBC 负责什么？',[opt('规定AES内部第几轮查S盒。',False,'这是AES算法层。'),opt('把128位密钥变成256位密钥。',False,'模式不改变AES密钥规格。'),opt('规定多个完整分组调用怎样链接。',True,'CBC使当前明文先异或前一密文。'),opt('把密文直接变成消息认证标签。',False,'CBC加密本身不提供认证。')],'一次完整E_K与内部一轮是不同层次。')])

start('t02')
body=r'''
<h3>1. 只有异或，为什么还不够</h3>
<p>逐位异或 ⊕ 的规则是相同为0、不同为1。例如课件的 01101001⊕11000101=10101100；再异或同一密钥就恢复，因为 a⊕a=0、a⊕0=a。它很适合把密钥混入状态，但若整个算法只有线性变换与异或，攻击者可利用简单代数关系。设计需要非线性部件。</p>
<p><strong>混淆</strong>使密钥与密文之间的关系难以分析，常由密钥混合和非线性 S 盒实现。S 盒就是输入小块后按表查出另一个值。<strong>扩散</strong>使某个输入位置的影响传播到多个输出位置，常由位置重排、线性混合和多轮迭代实现。香农1949年的思想奠定基础，DES、AES、SM4分别展示不同组合。</p>
<p>雪崩效应通常描述输入翻转1位后，输出约一半位发生变化的统计趋势。它不是“每个输入、每把密钥都恰好翻一半”，也不能独自证明安全。纯位置置换只移动比特，不会把一个1复制成多个1；必须和其他部件、多轮迭代配合。</p>
<h3>2. Feistel：整体可逆，但 F 可以不可逆</h3>
<p>把分组分成左右两半。第 i 轮先把旧右半搬到左边，再用 F 处理旧右半和轮密钥，与旧左半异或：</p>
<div class="formula">\[L_i=R_{i-1},\qquad R_i=L_{i-1}\oplus F(R_{i-1},K_i).\]
\[R_{i-1}=L_i,\qquad L_{i-1}=R_i\oplus F(L_i,K_i).\]</div>
<p>逆推第一式直接给出旧右半。将它代入第二式，再异或一次相同 F 值即可得到旧左半。整个过程没有求 F 的逆。因此 DES 的6进4出 S 盒虽然不可逆，DES 整体仍然可逆。多轮时，配合约定的首尾交换，解密可以复用同一轮结构，只把轮密钥逆序。</p>
'''+diagram('flowchart LR\n L["旧左半 L"] --> X["异或"]\n R["旧右半 R"] --> F["F：用轮密钥 K"]\n F --> X\n R --> NL["新左半等于旧右半"]\n X --> NR["新右半"]','Feistel保留旧右半，用同一个F值的两次异或消去影响，不需要逆F。')+r'''
<h3>3. SPN：逐层逆转，两个部件都要可逆</h3>
<p>代换—置换网络 SPN 通常对整个状态做密钥加、并行 S 盒代换与可逆线性变换 L。其一般示意式和逆式如下。逆序拆掉最外层 L，再拆 S，最后撤销异或密钥。</p>
<div class="formula">\[X_i=L(S(X_{i-1}\oplus K_i)),\quad X_{i-1}=S^{-1}(L^{-1}(X_i))\oplus K_i.\]</div>
<p>这里独立作用在整个状态的 S 和 L 都必须可逆。AES 是 SPN 的典型实例，但具体轮顺序、初始与末轮安排要以 AES 定义为准。不能把一般示意式逐字当作AES每轮流程。PRESENT 则展示4位 S 盒与比特置换的轻量设计。</p>
<h3>4. 原课件16位 SPN：不跳过中间轮</h3>
<p>每个 S 盒处理4位，4个并行，共16位；4轮，32位主密钥为3A94D63F。十六进制1位表示4个比特。轮密钥是长度16位的滑动窗口：K₁=3A94、K₂=A94D、K₃=94D6、K₄=4D63、K₅=D63F。这里的密钥编排用于教学，不能据此推断安全。</p>
'''+source_table(15,0,'第四讲P15：完整4位 S 盒')+source_table(15,1,'第四讲P15：输入位 z 移至输出位 πP(z)')+r'''
<p>P 的公式是 πP(i+4j)=4i+j−3，其中1≤i≤4、0≤j≤3。按最高位开始编号1到16，把16位排成4行后转置。因此输入第2位送往输出第5位；此处 P 恰好满足 P<sup>−1</sup>=P。主密钥位从最高位编号1到32，Kᵣ取第4r−3位到第4r+12位（含两端），r=1到5。前3轮 uᵣ=wᵣ₋₁⊕Kᵣ、vᵣ=S(uᵣ)、wᵣ=P(vᵣ)，w₀=x。末轮不做 P，输出 y=S(w₃⊕K₄)⊕K₅；额外的K₅异或叫输出白化。</p>
'''+table('x=26B7：四轮完整计算','轮次 Kᵣ uᵣ vᵣ wᵣ或密文'.split(),sprows)+r'''
<p>第一轮：26B7⊕3A94=1C23；逐半字节查表，1→4、C→5、2→D、3→1，得到45D1。将0100/0101/1101/0001排成四行，按列读为0010/1110/0000/0111，得到2E07。第二轮2E07⊕A94D=874A→3826→41B8。第三轮41B8⊕94D6=D56E→9FB0→E46E。第四轮E46E⊕4D63=A90D→6AE9，再⊕D63F=BCD6。</p>
<div id="spn-lab" class="native-lab"></div>
'''+f'<p>Feistel的 F {cloze("不必可逆")}；独立SPN代换层则{cloze("必须可逆")}。位置置换和雪崩效应{cloze("不能单独证明密码安全")}。</p>'+think('把密文 BCD6 一步步解回26B7','<p>先撤销白化：BCD6⊕D63F=6AE9。逆S表为 E 3 4 8 1 C A F 7 D 9 6 B 2 0 5（输入0至F依次对应）。S<sup>−1</sup>(6AE9)=A90D，⊕4D63=E46E。再逆P得9FB0，逆S得D56E，⊕94D6=41B8。再逆P得3826，逆S得874A，⊕A94D=2E07。最后逆P得45D1，逆S得1C23，⊕3A94=26B7。</p>','撤销白化1分；末轮逆S及密钥加2分；三轮按逆P、逆S、异或的顺序还原各2分；恢复原文1分。')
add('t02','混淆、扩散、Feistel 与完整 SPN 演算','第四讲','P11–16','看清可逆性从哪里来，完整走完16位SPN的加密和解密，并用实验台检查每一轮。',body,[
('为什么DES的6进4出 S 盒不妨碍解密？',[opt('6位与4位的集合大小相同。',False,'分别64种与16种，不相同。'),opt('每个S盒都是双射。',False,'6进4出无法单射。'),opt('Feistel逆推保留一半，并重新计算F后异或。',True,'不需要求S盒或F的逆。'),opt('解密靠穷举被丢弃的位。',False,'标准Feistel逆推是确定性操作。')],'可逆性取决于整体结构。'),
('SPN原例中，45D1经P置换以及最终密文分别是什么？',[opt('2E07与BCD6。',True,'P是4×4转置，末轮白化得到BCD6。'),opt('2E07与6AE9。',False,'6AE9只是末轮S输出，尚未⊕K5。'),opt('1C23与BCD6。',False,'1C23是第一轮密钥加后，不是P输出。'),opt('45D1与D63F。',False,'前者没有执行P，后者是K5。')],'分清u、v、w与最终y。'),
('哪些说法成立？',[opt('雪崩效应是统计性质，不能独自证明安全。',True,'局部扩散测试不能替代完整密码分析。'),opt('密钥扩展把128位密钥变成独立的1408位秘密。',False,'确定性展开不增加独立熵。'),opt('纯比特置换保持1的个数。',True,'只移动位置。'),opt('SPN中不可逆的独立S层也能无条件唯一解密。',False,'整层丢失信息后后续可逆层不能补回。')],'把混淆、扩散和安全结论分层判断。')])

start('t03')
body=r'''
<h3>1. DES 的三种长度</h3>
<p>DES由IBM设计，1977年发布为FIPS 46。每块64位，16轮Feistel。密钥常用64位表示，但第8、16、24、32、40、48、56、64位是奇校验位，只有56位参与加密。每个8位字节的1总数需为奇数，例如1010101有4个1，末尾补1得到10101011，共5个1。</p>
<p>奇校验只用于发现部分错误，两个比特同时翻转可能躲过检查；它不提供消息认证，也不增加有效密钥空间。DES的轮密钥是48位，数据左右半块各32位。这四个数字要分别记：64位分组、56位有效主密钥、48位轮密钥、32位半块。</p>
<h3>2. 总流程：IP、16轮、交换、IP逆置换</h3>
<p>IP只重排64位，不改变比特取值。输出前32位是L₀，后32位是R₀。依次执行Lᵢ=Rᵢ₋₁、Rᵢ=Lᵢ₋₁⊕F(Rᵢ₋₁,Kᵢ)，i=1到16。最后把两半交换，拼成R₁₆‖L₁₆，再做IP<sup>−1</sup>。下表均按行读，表项是“该输出位取输入的第几位”，不是字节数值。</p>
'''+diagram('flowchart TD\n A["64位明文"] --> B["IP初始置换"]\n B --> C["L0 与 R0：各32位"]\n C --> D["16轮Feistel：K1到K16"]\n D --> E["R16拼接L16：末尾交换"]\n E --> F["IP逆置换"]\n F --> G["64位密文"]','解密用相同流程，把轮密钥顺序改成K16到K1。IP与末置换互逆。')+source_table(20,0,'第四讲P20：IP完整表（输出取输入）',False)+source_table(20,1,'第四讲P20：IP逆置换完整表（输出取输入）',False)+r'''
<p>第一张表首项58表示新第1位取原第58位；末置换第58项为1，正好把它送回。IP<sup>−1</sup>(IP(X))=X。独立IP没有秘密，只知道IP规则不能得到密钥，也不能把IP当作安全来源。</p>
<h3>3. 轮函数 F 的四步</h3>
<div class="formula">\[F(R,K_i)=P\big(S(E(R)\oplus K_i)\big).\]</div>
<p><strong>E扩展：</strong>把32位分为8组，每4位左右各补一个相邻位，形成8组6位，共48位。首尾循环相邻。第一组 r₁r₂r₃r₄ 变为 r₃₂r₁r₂r₃r₄r₅。原例 r₃₂=1、中间0110、r₅=0，输出101100。</p>
'''+source_table(21,0,'第四讲P21：E扩展完整表',False)+r'''
<p><strong>混入密钥：</strong>48位 E(R) 与48位Kᵢ逐位异或，再切成8个6位小组。原例第一组101100⊕110111=011011。</p>
<p><strong>S盒代换：</strong>第j组进入第j张S盒，6位变4位；8张并行，得到32位。设输入b₁b₂b₃b₄b₅b₆，行号是首尾拼成的b₁b₆，列号是中间b₂b₃b₄b₅，行0到3，列0到15。表内是十进制，输出再写成4位二进制。课件提供S₁与S₂完整表，其余六张不同，不能拿S₁重复八次冒充DES。</p>
'''+source_table(24,0,'第四讲P23–25：S₁完整表')+source_table(24,1,'第四讲P24–25：S₂完整表')+r'''
<p>011011：首尾01=行1，中间1101=列13，查S₁得到5，即0101。注意行号不是最左两位，列号也不是最右四位。课堂题111111：行3、列15，S₁=13=1101；100110：行2、列3，S₂=11=1011。</p>
<p><strong>P置换：</strong>对S层的32位结果重排。新前四位取旧x₁₆、x₇、x₂₀、x₂₁。它把同一个S盒的输出分散到后续轮不同位置，配合E和多轮迭代扩散影响。</p>
'''+source_table(26,0,'第四讲P26：P置换完整表',False)+r'''
<h3>4. 为什么同一结构能解密</h3>
<p>对密文先IP，得到R₁₆‖L₁₆。令CL₀=R₁₆、CR₀=L₁₆。解密首轮用K₁₆，有CL₁=L₁₆=R₁₅，CR₁=R₁₆⊕F(L₁₆,K₁₆)=L₁₅。每轮又倒退一层。16轮后得到R₀‖L₀，再交换、IP逆置换，恢复明文。一般递推为CLᵢ=CRᵢ₋₁、CRᵢ=CLᵢ₋₁⊕F(CRᵢ₋₁,K₁₇₋ᵢ)。</p>
'''+f'<p>DES的S盒用{cloze("首尾两位选行，中间四位选列")}；解密轮函数{cloze("仍用同一个 F，轮密钥逆序")}；校验位{cloze("不增加有效密钥空间")}。</p>'+think('把第一组101100变成S₁输出0101，写出每步长度','<p>101100是E扩展结果中的6位。与本轮子密钥前6位110111异或，得到011011。首尾01为行1，中间1101为列13，S₁[1][13]=5，转为4位0101。八组各自查相应S盒，拼接32位后再做P；本题的4位不是整个F输出。</p>','异或2分；行列各1分；查值及4位表示1分；说明局部输出与整体长度1分。')
add('t03','DES：位表、轮函数与可逆结构','第四讲','P18–27','按位跟踪DES，完整保留课件中的IP、E、S₁/S₂、P表与课堂查表题。',body,[
('S₁输入011011时，行号、列号、输出依次是？',[opt('1、13、0101。',True,'首尾01，中间1101，查得5。'),opt('1、11、1011。',False,'把中间位误读成1011。'),opt('3、3、0010。',False,'首两位不能直接作行号。'),opt('0、27、11011。',False,'列号最多15，输出只有4位。')],'6位S盒输入的首尾和中间承担不同索引角色。'),
('关于DES长度，哪些成立？',[opt('明文分组64位。',True,'经过完整DES后仍64位。'),opt('每轮子密钥56位。',False,'每轮48位，56位是有效主密钥。'),opt('E把32位变48位。',True,'部分输入位重复选取。'),opt('8个校验位带来额外8位秘密。',False,'校验位由其他位确定。')],'读每个部件时同时标输入与输出宽度。'),
('DES解密的正确安排是什么？',[opt('改用逆S盒。',False,'6进4出S盒没有唯一逆；整体Feistel不需要它。'),opt('保留结构，K₁₆到K₁逆序使用。',True,'配合首尾置换和交换可还原。'),opt('只做一次IP逆置换。',False,'那只撤销位置重排，不能撤销16轮。'),opt('所有轮都用K₁。',False,'无法与原来对应轮的F抵消。')],'逆序的是轮密钥，不是把F四步逐项求逆。')])

start('t04')
body=r'''
<h3>1. DES 的密钥怎样变成16个子密钥</h3>
<p>PC-1从64位表示中去掉8个校验位，并重排剩余56位。前28位记C₀、后28位记D₀。第i轮两半分别循环左移sᵢ位得到Cᵢ、Dᵢ，拼接56位后按PC-2选出48位作为Kᵢ。送往下一轮的是Cᵢ、Dᵢ；48位Kᵢ不会反馈进移位链。这里Cᵢ是密钥编排半块，别与工作模式的密文Cᵢ混淆。</p>
'''+table('第四讲P28：每轮各自循环左移位数',[str(i) for i in range(1,17)],[[1,1,2,2,2,2,2,2,1,2,2,2,2,2,2,1]])+source_table(29,0,'第四讲P29：PC-1（输出取输入）',False)+source_table(29,1,'第四讲P29：PC-2（输出取输入）',False)+r'''
<p>PC-2排除Cᵢ‖Dᵢ中的第9、18、22、25、35、38、43、54位；它不仅删除，还改变选出位的顺序。首轮例K=133457799BBCDFF1，经PC-1得到F0CCAAF‖556678F；各左循环1位得到E19955F‖AACCF1E；PC-2选得K₁=1B02EFFC7072。所有数值是十六进制，前导0不能丢失位宽含义。</p>
<h3>2. 穷举强度：密钥组合数与攻击工作量</h3>
<p>有效密钥空间2<sup>56</sup>≈7.2×10<sup>16</sup>。若均匀选钥、遍历不重复，且已有足够明密文判别正确候选，最坏2<sup>56</sup>次，平均约2<sup>55</sup>次。8个校验位不会把它变成2<sup>64</sup>。DES已不适合现代保护需求，64位分组也限制同钥数据量；“密钥更长”与“分组更长”解决的是不同风险。</p>
<h3>3. DES-XEX3：前后异或白化</h3>
<div class="formula">\[C=K_3\oplus E_{K_2}(M\oplus K_1),\quad M=K_1\oplus D_{K_2}(C\oplus K_3).\]</div>
<p>先撤销最后的K₃异或，再做DES解密，最后撤销K₁异或。K₁、K₃是64位白化量，K₂是DES有效56位密钥。它只在中间调用一次完整DES。参数总位数不能直接当成已证明的安全强度，分组仍为64位。</p>
<h3>4. 双重 DES 为什么挡不住中间相遇</h3>
<div class="formula">\[C=E_{K_2}(E_{K_1}(M)),\quad M=D_{K_1}(D_{K_2}(C)),\quad E_{K_1}(M)=D_{K_2}(C).\]</div>
<p>攻击者知道一对(M,C)。先遍历2<sup>56</sup>个K₁，把E<sub>K₁</sub>(M)存表；再遍历2<sup>56</sup>个K₂，算D<sub>K₂</sub>(C)，查找相同中间值。每个匹配给出候选密钥对，再用其他明密文对排除误配。时间约2<sup>57</sup>次DES加解密，存储约2<sup>56</sup>项，而不是朴素2<sup>112</sup>。</p>
<p>一对64位明密文下，约2<sup>112</sup>/2<sup>64</sup>=2<sup>48</sup>个随机密钥对可能误配；所以“找到一个相遇”不能直接宣布找到唯一密钥。这个例子说明算法组合数2<sup>112</sup>不自动意味着112位攻击难度。</p>
'''+diagram('flowchart LR\n M["已知明文 M"] --> L["枚举K1：E_K1(M)"]\n L --> T["中间值表：查相等"]\n C["已知密文 C"] --> R["枚举K2：D_K2(C)"]\n R --> T\n T --> V["用更多明密文对筛候选"]','两端分别枚举，在中间相遇，减少时间但消耗巨大存储。')+r'''
<h3>5. 三重 DES 的四种形式</h3>
'''+table('第四讲P32：表中从左到右是执行顺序',['形式','加密执行顺序','对应解密顺序'],[['EEE3','E(K₁) → E(K₂) → E(K₃)','D(K₃) → D(K₂) → D(K₁)'],['EDE3','E(K₁) → D(K₂) → E(K₃)','D(K₃) → E(K₂) → D(K₁)'],['EEE2','E(K₁) → E(K₂) → E(K₁)','D(K₁) → D(K₂) → D(K₁)'],['EDE2','E(K₁) → D(K₂) → E(K₁)','D(K₁) → E(K₂) → D(K₁)']])+r'''
<p>3与2表示独立密钥数量。两密钥配置是K₃=K₁；组合数分别为2<sup>168</sup>和2<sup>112</sup>，不能直接等同安全强度。EDE且三把钥相同，E<sub>K</sub>(D<sub>K</sub>(E<sub>K</sub>(M)))=E<sub>K</sub>(M)，实现与单DES兼容。EEE同钥通常不能这样抵消。三重DES每块需执行三次DES，分组仍64位，是历史过渡技术。</p>
'''+f'<p>DES子密钥来自{cloze("PC-1、28位半块循环移位、PC-2")}；双重DES的关键反例是{cloze("中间相遇攻击")}；EDE兼容单DES的条件是{cloze("三把DES密钥相同")}。</p>'+think('为什么相遇表需要用更多明密文对验证？','<p>中间值只有64位，但密钥对有112位。对固定(M,C)，多个密钥对可能碰巧产生该映射。查表只筛选满足第一对数据的候选，另一独立明密文对会再施加64位约束，显著减少误配。额外数据的独立性和候选验证成本都属于攻击分析的一部分。</p>','写清相遇只给候选2分；指出64位与112位的差异2分；用额外数据排除1分。')
add('t04','DES 密钥编排与多重加密的安全陷阱','第四讲','P28–32','完整核对PC表与首轮子密钥，理解穷举、中间相遇和EEE/EDE的组合顺序。',body,[
('主密钥133457799BBCDFF1的首轮DES子密钥是什么？',[opt('F0CCAAF556678F。',False,'这是PC-1后的56位。'),opt('1B02EFFC7072。',True,'移位后再PC-2得到48位子密钥。'),opt('E19955FAACCF1E。',False,'这是左移后的56位拼接。'),opt('133457799BBCDFF1。',False,'这是含校验位的64位输入。')],'每个中间量都有不同长度和用途。'),
('双重DES的中间相遇分析，哪些正确？',[opt('时间约2⁵⁶+2⁵⁶次分组操作。',True,'忽略查表排序等额外开销。'),opt('不需要任何已知明密文。',False,'两端计算需要M和C。'),opt('匹配候选仍应由更多数据验证。',True,'64位中间值可能带来误配。'),opt('两把56位钥就必然112位安全。',False,'具体攻击可低于遍历全部密钥对的成本。')],'密钥空间大小只是穷举上界的一部分。'),
('三把密钥相同的EDE为什么等价单次DES？',[opt('中间D撤销内层E。',True,'E(D(E(M)))=E(M)。'),opt('任意三次E都等于一次E。',False,'EEE没有相邻逆操作抵消。'),opt('把分组长度扩大了三倍。',False,'仍处理64位分组。'),opt('K₃=K₁就一定等价单DES。',False,'若K₂不同，不能抵消；需三把都相同。')],'从右向左读复合公式，从左向右读执行顺序。')])

start('t05')
body=r'''
<h3>1. 16字节怎样放进状态矩阵</h3>
<p>AES基于Rijndael，2001年成为FIPS 197。三种密钥长度128/192/256位分别对应10/12/14轮，但分组一律128位=16字节。AES按字节运算，输入按<strong>列</strong>装进4×4状态矩阵：s[r,c]=input[4c+r]，r、c均从0开始。第一列是输入的前4个字节。</p>
<p>以明文00112233445566778899AABBCCDDEEFF、密钥000102030405060708090A0B0C0D0E0F为例。先做初始轮密钥加，字节00⊕00=00、11⊕01=10、22⊕02=20、33⊕03=30，因此首列为00、10、20、30。AES中的“字”是4字节，“半字节”是4位，“行移位1格”移动的是1字节，三者不能混用。</p>
'''+table('初始轮密钥加后的状态；按行显示、按列装入',['行/列','c=0','c=1','c=2','c=3'],[['r=0','00','40','80','C0'],['r=1','10','50','90','D0'],['r=2','20','60','A0','E0'],['r=3','30','70','B0','F0']])+diagram('flowchart TD\n A["128位明文"] --> B["初始AddRoundKey：rk0"]\n B --> C["第1到9轮：SubBytes"]\n C --> D["ShiftRows"]\n D --> E["MixColumns"]\n E --> F["AddRoundKey"]\n F --> G["第10轮：SubBytes、ShiftRows、AddRoundKey"]\n G --> H["128位密文"]','AES-128：初始密钥加不算一次完整轮，前9轮有列混淆，第10轮没有列混淆。')+r'''
<h3>2. SubBytes：改变值，保留位置</h3>
<p>16个字节各查同一张8进8出S盒。高4位是行号，低4位是列号，全部十六进制。00→63、10→CA。例如第一列00、10、20、30变为63、CA、B7、04。AES S盒可逆，与DES的6进4出盒不同；两者索引规则也不同。完整AES S盒在<a href="#t06-aes-sbox">下一章</a>，便于做密钥作业。</p>
<h3>3. ShiftRows：改变位置，保留值</h3>
<p>第0、1、2、3行分别循环左移0、1、2、3字节。第1行CA 53 60 70变为53 60 70 CA，第2行B7 D0 E0 E1变为E0 E1 B7 D0，第3行04 51 E7 8C变为8C 04 51 E7。它把原来同列字节送到不同列，下一步列内混合就能逐轮扩散到整个状态。</p>
'''+table('SubBytes与ShiftRows的完整对照',['行','SubBytes后','ShiftRows后'],[['0','63 09 CD BA','63 09 CD BA'],['1','CA 53 60 70','53 60 70 CA'],['2','B7 D0 E0 E1','E0 E1 B7 D0'],['3','04 51 E7 8C','8C 04 51 E7']])+r'''
<h3>4. MixColumns 的数学铺垫：字节是多项式</h3>
<p>GF(2⁸)不是“普通整数除以256取余”。把一个字节的8位作为x⁷到x⁰的系数，系数在模2下运算，所以加减都是异或。例如63=01100011，表示x⁶+x⁵+x+1。乘法先乘多项式，再模不可约多项式m(x)=x⁸+x⁴+x³+x+1，编码为11B，得到仍小于8次的多项式。</p>
<p>乘02就是乘x：先左移，若原最高位是1，产生的x⁸按x⁸≡x⁴+x³+x+1约减。实现为xtime(a)=((a&lt;&lt;1)&amp;FF)⊕(最高位为1时取1B，否则取00)。乘03=(02×a)⊕a。例如02×63=C6，没有溢出；02×E0=1C0⊕11B=DB，不能只截断为C0。</p>
<div class="formula">\[\begin{pmatrix}b_0\\b_1\\b_2\\b_3\end{pmatrix}=\begin{pmatrix}02&03&01&01\\01&02&03&01\\01&01&02&03\\03&01&01&02\end{pmatrix}\begin{pmatrix}a_0\\a_1\\a_2\\a_3\end{pmatrix}\quad\text{（矩阵元素为十六进制，运算在GF}(2^8)\text{）}.\]</div>
<h3>5. 把课件第一列的四个结果全部算出</h3>
<p>ShiftRows后的首列为63、53、E0、8C。先算所需乘积：02×63=C6，03×63=C6⊕63=A5；02×53=A6，03×53=A6⊕53=F5；02×E0=DB，03×E0=DB⊕E0=3B；02×8C=118⊕11B=03，03×8C=03⊕8C=8F。</p>
<div class="formula">\[\begin{aligned}b_0&=C6\oplus F5\oplus E0\oplus8C=33\oplus E0\oplus8C=D3\oplus8C=5F,\\b_1&=63\oplus A6\oplus3B\oplus8C=C5\oplus3B\oplus8C=FE\oplus8C=72,\\b_2&=63\oplus53\oplus DB\oplus8F=30\oplus DB\oplus8F=EB\oplus8F=64,\\b_3&=A5\oplus53\oplus E0\oplus03=F6\oplus E0\oplus03=16\oplus03=15.\end{aligned}\]</div>
<p>其余三列各自乘同一矩阵，得到下面状态。MixColumns是线性变换，名字中的“混淆”不表示它负责非线性Confusion；它主要扩散。AddRoundKey则把当前状态与128位轮密钥逐字节异或。</p>
'''+table('第四讲P39：第一轮的最后一步',['行','MixColumns后','rk₁','异或后的第一轮输出'],[['0','5F 57 F7 1D','D6 D2 DA D6','89 85 2D CB'],['1','72 F5 BE B9','AA AF A6 AB','D8 5A 18 12'],['2','64 BC 3B F9','74 72 78 76','10 CE 43 8F'],['3','15 92 29 1A','FD FA F1 FE','E8 68 D8 E4']])+r'''
<h3>6. 解密：严格按反方向撤销</h3>
<p>先与rk₁₀异或，撤销末轮密钥加；再InvShiftRows（各行右移）、InvSubBytes。然后对于r=9到1依次AddRoundKey(rkᵣ)、InvMixColumns、InvShiftRows、InvSubBytes。最后与rk₀异或恢复明文。逆S盒用于还原查表，逆列混淆矩阵如下。不同实现可以等价重排，但不能只把密钥倒序、仍原样运行AES加密轮。</p>
<div class="formula">\[\operatorname{InvMixColumns}:\quad\begin{pmatrix}0E&0B&0D&09\\09&0E&0B&0D\\0D&09&0E&0B\\0B&0D&09&0E\end{pmatrix}.\]</div>
'''+f'<p>AES状态按{cloze("列装入")}；末轮{cloze("没有MixColumns")}；GF(2⁸)加法是{cloze("异或，乘法还要模不可约多项式约减")}。</p>'+think('为什么03×53=F5，而不是普通整数乘法的F9？','<p>03表示x+1，53表示x⁶+x⁴+x+1。先02×53=A6，再把A6与53做异或，得到F5。GF(2)同次幂系数相加不进位，而普通整数乘法带进位，所以结果通常不同。</p>','多项式解释1分；02乘积1分；异或结果1分；说明无进位1分。')
add('t05','AES：状态矩阵、有限域与四种轮变换','第四讲','P33–39','从位、字节、行列坐标开始，理解AES-128流程，并手算列混淆的每个输出字节。',body,[
('AES输入前四字节在状态矩阵的哪里？',[opt('第0行。',False,'连续四个输入字节装入第0列。'),opt('第0列。',True,'s[r,c]=input[4c+r]。'),opt('主对角线。',False,'没有按对角线填充。'),opt('任意放置都不影响AES结果。',False,'布局决定行移位和列混淆作用对象。')],'矩阵显示方式必须与序列化约定一致。'),
('关于AES操作，哪些正确？',[opt('SubBytes改变字节值。',True,'位置保持，逐字节查表。'),opt('ShiftRows把字节内比特左移。',False,'移动的是整字节。'),opt('第10轮AES-128省略MixColumns。',True,'仍有SubBytes、ShiftRows和密钥加。'),opt('GF(2⁸)乘法就是整数乘法mod256。',False,'按多项式模11B约减。')],'值的代换、位置的移动和域内混合承担不同角色。'),
('GF(2⁸)中02×E0等于什么？',[opt('C0。',False,'只截断丢了约减项1B。'),opt('DB。',True,'1C0⊕11B=DB。'),opt('1C0。',False,'这是约减前9位结果，尚非字节。'),opt('E2。',False,'把乘法误当普通加法。')],'最高位为1时，左移后必须处理x⁸项。')])

start('t06')
hw=bytes(int(rawtable(55,2)[r][c],16) for c in range(4) for r in range(4))
hwkeys=aeskeys(hw)
body=r'''
<h3>1. 128位主密钥展开成44个字</h3>
<p>主密钥先分成4个32位字W₀到W₃，每个字对应状态矩阵的一列。再递推得到W₄到W₄₃。连续4个字组成一个128位轮密钥，共11个：rk₀=主密钥，rk₁到rk₁₀分别用于10轮。44不是轮数，11也不是AES-128的完整轮数。</p>
<div class="formula">\[W_i=\begin{cases}W_{i-4}\oplus T_{i/4}(W_{i-1}),&i\equiv0\pmod4,\\W_{i-4}\oplus W_{i-1},&i\not\equiv0\pmod4,\end{cases}\quad4\le i\le43.\]
\[T_j(w)=\operatorname{SubWord}(\operatorname{RotWord}(w))\oplus\operatorname{Rcon}[j].\]</div>
<p>RotWord循环左移<strong>一个字节</strong>；SubWord把4个字节分别查AES S盒；Rcon[j]=(RC[j],00,00,00)只影响第一个字节。先完成这三个操作得到临时字，再与Wᵢ₋₄异或；后续非特殊位置使用刚计算出的新字，不能一直拿原来的W₃。</p>
'''+table('第四讲P40：Rcon首字节，全部十六进制',['j']+list(range(1,11)),[['RC','01','02','04','08','10','20','40','80','1B','36']])+r'''
<h3 id="t06-aes-sbox">2. 完整AES S盒：查表前确认是哪种算法</h3>
<p>行、列均为十六进制，高半字节选行、低半字节选列。AES、SM4的查表方式相似，但表值不同；DES的索引方式则不同。</p>
'''+source_table(55,0,'第四讲P55：AES完整S盒')+r'''
<h3>3. 课堂原例：00010203…0F 的首轮轮密钥</h3>
<p>W₀=00010203、W₁=04050607、W₂=08090A0B、W₃=0C0D0E0F。对W₃先RotWord得0D0E0F0C，查表得D7AB76FE，再⊕01000000得D6AB76FE。</p>
<div class="formula">\[\begin{aligned}W_4&=00010203\oplus D6AB76FE=D6AA74FD,\\W_5&=04050607\oplus D6AA74FD=D2AF72FA,\\W_6&=08090A0B\oplus D2AF72FA=DAA678F1,\\W_7&=0C0D0E0F\oplus DAA678F1=D6AB76FE.\end{aligned}\]</div>
<p>rk₁=D6AA74FD D2AF72FA DAA678F1 D6AB76FE。每4字节竖排为一列，恰好得到上一章的rk₁矩阵。不要先按行拼接再套字递推。</p>
'''+table('核验补充：该主密钥的全部11个轮密钥',['轮密钥','按列拼接的16字节十六进制'],[[i,k.hex().upper()] for i,k in enumerate(aeskeys(bytes(range(16))))])+r'''
<h3>4. 课程作业1：先自己算，再展开</h3>
<p>第四讲P55的密钥以4×4矩阵给出。这里保留原来的行列位置；每列构成一个字。作业要求第一轮轮密钥。</p>
'''+table('P55原始密钥状态',['行','第0列','第1列','第2列','第3列'],[[i]+r for i,r in enumerate(rawtable(55,2))])+think('求P55矩阵对应的W₄到W₇及rk₁',r'''<p>按列读取：W₀=A12EF3A0，W₁=B53A5B7D，W₂=942517C9，W₃=D0445221。RotWord(W₃)=445221D0；逐字节查表44→1B、52→00、21→FD、D0→70，得1B00FD70；异或Rcon₁得1A00FD70。</p>
<div class="formula">\[\begin{aligned}W_4&=A12EF3A0\oplus1A00FD70=BB2E0ED0,\\W_5&=B53A5B7D\oplus BB2E0ED0=0E1455AD,\\W_6&=942517C9\oplus0E1455AD=9A314264,\\W_7&=D0445221\oplus9A314264=4A751045.\end{aligned}\]</div>
<p>所以rk₁=BB2E0ED0 0E1455AD 9A314264 4A751045。按行显示的矩阵是第一行BB 0E 9A 4A；第二行2E 14 31 75；第三行0E 55 42 10；第四行D0 AD 64 45。</p>''','按列读出四字2分；RotWord和SubWord各1分；Rcon异或1分；W₄至W₇各1分；输出矩阵方向1分。')+call('trap','易错陷阱 · 适用范围','本章递推仅针对AES-128。AES-192和AES-256的初始字数、总字数和特殊分支不同，不能只把轮数改大就认为实现完整。')+f'<p>AES-128扩展得到{cloze("44个32位字、11个128位轮密钥")}；RotWord移动{cloze("1字节，等于8位")}；确定性密钥编排{cloze("不增加主密钥的独立熵")}。</p>'
assert hwkeys[1].hex()=='bb2e0ed00e1455ad9a3142644a751045'
add('t06','AES 密钥扩展与原课件作业完整解答','第四讲','P40–41、P55','掌握字递推、RotWord/SubWord/Rcon，并完成课件中两个不同主密钥的首轮展开。',body,[
('AES-128初始密钥加加上10轮，一共需要几个128位轮密钥？',[opt('10个。',False,'漏掉初始rk0。'),opt('11个。',True,'rk0到rk10，共44个32位字。'),opt('44个。',False,'44是字数而非128位轮密钥数。'),opt('128个。',False,'128表示主密钥比特长度。')],'不要把字数、轮密钥数与轮数混用。'),
('P55作业按列读出的W₃是什么？',[opt('A07DC921。',False,'这是矩阵最后一行。'),opt('D0445221。',True,'最后一列依次D0、44、52、21。'),opt('A1B594D0。',False,'这是第一行。'),opt('214452D0。',False,'颠倒了列中字节顺序。')],'先固定状态的列方向，再计算。'),
('哪些操作正确？',[opt('RotWord(0C0D0E0F)=0D0E0F0C。',True,'循环左移一字节。'),opt('Rcon会对四字节都异或01。',False,'第一轮Rcon是01000000。'),opt('W₅=W₁⊕W₄。',True,'不是W1⊕W3。'),opt('SubWord使用SM4的S盒也一样。',False,'两种密码的S盒不同。')],'轮密钥的依赖顺序必须逐步推进。')])

start('t07')
body=r'''
<h3>1. SM4 的单位和结构</h3>
<p>SM4公开时称SMS4，课件列出2006年公开、2012年行业标准GM/T 0002、2016年国标GB/T 32907及2021年纳入ISO/IEC 18033-3的沿革。它是一种128位分组、128位主密钥、32轮的对称分组密码。主密钥和明文各分成4个32位字：MK₀‖MK₁‖MK₂‖MK₃、X₀‖X₁‖X₂‖X₃。</p>
<p>SM4采用四分支广义Feistel式结构。每轮保留后三个字，计算新字替换最早的一个。把状态想成长度为4的滑动窗口：第一轮从X₀…X₃变为X₁…X₄，第二轮变为X₂…X₅，直到X₃₂…X₃₅。轮号i从0开始，课件“第1轮”对应i=0、使用rk₀。</p>
<div class="formula">\[\begin{aligned}A_i&=X_{i+1}\oplus X_{i+2}\oplus X_{i+3}\oplus rk_i,\\B_i&=\tau(A_i),\quad T(A_i)=L(B_i),\\X_{i+4}&=X_i\oplus T(A_i),\qquad0\le i\le31,\\C&=X_{35}\parallel X_{34}\parallel X_{33}\parallel X_{32}.\end{aligned}\]</div>
<p>反序输出只交换四个字的位置，不翻转字内部的比特或字节。课件P46把合成置换结果记Bᵢ；这里为清楚展示中间步骤，把Bᵢ用作τ输出，T(Aᵢ)另写，计算含义相同。</p>
'''+diagram('flowchart LR\n A["后三字与轮密钥异或"] --> S["τ：四字节各查SM4 S盒"]\n S --> L["L：32位循环移位后异或"]\n L --> X["与旧第一字异或"]\n X --> N["生成新字，窗口前移"]','每轮只生成一个新32位字；其余三个字保留。32轮后四字反序输出。')+r'''
<h3>2. τ 与完整SM4 S盒</h3>
<p>32位拆成从高到低4字节，每字节高4位定行、低4位定列。四个盒使用同一张表，和DES八张不同的盒区别开。例82→8A、83→D2、CB→41、69→22，所以τ(8283CB69)=8AD24122。不要拿AES盒查这些字节。</p>
'''+source_table(47,0,'第四讲P47–48：SM4完整S盒')+r'''
<h3>3. L：让一个字内部的影响扩散</h3>
<div class="formula">\[L(B)=B\oplus\operatorname{ROL}_{32}(B,2)\oplus\operatorname{ROL}_{32}(B,10)\oplus\operatorname{ROL}_{32}(B,18)\oplus\operatorname{ROL}_{32}(B,24).\]</div>
<p>ROL₃₂把移出高端的位送回低端，始终保持32位；逻辑左移则丢弃高位、低端补0。原例B=00000001，五项依次为00000001、00000004、00000400、00040000、01000000，异或得到01040405。只有1个1的输入变成分布在4字节中的5个1。L是线性层，τ是非线性层，两者合成T。</p>
<h3>4. 密钥编排：相似窗口，不同线性变换</h3>
<p>先用公开FK常量与主密钥四字分别异或。之后每轮用公开CKᵢ、同一τ，但把L换成L′。新生成的Kᵢ₊₄就是本轮rkᵢ；不能把加密状态X和编排状态K混成一个数组。</p>
<div class="formula">\[\begin{aligned}K_j&=MK_j\oplus FK_j\quad(0\le j\le3),\\rk_i=K_{i+4}&=K_i\oplus T'(K_{i+1}\oplus K_{i+2}\oplus K_{i+3}\oplus CK_i),\\T'(A)&=L'(\tau(A)),\quad L'(B)=B\oplus\operatorname{ROL}_{32}(B,13)\oplus\operatorname{ROL}_{32}(B,23).\end{aligned}\]</div>
'''+table('第四讲P51：初始化FK',['j','FKⱼ'],[[i,f'{v:08X}'] for i,v in enumerate(FK)])+r'''
<p>CKᵢ的四字节定义为ckᵢ,ⱼ=((4i+j)×7) mod256，j=0到3，按高到低拼接。i=1得到28、35、42、49（十进制），对应1C、23、2A、31，所以CK₁=1C232A31。CK₀=00070E15，CK₃₁=646B7279。所有FK、CK和S盒都公开；秘密只有主密钥。</p>
<h3>5. 解密为什么仍用同一个T</h3>
<p>加密最后一轮X₃₅=X₃₁⊕T(X₃₂⊕X₃₃⊕X₃₄⊕rk₃₁)。密文已经给出X₃₅、X₃₄、X₃₃、X₃₂；再算相同T并与X₃₅异或，就恢复X₃₁。XOR交换律允许后三字反序出现，不必求T逆。</p>
<div class="formula">\[Y_0,Y_1,Y_2,Y_3=X_{35},X_{34},X_{33},X_{32},\]
\[Y_{j+4}=Y_j\oplus T(Y_{j+1}\oplus Y_{j+2}\oplus Y_{j+3}\oplus rk_{31-j}),\quad0\le j\le31.\]
\[(Y_{35},Y_{34},Y_{33},Y_{32})=(X_0,X_1,X_2,X_3).\]</div>
'''+f'<p>SM4每轮生成{cloze("1个32位新字")}；输出{cloze("最后四字反序，字内部不变")}；密钥编排使用{cloze("L′的13、23位循环移位，不是L的2、10、18、24")}；解密{cloze("同结构、轮密钥逆序")}。</p>'+think('推导CK₃₁，解释为何常量不必保密','<p>四字节是124×7 mod256=100=64；125×7 mod256=107=6B；126×7 mod256=114=72；127×7 mod256=121=79。因此CK₃₁=646B7279。常量用于区分轮次，不承担秘密角色；把公开常量隐藏也不能替代随机主密钥。</p>','四字节计算各1分；按顺序拼接1分；区分常量和秘密1分。')
add('t07','SM4：四字窗口、S 盒、线性层与密钥编排','第四讲','P42–52','逐个区分字、字节、循环移位、公开常量与轮密钥，推导加密和解密的一致结构。',body,[
('SM4最终密文输出顺序是什么？',[opt('X₃₂‖X₃₃‖X₃₄‖X₃₅。',False,'漏掉最后的四字反序。'),opt('X₃₅‖X₃₄‖X₃₃‖X₃₂。',True,'字内部的字节次序不改。'),opt('每个字的32位也反转。',False,'反序变换针对四个字。'),opt('仅输出X₃₅。',False,'密文128位，需要4个32位字。')],'窗口末状态与密文序列不同。'),
('SM4的L与L′，哪些正确？',[opt('L用于加密轮，移位2、10、18、24。',True,'还包含未移位的B。'),opt('L′用于编排，移位13、23。',True,'同样与B自身异或。'),opt('二者都调用AES的S盒。',False,'τ查SM4 S盒。'),opt('解密必须先求L的逆。',False,'广义Feistel逆推只需重算T再异或。')],'两个线性层不同，整体解密也不是逐项逆T。'),
('τ(8283CB69)等于什么？',[opt('8AD24122。',True,'82、83、CB、69依次查SM4表。'),opt('01040405。',False,'这是L(00000001)。'),opt('F12186F9。',False,'这是标准例rk0。'),opt('8283CB69。',False,'没有执行代换。')],'逐字节查表，保留高到低拼接顺序。')])

start('t08')
body=r'''
<h3>1. 国标示例：明文与主密钥恰好相同</h3>
<p>第四讲P53使用X=MK=01234567 89ABCDEF FEDCBA98 76543210。它是用于检验实现的固定测试值，不表示日常加密要把明文当密钥。按32位字分割，全部十六进制。</p>
<p>密钥初始化：01234567⊕A3B1BAC6=A292FFA1；89ABCDEF⊕56AA3350=DF01FEBF；FEDCBA98⊕677D9197=99A12B0F；76543210⊕B27022DC=C42410CC。首轮编排输入DF01FEBF⊕99A12B0F⊕C42410CC⊕00070E15=8283CB69，经τ得8AD24122。</p>
<div class="formula">\[\begin{aligned}\operatorname{ROL}_{32}(8AD24122,13)&=4824515A,\\\operatorname{ROL}_{32}(8AD24122,23)&=91456920,\\L'(8AD24122)&=8AD24122\oplus4824515A\oplus91456920=53B37958,\\rk_0&=A292FFA1\oplus53B37958=F12186F9.\end{aligned}\]</div>
<p>加密第一轮：X₁⊕X₂⊕X₃=01234567，再⊕rk₀=F002C39E；四字节代换F0→18、02→E9、C3→92、9E→B1，得18E992B1。L所需循环移位分别是63A64AC4、A64AC463、4AC463A6、B118E992。</p>
<div class="formula">\[\begin{aligned}L(18E992B1)&=18E992B1\oplus63A64AC4\oplus A64AC463\\&\quad\oplus4AC463A6\oplus B118E992=26D99622,\\X_4&=01234567\oplus26D99622=27FAD345.\end{aligned}\]</div>
<h3>2. 全部32轮：不给中间步骤留省略号</h3>
<p>下表每行按上一章公式执行：Aᵢ=后三字⊕rkᵢ，τ(Aᵢ)查表，T(Aᵢ)=L(τ(Aᵢ))，Xᵢ₊₄=Xᵢ⊕T(Aᵢ)。表中的CK参与同轮轮密钥编排，不是加密轮直接异或的量。可任挑一行按公式复算。</p>
'''+table('国标示例32轮完整核验；所有字均8位十六进制',['i','CKᵢ','rkᵢ','Aᵢ','τ(Aᵢ)','T(Aᵢ)','Xᵢ₊₄'],smrows)+r'''
<p>末状态X₃₂=536E4246、X₃₃=86B3E94F、X₃₄=D206965E、X₃₅=681EDF34。反序输出C=681EDF34 D206965E 86B3E94F 536E4246。把同一密文送入同结构、使用rk₃₁到rk₀，应精确得到初始明文。</p>
<h3>3. 三个分组密码的比较</h3>
'''+source_table(54,0,'第四讲P54：DES、AES、SM4完整对照')+call('correction','应用边界 · 算法名称不能代替完整协议','SM4适合批量加密，SM2用于相应签名或密钥协商机制，SM3提供摘要。单独SM3不能认证来源；实际防篡改还需要带密钥的认证或签名。SM4也需要合适的工作模式。')+r'''
<h3>4. 课程作业2：对称轮密钥让加密器变成解密器</h3>
<p>设某个<strong>SM4变体</strong>保留轮函数与反序输出，但对任意主密钥都满足rk₃₁₋ᵢ=rkᵢ（i=0到15）。证明其加密与解密相同，并考虑攻击者能提交任意一个分组给同钥加密接口的情况。</p>
'''+think('为什么一次选择明文请求就能解密截获的C？','<p>标准结构中Dec<sub>K</sub>调用相同变换，只把轮密钥序列逆序。变体序列正反相同，因此Dec<sub>K</sub>=Enc<sub>K</sub>。由正确性Enc<sub>K</sub>(Enc<sub>K</sub>(M))=M，这个变换是对合。攻击者把截获的密文C当作“待加密明文”提交一次，得到Enc<sub>K</sub>(C)=Dec<sub>K</sub>(C)=M。</p><p>条件是同一把密钥、同一分组变换接口、确实允许把C作为合法输入。这里讨论课件刻意设定的缺陷变体；不能声称标准SM4对任意主密钥都存在此缺陷，也不能跳过模式与协议输入限制把攻击直接套到所有系统。</p>','证明逆序序列等于原序列2分；推出Enc=Dec及对合2分；提交C得到M2分；说清同钥接口与变体条件1分。')+f'<p>标准示例首轮rk₀={cloze("F12186F9")}、X₄={cloze("27FAD345")}；对称轮密钥变体的攻击请求是{cloze("把截获密文C提交给同钥加密接口")}。</p>'
add('t08','SM4 完整32轮验算与选择明文攻击作业','第四讲','P53–57','从首轮位运算到完整密文，核对32轮数据，并证明缺陷密钥编排的攻击。',body,[
('SM4国标示例密文第一字为什么是681EDF34？',[opt('它就是明文X₀。',False,'明文首字是01234567。'),opt('它是最后生成的X₃₅，因反序输出排在第一。',True,'不是按X32到X35正序输出。'),opt('它是第一个轮密钥。',False,'rk0是F12186F9。'),opt('只需要算一轮就得到它。',False,'这是32轮后结果。')],'轮状态、轮密钥和输出字序分别跟踪。'),
('对于课件的对称轮密钥SM4变体，哪些成立？',[opt('同钥下Enc=Dec。',True,'正反轮密钥序列相同。'),opt('对截获C请求一次同钥加密可得到M。',True,'Enc(C)=Dec(C)。'),opt('证明了所有标准SM4密钥都如此。',False,'课件讨论缺陷变体的额外条件。'),opt('必须先恢复128位主密钥。',False,'直接使用加密接口，不必恢复密钥。')],'攻击利用接口能力和结构缺陷，不是穷举。'),
('DES、AES、SM4对照哪项正确？',[opt('三者分组均128位。',False,'DES是64位。'),opt('AES解密只需原样使用加密轮。',False,'AES还需要相应逆部件。'),opt('SM4有32轮，每轮使用32位轮密钥。',True,'主密钥与分组均128位。'),opt('SM4轮数更多就必然比AES更安全。',False,'轮复杂度和结构不同，不能只比轮数。')],'算法参数要在具体结构下理解。')])

start('m01')
body=r'''
<h3>1. 为什么需要模式和填充</h3>
<p>分组密码只承诺处理一个固定长度块。第五讲在这个核心外面加上组织规则，处理文件、图片和网络消息。记Mᵢ为第i个明文块、Cᵢ为第i个密文块，E<sub>K</sub>和D<sub>K</sub>分别是一次<strong>完整</strong>分组加解密，不是AES的一轮。</p>
<p>ECB、普通CBC要求输入长度为整块数。若末块不足，需要无歧义填充。以按字节的PKCS#7为理解补充：块长b字节，缺p字节就补p个值为p的字节。AES的b=16；若消息14字节，补02 02；若正好16字节，仍需补整块16个10（十六进制），防止把数据末尾误认成填充。末尾补0若无原长约定，会无法区分原文0与填充0。</p>
<h3>2. ECB：同样的块得到同样的密文</h3>
<div class="formula">\[C_i=E_K(M_i),\qquad M_i=D_K(C_i).\]</div>
<p>各块独立，可并行加解密。固定密钥的双射意味着Mᵢ=Mⱼ当且仅当Cᵢ=Cⱼ。即使不知道每块内容，也能认出重复位置。攻击者还可以复制、删除、交换整块，接收者仍逐块解密；没有额外认证，就不知道被改过。</p>
<p>课件图示的“图像仍可看出轮廓”是在每格代表一个相同大小明文块、保持块位置的教学模型。以下SVG把相等的块编码成相同图案，展示重复结构泄露；它不是对真实图片文件执行AES的实测截图。真实图像头部、编码压缩也影响外观。</p>
<figure class="diagram" id="ecb-pattern"><svg viewBox="0 0 650 235" role="img" aria-labelledby="ecb-title"><title id="ecb-title">ECB重复块泄露示意：相等明文块映射成相同密文块，保留空间结构</title><defs><pattern id="block-a" width="12" height="12" patternUnits="userSpaceOnUse"><rect width="12" height="12" fill="#d8e9ed"/><path d="M0 0L12 12" stroke="#087e8b" stroke-width="3"/></pattern><pattern id="block-b" width="10" height="10" patternUnits="userSpaceOnUse"><rect width="10" height="10" fill="#e6d9f5"/><circle cx="5" cy="5" r="3" fill="#7854a5"/></pattern></defs><g fill="#14343b" font-size="16" font-family="sans-serif"><text x="20" y="25">明文：两种重复块</text><text x="340" y="25">ECB：相等块仍相等</text></g><g id="pattern-left"></g><g id="pattern-right"></g></svg><figcaption class="caption">重绘第五讲P7的核心含义；CBC的随机外观不等于具备消息认证。</figcaption></figure>
<h3>3. CBC：先异或，再加密</h3>
<div class="formula">\[C_0=IV,\qquad C_i=E_K(M_i\oplus C_{i-1}),\qquad M_i=D_K(C_i)\oplus C_{i-1}.\]</div>
<p>第一块没有前块，使用初始化向量IV，它与分组等长。IV一般公开随密文发送，但应每次按要求新生成且在选定消息时不可预测，并保护其完整性。解密不要求IV保密，要求双方用同一个正确IV。CBC加密要等Cᵢ₋₁，所以串行；密文全已获得时，各D<sub>K</sub>(Cᵢ)可同时算，再与已知前块异或，因此可并行解密。</p>
'''+diagram('flowchart LR\n I["IV等于C0"] --> X1["M1与C0异或"]\n X1 --> E1["E_K"]\n E1 --> C1["C1"]\n C1 --> X2["M2与C1异或"]\n X2 --> E2["E_K"]\n E2 --> C2["C2"]','CBC中反馈的是上一个密文，异或发生在分组加密之前。解密先D_K，再异或前一密文。')+call('trap','易错陷阱 · CBC并不保证相等块永不碰撞','同一消息内，即使Mᵢ=Mⱼ，两个前块一般不同，所以CBC通常隐藏ECB式重复关系；但不能断言所有位置密文必不同。若Mᵢ⊕Cᵢ₋₁=Mⱼ⊕Cⱼ₋₁，双射仍给相同密文。同钥同IV重复加密相同消息，则整个密文重复。')+r'''
<h3>4. 课堂题：4位玩具密码完整计算</h3>
<p>E<sub>K</sub>(X)=ROL₄(X⊕K,1)，K=1010，IV=0110，M=1001‖1001‖0101。先异或再4位循环左移；例如0101循环左移1位变1010。逆函数D<sub>K</sub>(Y)=ROR₄(Y,1)⊕K。</p>
'''+table('第五讲P13–14：课堂题补全',['i','Mᵢ','ECB：Mᵢ⊕K','ECB：Cᵢ','CBC：Mᵢ⊕Cᵢ₋₁','CBC：再⊕K','CBC：Cᵢ'],[[1,'1001','0011','0110','1111','0101','1010'],[2,'1001','0011','0110','0011','1001','0011'],[3,'0101','1111','1111','0110','1100','1001']])+r'''
<p>CBC第二块：1001⊕1010=0011，再⊕1010=1001，ROL₄得0011。第三块0101⊕0011=0110，再⊕1010=1100，ROL₄得1001。反验首块：D(1010)=0101⊕1010=1111；再⊕IV 0110得到1001。</p>
'''+f'<p>ECB暴露{cloze("相同位置块的相等关系")}；CBC加密时{cloze("明文先异或前一密文，再调用E_K")}；CBC的IV需要{cloze("不可预测，并保护完整性")}。</p>'+think('为什么篡改CBC的IV能定向改变第一块明文？',r'<p>第一块M₁=D<sub>K</sub>(C₁)⊕IV。若攻击者替换IV′=IV⊕Δ、密文不变，则M₁′=D<sub>K</sub>(C₁)⊕IV⊕Δ=M₁⊕Δ。攻击者不需要密钥就能定向翻位。这说明IV公开可接受，但完整性保护必不可少。</p>','写出解密式1分；代入IV′并化简2分；说明公开与完整性是不同要求1分。')
add('m01','ECB 与 CBC：长消息、填充与重复泄露','第五讲','P1–8、P13–14','用四位玩具密码学习完整分组调用、CBC链式依赖与填充规则，理解图像轮廓为何泄露。',body,[
('课堂题CBC的C₁、C₂、C₃是什么？',[opt('0110、0110、1111。',False,'这是ECB输出。'),opt('1010、0011、1001。',True,'每块先异或前块，再调用玩具密码。'),opt('1111、0011、0110。',False,'这是CBC进入E前的中间量。'),opt('0101、1001、1100。',False,'尚未执行最后循环左移。')],'每一层中间值都应清楚标名。'),
('关于CBC，哪些正确？',[opt('密文全部已知时可并行解密。',True,'前一密文是已知输入，不必等前块解密结果。'),opt('CBC能自动发现恶意篡改。',False,'它只提供相应条件下的保密，未提供认证。'),opt('IV可以公开，但应满足不可预测性并受认证保护。',True,'公开与可修改是不同概念。'),opt('相同明文块保证得到互不相同的CBC密文块。',False,'不是绝对保证，异或后的分组输入仍可能相等。')],'链式依赖隐藏重复的条件性，不能替代认证。'),
('AES采用PKCS#7，明文长度正好16字节时怎样填充？',[opt('不补任何字节。',False,'该填充约定需要整块填充以避免歧义。'),opt('补16个十六进制10。',True,'每个字节值16，整块补齐。'),opt('补1个00。',False,'长度仍非整块且规则错误。'),opt('删去末字节。',False,'会破坏消息。')],'这是PKCS#7的具体约定，不是所有模式都需要填充。')])

start('m02')
body=r'''
<h3>1. 三种模式都先产生掩码，再异或消息</h3>
<p>CFB、OFB、CTR把完整E<sub>K</sub>当作密钥流生成部件。消息只参与异或，因此解密仍调用E<sub>K</sub>，不用D<sub>K</sub>。但“喂给E的输入是什么”决定依赖关系与安全要求。以下CFB按课件采用<strong>整块反馈CFB-n</strong>；CFB-8、CFB-1等分段变体要另外分析移位寄存器。</p>
<h3>2. CFB：加密前一个密文</h3>
<div class="formula">\[C_0=IV,\qquad C_i=M_i\oplus E_K(C_{i-1}),\qquad M_i=C_i\oplus E_K(C_{i-1}).\]</div>
<p>第一步E(IV)得到第一段掩码，第二步E(C₁)得到第二段。反馈的是密文，所以加密依赖前一块结果，不能提前生成完整密钥流。密文已知时可并行解密。IV应新生成且不可预测，同钥复用IV会重复首段掩码，泄露首段明文异或。</p>
<h3>3. OFB：加密前一次输出</h3>
<div class="formula">\[O_0=IV,\quad O_i=E_K(O_{i-1}),\quad C_i=M_i\oplus O_i,\quad M_i=C_i\oplus O_i.\]</div>
<p>O₂=E(E(IV))，反馈完全不看明文或密文，所以可在消息到达前生成。密钥流内部仍是串行：O₂要等O₁。已有O序列时，异或消息可并行。同一密钥下IV必须唯一，且应避免一个消息的内部O值被另一个消息当IV，防止密钥流重叠。通常采用规范规定的IV生成方法。</p>
<h3>4. CTR：加密互不重复的计数器块</h3>
<div class="formula">\[S_i=E_K(\operatorname{ctr}_i),\quad C_i=M_i\oplus S_i,\quad M_i=C_i\oplus S_i.\]</div>
<p>ctrᵢ与分组等长，常由nonce（每消息唯一量）与块计数器拼成。各块直接算E(ctrᵢ)，可并行、预计算、随机访问。接收第100块时，知道该块计数器就可直接解密。计数器不必秘密；关键是同一K下所有完整计数器输入不得重复，约束横跨不同消息。计数器溢出回绕、两个消息使用重叠计数范围，都会重用密钥流。</p>
'''+diagram('flowchart LR\n A["独立ctr1"] --> E1["E_K"]\n E1 --> S1["密钥流S1"]\n S1 --> X1["与M1异或得到C1"]\n B["独立ctr2"] --> E2["E_K"]\n E2 --> S2["密钥流S2"]\n S2 --> X2["与M2异或得到C2"]','CTR各分组的E_K输入可独立构造；不得在同钥下复用任何计数器块。')+r'''
<h3>5. 最后不足一块怎么办</h3>
<p>若剩下r位，取该段密钥流最前r位与明文异或即可，无须补足整个明文块。CFB整块反馈采用课件所述“允许最后一段截短”的约定；其他标准或API可能要求整段输入，应遵循具体接口。仅最后一段可截短时，不能把截短密文当完整反馈块继续后续处理。</p>
'''+table('第五讲P12：完整对照，CFB为整块反馈',['模式','E_K输入','加密并行？','解密并行？','初始量条件','末段'],[['ECB','当前明文','是','是','无IV','按约定填充'],['CBC','当前明文⊕前一密文','否','是','IV不可预测','按约定填充'],['CFB-n','前一密文或IV','否','是（密文已知）','IV不可预测','截取最终所需位，遵循接口约定'],['OFB','上一次输出或IV','密钥流否','密钥流否','IV不重复且避免流重叠','截取所需位'],['CTR','当前计数器块','是','是','同钥下所有块计数器唯一','截取所需位']])+call('exam','考点提示 · 用一个问题辨认模式','问“E<sub>K</sub>究竟在加密什么？”CBC加密的是明文与前块异或的结果；CFB加密前一密文；OFB加密前一次输出；CTR加密计数器。名字中都有反馈，反馈对象却不同。')+f'<p>CFB/OFB/CTR解密都调用{cloze("E_K，不调用D_K")}；OFB的反馈对象是{cloze("上一次E_K输出")}；CTR最重要的约束是{cloze("同钥下完整计数器块全局不重复")}。</p>'+think('CTR中只保证每条消息内部计数器不重复，为什么仍不安全？','<p>设两条消息都从ctr=0开始并使用同一K，那么首段密钥流均为E<sub>K</sub>(0)，第二段均为E<sub>K</sub>(1)。两密文异或时掩码抵消，得到对应明文异或。可为每条消息安排唯一nonce并划分计数范围，同时防止计数器回绕；具体格式应由协议定义。</p>','指出跨消息重复2分；写出掩码抵消2分；说明nonce与回绕约束1分。')
add('m02','CFB、OFB、CTR：反馈对象与 IV 条件','第五讲','P9–12','比较三种密钥流模式，区分串行生成、可并行解密和跨消息的唯一性约束。',body,[
('哪种模式把上一次E_K的输出再次送进E_K？',[opt('CBC。',False,'CBC先把当前明文与前一密文异或。'),opt('CFB。',False,'CFB反馈密文。'),opt('OFB。',True,'O_i=E_K(O_{i-1})。'),opt('CTR。',False,'CTR加密独立构造的计数器块。')],'通过E的输入判断模式。'),
('哪些约束正确？',[opt('CBC与CFB的IV应不可预测。',True,'单纯可见递增值通常不能满足该要求。'),opt('CTR可以每条消息都从0开始，只要消息内部不重。',False,'同钥跨消息也必须避免重复。'),opt('OFB同钥下IV不得重复。',True,'否则产生重复密钥流。'),opt('CFB解密要调用D_K。',False,'CFB两端都用E_K产生掩码。')],'IV一般公开，其生成规则仍关系保密性。'),
('需独立读取文件中第100块时，哪种模式最直接？',[opt('CTR：构造该块计数器并算E_K。',True,'无需先计算前99段密钥流。'),opt('OFB：每块密钥流互不依赖。',False,'OFB流是串行反馈生成的。'),opt('CBC加密：无需前一密文。',False,'CBC加密依赖前一密文。'),opt('CFB加密：可随意跳过前99块。',False,'产生密文前需对应反馈输入。')],'“可提前生成”不自动表示“任意段能独立生成”。')])

start('m03')
body=r'''
<h3>1. 用同一个玩具密码做作业</h3>
<p>第五讲P45：E<sub>K</sub>(X)=ROL₄(X⊕K,1)，K=1010，消息0011‖0011‖1100。CBC取IV=0101；CTR取ctr₁=0001，后续为0010、0011。左循环宽度严格为4位，不是无限长整数左移。</p>
'''+think('先独立完成CBC与CTR的全部中间量','<p>CBC第一块0011⊕0101=0110，再⊕1010=1100，左循环得1001。第二块0011⊕1001=1010，再⊕1010=0000，左循环仍0000。第三块1100⊕0000=1100，再⊕1010=0110，左循环得1100。</p><p>CTR三个掩码：E(0001)=ROL(1011)=0111；E(0010)=ROL(1000)=0001；E(0011)=ROL(1001)=0011。分别与0011、0011、1100异或，得到0100、0010、1111。</p>'+table('P45作业完整答案',['i','Mᵢ','CBC进入E前','CBC Cᵢ','ctrᵢ','CTR密钥流','CTR Cᵢ'],[[1,'0011','0110','1001','0001','0111','0100'],[2,'0011','1010','0000','0010','0001','0010'],[3,'1100','1100','1100','0011','0011','1111']]),'CBC三块各2分；CTR掩码和密文每块各1分；保留4位宽度。')+r'''
<h3>2. 出错的是密文某一位，哪些明文受影响</h3>
<p>先固定条件：只翻转Cⱼ的一位，不删除或插入数据，密钥与初始量不变，分组边界保持同步，且暂时不做认证检查。令Δ表示只有该位为1的差分。</p>
'''+table('单个密文比特翻转的传播范围（实际强分组密码的典型表现）',['模式','当前块Mⱼ','后续影响'],[['ECB','该块通常约一半位变化','其他块不变'],['CBC','该块通常约一半位变化','下一块对应位定向翻转，再后正常'],['CFB-n','该块对应位定向翻转','下一块通常约一半位变化，再后正常'],['OFB','对应位定向翻转','其他位不变'],['CTR','对应位定向翻转','其他位不变']])+r'''
<p>CBC：Mⱼ′=D(Cⱼ⊕Δ)⊕Cⱼ₋₁，当前块差分取决于D；下一块Mⱼ₊₁′=D(Cⱼ₊₁)⊕(Cⱼ⊕Δ)=Mⱼ₊₁⊕Δ。后面的前块没被改，恢复正常。CFB方向相反：当前块直接与密文异或，所以当前对应位翻；下一块用被改Cⱼ作为E输入，出现扩散。OFB/CTR的流不依赖密文，因此只影响被翻那一位。</p>
<h3>3. 原作业的玩具密码不会“约半块损坏”</h3>
<p>玩具D(C)=ROR₄(C,1)⊕K是线性的，翻1位只会在D输出中移动1位。CBC作业C₂=0000最高位翻转为1000：D(0000)=1010，D(1000)=0100⊕1010=1110。与C₁=1001异或，M₂′=0111，原M₂=0011，所以差分0100，即从左第2位翻转。M₃′=1100⊕1000=0100，差分1000，即最高位翻转。M₁不变。</p>
<p>CTR的C₂=0010变为1010，掩码仍0001，M₂′=1011，与原0011差1000；其他块不变。若换成AES等强分组密码，CBC当前块通常发生广泛变化，下一块仍只有对应最高位翻转；CTR仍仅一位翻转。这里“约一半”是典型统计预期，不是每次必须精确一半。</p>
<div id="modes-lab" class="native-lab"></div>
'''+call('correction','安全边界 · 错误传播不是完整性保护','把某块解成乱码不代表系统一定发现攻击。攻击者可能利用定向比特翻转，或利用系统暴露的填充错误形成攻击面。完整保护需要认证标签和正确的验证顺序，例如采用合适的认证加密方案。GCM在第四讲仅列举为模式，本两讲没有展开其内部推导。')+f'<p>CBC中翻Cⱼ的一位会使下一明文块{cloze("同一位置位定向翻转")}；CTR中则{cloze("只翻对应明文位")}；“出现乱码”{cloze("不等于检测到恶意篡改")}。</p>'+think('为什么不能把CFB的错误传播结论无条件套到CFB-8？','<p>本表使用整块CFB-n，每次把完整Cⱼ送进E，所以只污染下一整块的掩码。CFB-8每次把一个字节移入n位寄存器，一个错误字节会在寄存器中停留多次调用，影响范围取决于分组长度与反馈段长。先写清模式参数，再数它经过多少次移位才离开。</p>','说明整块假设1分；分段移位寄存器记忆2分；按段长确定范围1分。')
add('m03','模式演算实验、错误传播与完整作业','第五讲','P13–14、P33、P45','完成CBC/CTR作业，比较玩具算法与真实分组密码的错误传播，用交互实验验证差异。',body,[
('P45作业CBC密文与CTR密文分别是？',[opt('1001 0000 1100；0100 0010 1111。',True,'按IV和起始计数器逐块计算。'),opt('0100 0010 1111；1001 0000 1100。',False,'颠倒两模式。'),opt('1001 1001 1100；0111 0001 0011。',False,'前者反馈计算错误，后者仅为密钥流。'),opt('0011 0011 1100；相同。',False,'没有执行加密。')],'掩码与密文分开记录。'),
('CBC密文C₂最高位翻转，哪些判断正确？',[opt('M₃最高位会定向翻转。',True,'前块异或直接把差分带入下一明文。'),opt('所有后续明文都会损坏。',False,'只改一块密文且对齐未变时，M4起恢复。'),opt('玩具密码下M₂从左第2位翻转。',True,'4位右循环把1000差分变0100。'),opt('AES下M₂必定精确翻64位。',False,'约一半是统计趋势，不保证精确数值。')],'区分代数必然结论和实际密码的统计扩散。'),
('哪种说法说明了完整性问题？',[opt('CTR一位密文篡改会被自动拒绝。',False,'没有标签时不会自动发现。'),opt('CBC出现错误传播就相当于MAC。',False,'乱码不等于验证。'),opt('CTR中c′=c⊕Δ会使m′=m⊕Δ。',True,'攻击者可定向修改而无需知道密钥流。'),opt('填充消除了所有篡改。',False,'填充是长度约定，不是认证。')],'正确解密与验证来源完整性是不同功能。')])

start('f01')
body=r'''
<h3>1. 一次一密怎样变成可部署的流密码</h3>
<p>一次一密要求密钥流均匀随机、与消息独立、保密、至少与消息等长、绝不复用。满足条件可得到信息论保密，但密钥分发成本很高。课件“目前只有一次性密码本无条件安全”需收窄：OTP是典型实例，并非所有信息论安全构造的唯一形式。</p>
<p>流密码以较短的秘密种子和公开IV初始化确定性生成器，输出很长的密钥流。两端输入相同、运行顺序相同，就得到相同流。现代密钥流需要密码学不可预测性；0/1大致均衡、周期很长、统计检验通过都只是部分要求，不能保证安全。一个简单交替0101…也很均衡，却完全可预测。</p>
<div class="formula">\[c_i=m_i\oplus k_i,\qquad c_i\oplus k_i=m_i\oplus(k_i\oplus k_i)=m_i.\]</div>
<p>这里种子密钥小写k与密钥流位kᵢ角色不同。流密码的长输出依然受短种子能提供的计算安全限制；它不会自动继承OTP的信息论保证。ZUC每拍输出32位，再把这些字展开成比特流。</p>
<h3>2. 复用流的代数灾难</h3>
<div class="formula">\[c=m\oplus s,\quad c'=m'\oplus s\quad\Longrightarrow\quad c\oplus c'=m\oplus m'.\]
\[m\text{ 已知时：}\quad s=c\oplus m,\qquad m'=c'\oplus s.\]</div>
<p>已知一个位置的明文只能恢复对应位置的密钥流，不能据此直接断言整个安全生成器种子已泄露；但另一条复用同段流的明文已无法保密。同钥不同IV若保证规范要求的唯一性，可避免生成相同流。这里风险是“同一段流复用”，不是“秘密种子只允许加密一条消息”。</p>
<h3>3. 课件LCG反例需要条件</h3>
<p>线性同余Xₙ₊₁=(aXₙ+c) modm不适合直接做密钥流。若m已知且相邻输出差可逆，几个输出便可恢复a、c。若m未知，可构造差分tᵢ=Xᵢ₊₁−Xᵢ、uᵢ=tᵢ₊₂tᵢ−tᵢ₊₁²，真实m整除各uᵢ，再用多个uᵢ的最大公因数得到m候选。四个输出仅提供一个u，不能普遍唯一恢复m；u=0、差分不可逆或截断输出都会改变分析。完整基础推导见<a href="#r02">第三讲LCG章节</a>。</p>
<h3>4. 分组密码与流密码的完整对照</h3>
'''+table('第五讲P19：两类对称密码的比较',['比较项','分组密码','二元加法流密码'],[['处理单位','固定长度分组，例如128位','消息按位异或，核心可按字输出，例如ZUC32位'],['核心部件','多轮轮函数、S盒、线性层','密钥流生成器，例如LFSR与非线性F'],['加密与解密','整体互逆，具体结构可复用','同一个异或操作'],['错误传播','由工作模式决定','保持同步时一位传输错误对应一位明文错误'],['主要风险','模式及初始量误用、缺少认证','流重用、可预测流、比特翻转'],['实现特点','通用，可作为MAC或一些哈希构造的基础','可实现低时延高吞吐，实际软硬件效率依算法而异'],['典型算法','DES、AES、SM4','ZUC']])+r'''
<p>CFB、OFB、CTR把分组加密输出用于生成掩码，所以两类的使用方式存在联系。并非所有流密码都由LFSR构成，也不能仅按名称断言哪一类总是更快。</p>
<h3>5. 课程作业：hello 与重复字节99</h3>
<p>按ASCII，hello的字节为68 65 6C 6C 6F（十六进制）。密钥流10011001反复出现，即每字节99。</p>
'''+table('第五讲P46：完整逐字节异或',['字符','ASCII十六进制','明文二进制','密钥流','密文二进制','密文十六进制'],[['h','68','01101000','10011001','11110001','F1'],['e','65','01100101','10011001','11111100','FC'],['l','6C','01101100','10011001','11110101','F5'],['l','6C','01101100','10011001','11110101','F5'],['o','6F','01101111','10011001','11110110','F6']])+r'''
<p>密文为F1 FC F5 F5 F6。两个l得到相同F5，字节重复关系直接保留；这个流周期仅8位。攻击者可穷举256种重复字节，再用语言或格式排除，也可以由任一已知明文字节求出99并解全部字节。99在题面公开时，当然可直接解密；分析其安全性时是假设攻击者知道“重复一个字节”的构造，但尚不知道该字节值。</p>
<h3>6. 保密之外还有篡改</h3>
<p>攻击者把密文某字节异或Δ=08，则解密明文同样异或08。ASCII的1=31，9=39，31⊕39=08；若知道金额首位位置，可把“100”首字符改为“900”，无需知道流。OFB、CTR同样存在这个代数性质。丢位、插位造成流不同步是另一类情况，不能套“一位错误只影响一位”的固定对齐结论。</p>
'''+f'<p>长周期与统计均衡{cloze("不是密码学不可预测性的充分条件")}；重复流导致{cloze("两密文异或等于两明文异或")}；加密自身{cloze("不提供完整性认证")}。</p>'+think('hello作业中，攻击者知道第一个字符h，会发生什么？','<p>从F1⊕68=99得到第一个密钥流字节。题设流为同一字节反复，所以所有后续字节都使用99，依次得到65、6C、6C、6F。如果流来自安全生成器，知道一个流字节通常不能预测剩余流；这道攻击依赖极短周期，不应概括为“任意流密码已知一字节就全破”。</p>','恢复99两分；解释重复流结构两分；指出对安全生成器不能推广一分。')
add('f01','流密码：一次一密、密钥流重用与 hello 作业','第五讲','P15–19、P33、P46','理解计算安全与信息论安全的差别，完成重复密钥流作业，并推导已知明文与定向篡改。',body,[
('hello与重复99字节流异或的密文是？',[opt('F1 FC F5 F5 F6。',True,'68、65、6C、6C、6F分别⊕99。'),opt('F1 FC F5 F6 F5。',False,'最后两个字节位置错误。'),opt('68 65 6C 6C 6F。',False,'这是ASCII明文。'),opt('99 99 99 99 99。',False,'这是密钥流。')],'按字节异或，密文不必是可打印字符。'),
('哪些说法成立？',[opt('密钥流复用时，已知一条明文可恢复另一条对应位置。',True,'由流抵消等式推出。'),opt('统计检测通过就足以当密码流。',False,'可预测序列也可能呈现不错的统计外观。'),opt('短种子确定性扩展一般只主张计算安全。',True,'独立信息熵不会凭空增长。'),opt('四个LCG输出对任意情况都唯一确定m、a、c。',False,'需要可逆性等条件，未知模数也可能多解。')],'不要把测试外观或特例恢复方法当普遍保证。'),
('知道金额首字符1的位置，想改成9，流密码攻击者怎样改密文？',[opt('把对应密文字节异或08。',True,'31⊕39=08，密文差分直接传给明文。'),opt('必须先知道种子密钥。',False,'定向异或篡改无需密钥。'),opt('把整个密文反转。',False,'不能产生所需局部字符变化。'),opt('异或99。',False,'99是hello作业流，不是1到9的差分。')],'消息认证用于阻止这种可检测性缺失。')])

start('z01')
body=r'''
<h3>1. ZUC 的三层分工</h3>
<p>ZUC以祖冲之的名字命名。本课介绍128位KEY与128位IV输入、每拍32位密钥字输出的版本。课件列出2011年LTE、2012年GM/T 0001、2016年GB/T 33133及2020年ISO/IEC 18033-4的标准化沿革。这里研究产生流的核心，下一章再研究如何构造机密性算法的IV。ZUC也可用于完整性算法，但本课不推导128-EIA3。</p>
<p>上层是LFSR（线性反馈移位寄存器），16个单元s₀到s₁₅，每个31位，负责线性状态演化。中层BR从分散位置抽位，组合成4个32位字X₀到X₃。下层F含两个32位记忆单元R₁、R₂及非线性S盒，输出W。正常输出阶段密钥字Z=W⊕X₃。</p>
'''+diagram('flowchart TD\n K["128位KEY与128位IV"] --> L["LFSR：16个31位单元"]\n L --> B["BR：抽取并拼成X0到X3"]\n B --> F["X0 X1 X2进入F：R1 R2与S盒"]\n F --> W["32位W"]\n B --> X["32位X3"]\n W --> Z["Z等于W异或X3"]\n X --> Z\n F -. "初始化时W逻辑右移1位反馈" .-> L','LFSR使用模2³¹−1，F使用32位字和模2³²加法。初始化反馈只在初始化阶段使用。')+r'''
<h3>2. 两种加法、两种位宽</h3>
<p>⊕是逐位异或，不进位。⊞是普通整数加法后mod2³²，保留32位。LFSR使用的是mod(2³¹−1)，不能截断31位后就认为完成取模。‖表示拼接，不是加法。ROL₃₁和ROL₃₂的环宽不同；W≫1在初始化中是<strong>逻辑右移</strong>，丢弃最低位、高端补0，不是循环右移。</p>
<p>对32位字a，高半a<sub>H</sub>是31到16位，低半a<sub>L</sub>是15到0位。对31位sᵢ，高16位是<strong>30到15位</strong>，低16位是15到0位，两部分共享第15位。这看似重叠，却正是算法定义；不能照搬32位的高半取法。</p>
<h3>3. 密钥装入：8+15+8=31</h3>
<div class="formula">\[s_i=\operatorname{key}_i\parallel d_i\parallel\operatorname{iv}_i=(\operatorname{key}_i\ll23)\;|\;(d_i\ll8)\;|\;\operatorname{iv}_i.\]</div>
<p>keyᵢ、ivᵢ各8位，dᵢ是固定15位公开常量。16个非零常量不同，避免全零KEY/IV直接装成全零寄存器。课件d₀=100010011010111₂=44D7，d₁=010011010111100₂=26BC，d₁₅=100011110101100₂=47AC。原例key₀=01、iv₀=02：s₀=(01&lt;&lt;23)|(44D7&lt;&lt;8)|02=00800000|0044D700|02=00C4D702。</p>
<h3>4. LFSR反馈与两种模式</h3>
<div class="formula">\[p=2^{31}-1,\quad v=\big(2^{15}s_{15}+2^{17}s_{13}+2^{21}s_{10}+2^{20}s_4+(1+2^8)s_0\big)\bmod p.\]
\[s_{16}=\begin{cases}(v+u)\bmod p,&\text{初始化，}u=W\gg1,\\v,&\text{工作模式。}\end{cases}\]</div>
<p>若结果s₁₆=0，则置s₁₆=p，用全1的31位字表示模p的零类，维持寄存器取值在1到p。然后把旧(s₁,…,s₁₅,s₁₆)作为新(s₀,…,s₁₅)。每次移位是一整个31位单元前移，不是把整个数组按1个bit左移。</p>
<h3>5. 为什么快速模运算可以“绕回来”</h3>
<p>因为2³¹≡1(modp)，乘2ᵏ可按31位循环左移实现。若把a分为将移出高端的部分q和留在低端的部分r，则2ᵏa=q·2³¹+r≡q+r；高端移出部分回到低端正好实现同余。对全1表示p的零类，循环移位仍是全1，应理解为规范表示，而非标准余数0。</p>
<div class="formula">\[2^k a\equiv\operatorname{ROL}_{31}(a,k)\pmod p.\]
\[a(2^i+2^j+2^k)\equiv\operatorname{ROL}_{31}(a,i)+\operatorname{ROL}_{31}(a,j)+\operatorname{ROL}_{31}(a,k)\pmod p.\]</div>
<p>模加时c=a+b，再c=(c&amp;7FFFFFFF)+(c≫31)，把溢出的2³¹进位作为1加回。两操作数在1…p内时结果也在1…p，是该表示体系下的模加；若另外使用标准余数约定，结果p还需解释为0。多个操作数要逐次做规范模加，不能无界累加后只折叠一次。</p>
<p>课件例7FFFFFF0+20=80000010；低31位为00000010，进位1加回得到11。因此反馈v可由4个31位循环移位、s₀本身、ROL₃₁(s₀,8)做模加组成，无须大整数除法。汉明重量是二进制1的个数；系数2ᵏ仅一个1。</p>
'''+f'<p>ZUC上层单元{cloze("31位，模2³¹−1")}，F的字运算{cloze("32位，⊞表示模2³²加法")}；初始化W≫1是{cloze("逻辑右移，丢弃最低位")}；sᵢ的高16位从{cloze("第30位到第15位")}。</p>'+think('为什么80000010折叠为11，而不是10？','<p>80000010=2³¹+10（十六进制低部分为10）。在模(2³¹−1)下2³¹相当于1，所以余类为10+1=11。只与7FFFFFFF按位与会漏掉这个进位贡献，那相当于mod2³¹，而不是mod(2³¹−1)。</p>','拆出高位2分；同余关系2分；区分两个模数1分。')
add('z01','ZUC 基础：三层结构、31位模运算与装入','第五讲','P20–25','补足ZUC符号与位宽基础，推导密钥装入、LFSR反馈以及快速模乘和模加。',body,[
('key₀=01、iv₀=02、d₀=44D7时，s₀为？',[opt('00C4D702。',True,'01左移23位，44D7左移8位，再拼入02。'),opt('0144D702。',False,'误把15位d当16位，变成8+16+8。'),opt('44D70102。',False,'拼接顺序错误。'),opt('0044D702。',False,'遗漏key0的高8位贡献。')],'按8+15+8位拼接，不能直接拼三段十六进制文本。'),
('ZUC中哪些说法成立？',[opt('LFSR的模数是2³¹−1。',True,'不是2³¹或2³²。'),opt('F中的⊞按mod2³²。',True,'保留32位字。'),opt('31位sᵢ的高16位是31到16位。',False,'最高有效位是30，应取30到15位。'),opt('W≫1把最低位循环回最高位。',False,'这里是逻辑右移。')],'必须同时看运算符、位宽和所属层。'),
('7FFFFFF0与20做ZUC式模加，结果是什么？',[opt('10。',False,'遗漏2³¹进位回加。'),opt('11。',True,'80000010低31位10加进位1。'),opt('80000010。',False,'尚未模约减。'),opt('0。',False,'该和值不属于零类。')],'mod(2³¹−1)与简单位掩码不同。')])

start('z02')
body=r'''
<h3>1. BR：从分散寄存器拼接128位</h3>
<div class="formula">\[\begin{aligned}X_0&=s_{15H}\parallel s_{14L},&X_1&=s_{11L}\parallel s_{9H},\\X_2&=s_{7L}\parallel s_{5H},&X_3&=s_{2L}\parallel s_{0H}.\end{aligned}\]</div>
<p>每次拼两个16位半字，前项放高半，后项放低半。sᵢH=(sᵢ≫15)&amp;FFFF，sᵢL=sᵢ&amp;FFFF。例如X₀=((s₁₅&amp;7FFF8000)&lt;&lt;1)|(s₁₄&amp;FFFF)，保持32位。X₀到X₂给F；X₃留作最终异或，不能把四个字全部作为F的直接输入。</p>
<h3>2. F：先读旧状态，再更新两个记忆单元</h3>
<div class="formula">\[\begin{aligned}W&=(X_0\oplus R_1)\boxplus R_2,\\W_1&=R_1\boxplus X_1,\quad W_2=R_2\oplus X_2,\\R_1^{new}&=S(L_1(W_{1L}\parallel W_{2H})),\\R_2^{new}&=S(L_2(W_{2L}\parallel W_{1H})).\end{aligned}\]</div>
<p>W、W₁、W₂全部用旧R₁、R₂计算，再一起更新。W₁、W₂是32位，所以H/L按31…16和15…0取，不同于31位s。交叉拼接时把低半放到新字高半，例如W₁L‖W₂H=(W₁&amp;FFFF)&lt;&lt;16 | (W₂≫16)。</p>
<div class="formula">\[\begin{aligned}L_1(X)&=X\oplus\operatorname{ROL}_{32}(X,2)\oplus\operatorname{ROL}_{32}(X,10)\oplus\operatorname{ROL}_{32}(X,18)\oplus\operatorname{ROL}_{32}(X,24),\\L_2(X)&=X\oplus\operatorname{ROL}_{32}(X,8)\oplus\operatorname{ROL}_{32}(X,14)\oplus\operatorname{ROL}_{32}(X,22)\oplus\operatorname{ROL}_{32}(X,30).\end{aligned}\]</div>
<p>S对高到低4字节依次用S₀、S₁、S₀、S₁；课件记S₂=S₀、S₃=S₁。这些是<strong>ZUC专用</strong>S盒，不能换成AES或SM4。课件未给两张具体256项表，本指南保留课件的完整算法表达；若进一步实现ZUC，需使用<a href="https://www.gsma.com/solutions-and-impact/technologies/security/wp-content/uploads/2019/05/eea3eia3zucv16.pdf" target="_blank" rel="noopener noreferrer">ETSI/SAGE ZUC v1.6的S₀/S₁参数表</a>，不要用自造表补齐。本指南的IV实验不调用ZUC核心，避免把演示误认成实际加密。</p>
<h3>3. 一次装入，32拍初始化，再丢弃1拍</h3>
<ol><li>装入16个sᵢ，置R₁=R₂=0。</li><li>重复32次：先BR；再F计算W并更新R；最后以u=W≫1运行LFSR初始化模式。这些W都不输出为密钥流。</li><li>工作阶段第1拍：BR；F更新但丢弃W；LFSR工作模式。只丢弃这一拍一次。</li><li>之后每拍：BR；F得到W；输出Z=W⊕X₃；LFSR工作模式。X₃来自<strong>本拍更新LFSR之前</strong>的BR结果。</li></ol>
<p>每拍顺序不能交换。若先移动LFSR再取BR，或把初始化的W发出去，就会得到不同序列；即使两端都犯同样错误也不是标准ZUC。</p>
<h3>4. 从核心到128-EEA3机密性算法</h3>
<p>输入COUNT=32位计数器、BEARER=5位承载标识、DIRECTION=1位方向、CK=128位机密性密钥、LENGTH=消息比特长度。KEY=CK，IV按16字节构造：</p>
<div class="formula">\[\begin{aligned}\operatorname{iv}_0\parallel\operatorname{iv}_1\parallel\operatorname{iv}_2\parallel\operatorname{iv}_3&=\operatorname{COUNT}\quad\text{（高字节先）},\\\operatorname{iv}_4&=\operatorname{BEARER}\parallel\operatorname{DIRECTION}\parallel00_2,\\\operatorname{iv}_5=\operatorname{iv}_6=\operatorname{iv}_7&=00_{16},\\\operatorname{iv}_{8+j}&=\operatorname{iv}_j,\quad0\le j\le7.\end{aligned}\]</div>
<p>“前8字节重复”是算法规定的IV内部结构，不等于允许同钥的不同消息复用同一个完整IV。同一CK下应由协议保证COUNT/BEARER/DIRECTION组合的使用满足唯一性；计数器回绕或状态回滚要重新管理密钥，不能凭这条构造公式就宣称永远不会重复。</p>
<p>产生L=⌈LENGTH/32⌉个32位密钥字；按每字最高位先展开，只取前LENGTH位与消息异或。若LENGTH=100，生成4字128位，取前三字96位及第4字最高4位，丢余28位。课件把初始化与流生成对应Gen是便于教学的抽象；实践中CK是事先提供的密钥，这两步主要是初始化和扩展密钥流，不能当成凭空产生新的独立秘密钥。</p>
<h3>5. 课堂题：COUNT=66035492、BEARER=01111、方向0、长度1000</h3>
'''+think('构造全部16个IV字节，并计算生成字数','<p>COUNT按高到低字节拆成66 03 54 92。第4字节=01111‖0‖00=01111000=78，其后3字节00。后8字节复制前8字节，因此IV=66 03 54 92 78 00 00 00 66 03 54 92 78 00 00 00。</p><p>L=⌈1000/32⌉=32，前31字共992位，第32字只用最高8位，丢弃24位。这里BEARER是二进制01111，不是十六进制1111；LENGTH是比特，不是字节。</p>','COUNT拆字节2分；iv4位拼接2分；复制后半2分；32字和末字8位各1分。')+r'<div id="zuc-iv-lab" class="native-lab"></div>'+f'<p>初始化重复{cloze("32次且不输出流")}，工作开始{cloze("再舍弃1拍")}；长度1000比特需{cloze("32个字，末字取最高8位")}。</p>'
add('z02','ZUC 完整运行顺序与机密性算法参数题','第五讲','P26–32','展开BR、F和初始化时序，逐位构造EEA3的IV，完成1000比特课堂例题。',body,[
('ZUC完成32次初始化后下一步是什么？',[opt('马上把第32次W当Z₁输出。',False,'初始化输出不作为流。'),opt('工作模式运行一拍并舍弃F输出，再开始输出Z。',True,'避免漏掉额外丢弃拍。'),opt('每个Z输出前都重新初始化32次。',False,'正常工作连续更新状态。'),opt('清空R₁、R₂再输出。',False,'会破坏初始化后的内部状态。')],'初始化、舍弃拍和输出拍是三个阶段。'),
('课堂题16字节IV中iv₄及iv₁₂是什么？',[opt('都为78。',True,'iv12复制iv4，01111‖0‖00=78。'),opt('分别78和00。',False,'后8字节重复前8。'),opt('都为0F。',False,'没有把BEARER移到字节高5位。'),opt('都为F0。',False,'拼接位移错误。')],'BEARER占高5位，方向占第2位，末两位为0。'),
('LENGTH=1000位时，哪些正确？',[opt('生成32个32位字。',True,'向上取整1000/32。'),opt('最后一字使用最高8位。',True,'1000−31×32=8。'),opt('共使用1000字节。',False,'题目单位为比特。'),opt('应把初始化阶段32个W直接当作消息流。',False,'初始化W不能输出。')],'先判断单位，再计算总字数和末字取位。')])

start('a01')
body=r'''
<h3>1. MAC 回答“有没有被改”，不隐藏内容</h3>
<p>消息认证码MAC，又称消息鉴别码，用共享密钥K把消息M变成标签T。发送M‖T，接收者以同一K重算标签并比较。旁观者仍可读到M，故MAC本身不提供保密。裸Hash人人都能重算，不能替代MAC抵抗恶意修改。</p>
<div class="formula">\[T=\operatorname{Co}(K,M),\quad \operatorname{Ver}(K,M,T)=\begin{cases}1,&T=\operatorname{Co}(K,M),\\0,&\text{否则}.\end{cases}\]
\[\operatorname{Ver}(K,M,\operatorname{Co}(K,M))=1.\]</div>
<p>这是课件定义2.4：K∈𝒦，M∈ℳ，标签T∈ℳ′；安全要求是攻击者即使获得许多消息及标签，也难以为未获认证的新消息生成合法标签。合法标签说明发送方持有共享K，并不能在多名持钥者之间证明具体是谁。双方都能生成标签，所以MAC不具有数字签名的公开验证与不可否认性质。</p>
'''+diagram('flowchart LR\n A["发送方：M和共享K"] --> B["Co生成标签T"]\n B --> C["发送M与T"]\n C --> D["接收方：同K重算T星"]\n D --> E["比较收到T：一致才接受"]','消息与标签均可公开，密钥不发送。验证成功必须建立在密钥安全和MAC安全的前提下。')+call('correction','课件边界 · 删除、重放与身份','MAC能发现被认证消息内部的删改和伪造，不能阻止整条消息被丢弃，也不能自动判断收到的是昨天的合法消息。防重放应把序号、会话标识等纳入认证，并维护接收状态。网络身份还依赖共享密钥怎样绑定到参与者。')+r'''
<h3>2. HMAC：内外两层带密钥的哈希</h3>
<p>HMAC的研究构造发表于1996年，RFC 2104发表于1997年。设哈希输入分组长B字节、输出长h字节（h≤B）。先将K规范化为B字节K′，然后内层处理消息，外层再封装内层摘要。ipad是B个36，opad是B个5C，值均为十六进制。</p>
<div class="formula">\[K'=\begin{cases}H(K)\parallel0^{8(B-h)},&|K|>B,\\K\parallel0^{8(B-|K|)},&|K|\le B,\end{cases}\quad |K|\text{ 以字节计}.\]
\[I=H((K'\oplus\operatorname{ipad})\parallel M),\qquad T=H((K'\oplus\operatorname{opad})\parallel I).\]</div>
<p>先比较的是密钥长度与<strong>哈希输入分组长度</strong>，不是摘要长度。HMAC-SM3中B=64字节、h=32字节，输出标签默认32字节。K′异或ipad后与M拼接，不是把M也与ipad异或；内层摘要必须以原始32字节接入外层，不能把64字符十六进制打印文本当32字节摘要。</p>
'''+diagram('flowchart TD\n K["K规范化成64字节K撇"] --> I["K撇异或ipad，与M拼接"]\n I --> H["SM3：内层32字节摘要"]\n K --> O["K撇异或opad"]\n O --> J["与内层原始摘要拼接"]\n H --> J\n J --> T["SM3：最终32字节HMAC"]','HMAC-SM3的64字节是输入分组长，32字节是输出长。两次完整哈希调用可能各含多次压缩。')+r'''
<h3>3. 课堂题：密钥规范化和两个字节异或</h3>
'''+table('第五讲P40–41：HMAC-SM3的K′',['原K长度','比较与处理','最终K′长度'],[['20字节','20≤64，K后补44个00','64字节'],['64字节','等于B，直接使用K，不补额外整块','64字节'],['100字节','100>64，先SM3(K)成32字节，再补32个00','64字节']])+r'''
<p>若K′首字节4B=01001011，ipad首字节36=00110110，异或得01111101=7D；opad首字节5C=01011100，异或得00010111=17。这里补0是HMAC密钥规范化，完全不同于SM3消息的“1、若干0、长度”填充，也不同于CBC的PKCS#7填充。</p>
<p>“内外两次哈希调用”不表示只做两次压缩函数。SM3内层前缀已有64字节，后面还要拼M并做消息填充；外层输入64+32=96字节，再由SM3正常填充、分块处理。哈希函数内部细节见<a href="#s01">第三讲SM3</a>。</p>
'''+f'<p>HMAC-SM3比较密钥长度时用{cloze("B=64字节，而非32字节摘要长")}；等长密钥{cloze("原样使用，不加额外0块")}；MAC不自动提供{cloze("机密性、防重放或不可否认性")}。</p>'+think('为什么消息与H(M)一起发送，不能替代M与HMAC(K,M)？','<p>攻击者改M为M′后能公开计算H(M′)，一起替换即可。HMAC多了秘密K，攻击者没有K就难以为M′生成合法标签。若攻击者把旧合法(M,T)原样重放，T仍正确，所以协议仍需验证被认证序号、会话与时效。</p>','裸哈希可重算2分；共享秘密的作用2分；指出重放仍需状态1分。')
add('a01','MAC 与 HMAC：来源、完整性与完整计算流程','第五讲','P34–41','区分加密、裸哈希与消息认证，掌握HMAC内外层、64字节密钥处理和全部课堂题。',body,[
('HMAC-SM3使用100字节密钥，K′怎样构造？',[opt('截取前64字节。',False,'规则要求先哈希，不是截断。'),opt('SM3(K)的32字节后补32个00。',True,'先压缩至摘要长，再补到分组长。'),opt('K后补28个00。',False,'不能因超过B就补到128。'),opt('SM3(K)后再补64个00。',False,'最终应64字节，不是96。')],'K′总长必须等于哈希输入分组长。'),
('HMAC-SM3中K′首字节4B与ipad、opad异或结果依次是？',[opt('7D、17。',True,'4B⊕36=7D，4B⊕5C=17。'),opt('17、7D。',False,'颠倒内外填充值。'),opt('81、A7。',False,'误用整数加法。'),opt('36、5C。',False,'那是ipad和opad原字节。')],'异或按比特无进位。'),
('关于MAC，哪些成立？',[opt('验证共享密钥控制下的消息完整性与来源。',True,'仍依赖安全密钥分发与正确实现。'),opt('MAC天然阻止所有重放。',False,'旧消息旧标签仍匹配。'),opt('接收方也能生成标签。',True,'共享K，所以不能用MAC区分双方的不可否认性。'),opt('标签正确就证明内容从未被读取。',False,'MAC没有保密性。')],'验证目标要和协议中的身份及新鲜性分开。')])

start('a02')
body=r'''
<h3>1. CBC-MAC：零初值，只输出最后一块</h3>
<div class="formula">\[C_0=0^n,\quad C_i=E_K(M_i\oplus C_{i-1}),\quad T=C_t.\]</div>
<p>它长得像CBC加密，但目标是认证，只发送最终标签，不发送完整加密链。基本CBC-MAC要求同一密钥认证的消息来自<strong>固定长度域</strong>，输入分块与编码也必须无歧义；填充不足一块并不自动修复跨长度伪造。IV固定为全0，不能像CBC加密那样随消息随机选IV。</p>
'''+diagram('flowchart LR\n Z["固定零初值"] --> A["异或M1"]\n A --> E1["E_K"]\n E1 --> B["异或M2"]\n B --> E2["E_K"]\n E2 --> C["继续链接至末块"]\n C --> T["只输出最后标签T"]','基本CBC-MAC适用固定长度域。变长消息应采用经设计的改进构造，例如CMAC。')+think('证明直接把基本CBC-MAC用于变长消息会被伪造',r'<p>已知一个单块消息M的标签T=E<sub>K</sub>(M)。构造两块新消息M‖(M⊕T)。第一块链值为T；第二块进入E的输入是(M⊕T)⊕T=M，末标签再次为T。攻击者没有密钥，却为新的两块消息给出正确标签。这个示例假定验证器直接接受不同整块长度的基本CBC-MAC；正规定长域或合适改进构造会排除这类跨长度攻击。</p>','已知标签关系1分；构造新消息2分；两块递推消去T两分；说明变长条件1分。')+r'''
<h3>2. KDF 的目标：提取、扩展和用途隔离</h3>
<p>课件定义2.5写作KDF:𝒦×𝒮×ℕ→𝒯，输入派生材料σ、盐salt、迭代次数r，输出密钥k。σ可以来自已有密钥、协商共享秘密或口令；salt用于区分输入实例、减少预计算复用，r是某些口令KDF的成本参数。不同KDF接口不同，不是每一种都有可省略的salt和r。HKDF还有info和输出长度L，但没有用于口令加固的r。</p>
'''+call('correction','必须修正 · KDF不能凭空创造秘密熵','P43“由低熵输入生成高熵输出”不能按字面记忆。对固定公开salt，确定性函数不会增加输入的秘密信息熵。若σ只有100万种候选，即使输出256位，攻击者仍可逐一派生这100万种候选。KDF可以整理已有熵、产生用途不同的伪随机密钥，口令KDF可提高每次猜测成本，但输出长不等于秘密更难猜。')+r'''
<p>从“随机来源”到“格式合适的密钥”是两个步骤：材料应有足够不可预测性，KDF负责转换及域隔离。把同一派生结果同时当加密钥、MAC钥可能混淆用途；应以规范定义的info或标签派生不同子密钥。</p>
<h3>3. HKDF：Extract 提取，再 Expand 扩展</h3>
<p>下面先按课件使用比特单位：底层HMAC输出l位，要求输出L位，N=⌈L/l⌉。Co的第一个参数始终是HMAC密钥。Extract用salt作HMAC密钥、σ作消息，得到prk；Expand用prk作HMAC密钥。info是公开上下文，例如协议版本、方向、用途标签。</p>
<div class="formula">\[\begin{aligned}\operatorname{prk}&=\operatorname{HMAC}(\operatorname{salt},\sigma),\\T(0)&=\varepsilon,\quad N=\lceil L/l\rceil,\\T(i)&=\operatorname{HMAC}(\operatorname{prk},T(i-1)\parallel\operatorname{info}\parallel\operatorname{I2OSP}(i,1)),\quad1\le i\le N,\\\operatorname{OKM}&=\operatorname{left}_L(T(1)\parallel\cdots\parallel T(N)).\end{aligned}\]</div>
<p>ε是空串，不是一个00字节。计数器i编码成<strong>单字节</strong>01、02、03…FF，而不是字符“1”或四字节整数。RFC 5869的L、HashLen都以<strong>字节</strong>计，最大L≤255×HashLen；课件以比特计的对应上限是L≤255l。若salt省略，RFC约定用HashLen个00。HKDF-SM3中HashLen=32字节，上限8160字节。HKDF适合已有足够熵的材料，不能直接替代专为低熵口令设计的慢速KDF。</p>
<h3>4. 两个长度例题：512位与600位</h3>
<p>HMAC-SM3的l=256位。L=512时N=2，拼T₁与T₂正好64字节。第五讲P47作业L=600时，600/256=2.34375，向上取整N=3。</p>
'''+think('写出600位派生的T(1)、T(2)、T(3)，以及截取位置',r'<div class="formula">\[\begin{aligned}T(1)&=\operatorname{HMAC}(\operatorname{prk},\operatorname{info}\parallel01_{16}),\\T(2)&=\operatorname{HMAC}(\operatorname{prk},T(1)\parallel\operatorname{info}\parallel02_{16}),\\T(3)&=\operatorname{HMAC}(\operatorname{prk},T(2)\parallel\operatorname{info}\parallel03_{16}).\end{aligned}\]</div><p>前两块共512位，还缺88位=11字节。因此输出T(1)的32字节、T(2)的32字节、T(3)最高的11字节，共75字节。不是把三块全部96字节输出，也不是取T(3)最低88位。标准HKDF要求按字节计长度，本题600恰好是8的倍数，可直接对应75字节；非整字节长度则需另明确截位协议。</p>','N=3一分；三个递推各2分；前两块512位一分；末块88位=11字节两分。')+r'''
<h3>5. 可运行参考代码：HMAC-SM3 与 HKDF-SM3</h3>
<p>下面代码使用Python标准库，需底层hashlib提供SM3；若环境缺少SM3会明确报错。它以字节处理，包含RFC长度上限检查，并独立实现内外层HMAC再与标准hmac库核对。代码用于核对推导，不要求零基础学习者先掌握编程。</p>
<pre><code class="language-python">import hashlib
import hmac

def sm3(data):
    return hashlib.new("sm3", data).digest()

def hmac_sm3(key, message):
    if len(key) &gt; 64:
        key = sm3(key)
    key = key + bytes(64 - len(key))
    inner_key = bytes(x ^ 0x36 for x in key)
    outer_key = bytes(x ^ 0x5C for x in key)
    return sm3(outer_key + sm3(inner_key + message))

def hkdf_sm3(ikm, salt, info, length):
    if not isinstance(length, int) or not 0 &lt;= length &lt;= 255 * 32:
        raise ValueError("length must be 0..8160 bytes")
    if salt is None:
        salt = bytes(32)
    prk = hmac_sm3(salt, ikm)
    previous = b""
    output = bytearray()
    count = (length + 31) // 32
    for i in range(1, count + 1):
        previous = hmac_sm3(prk, previous + info + bytes([i]))
        output.extend(previous)
    return bytes(output[:length])

key = bytes([0x4B]) * 20
message = b"hello"
assert hmac_sm3(key, message) == hmac.new(key, message, "sm3").digest()
derived = hkdf_sm3(bytes(range(32)), b"course-salt", b"encryption", 75)
assert len(derived) == 75
print(derived.hex())</code></pre>
'''+f'<p>基本CBC-MAC的消息域应{cloze("固定长度，IV固定全0")}；HKDF的Extract密钥是{cloze("salt")}，Expand密钥是{cloze("prk")}；600位SM3派生需要{cloze("3块，只取第3块前11字节")}。</p>'+call('exam','课程收束 · 你应能独立解释','为什么能解密？为什么这种模式隐藏或泄露重复？IV需要什么条件？密文被改会怎样？认证标签验证了什么？派生输出长度与秘密熵有什么区别？能把每个回答写出公式、条件和一个反例，才算把第二至第五讲连成体系。')
add('a02','CBC-MAC、KDF 与 HKDF：条件、推导和作业','第五讲','P42–44、P47–48','证明变长CBC-MAC的反例，修正熵的误解，完整推导HKDF并完成600位派生作业。',body,[
('已知单块M的CBC-MAC标签T，直接接受变长消息的基本构造会怎样？',[opt('M‖(M⊕T)也具有标签T。',True,'第二块入E的值是M⊕T⊕T=M。'),opt('把IV随机化就自动完全修复。',False,'基本CBC-MAC要求固定零IV，随意随机化可能引入其他伪造。'),opt('补到整块就一定安全。',False,'长度变化本身仍需处理。'),opt('MAC本来就是保密加密，无需分析。',False,'它提供认证，且构造条件必须满足。')],'原始CBC-MAC的固定长度条件有实质作用。'),
('HKDF-SM3派生600位，哪些正确？',[opt('N=3。',True,'向上取整600/256。'),opt('T(3)输入包含T(2)、info与单字节03。',True,'Expand递推链接前一个块。'),opt('第三块取全部32字节。',False,'只需前11字节，输出总75字节。'),opt('Extract把σ当HMAC密钥、salt当消息。',False,'顺序相反：salt是密钥，σ是消息。')],'位与字节单位、参数方向都不能省略。'),
('把100万种可能口令通过KDF输出256位密钥，哪项成立？',[opt('自动获得256位秘密熵。',False,'公开参数下仍可枚举原100万候选。'),opt('输出比输入长，所以无法穷举。',False,'攻击者枚举输入并重算。'),opt('确定性派生不能凭空增加秘密熵。',True,'成本参数可增加猜一次的代价，但不增加独立秘密信息。'),opt('公开salt使输出必然可直接反解。',False,'salt公开不等于泄露输入；其作用不同于密钥。')],'输出位数、伪随机性、输入熵与猜测成本是不同指标。')])

# Generate the repeated-block illustration as a real, self-contained vector.
cells=['001100','011110','110011','111111','110011','110011']
left=[];right=[]
for r,row in enumerate(cells):
    for c,value in enumerate(row):
        left.append(f'<rect x="{22+c*28}" y="{46+r*28}" width="26" height="26" fill="'+('#087e8b' if value=='1' else '#d8e9ed')+'"/>')
        right.append(f'<rect x="{342+c*28}" y="{46+r*28}" width="26" height="26" fill="url(#block-'+('b' if value=='1' else 'a')+')"/>')
raw='\n'.join(bodies).replace('<g id="pattern-left"></g>','<g>'+''.join(left)+'</g>').replace('<g id="pattern-right"></g>','<g>'+''.join(right)+'</g>')
(B/'additions.raw.html').write_text(raw,encoding='utf-8')
(B/'new-data.json').write_text(json.dumps(dict(modules=modules,quizzes=quizzes),ensure_ascii=False,indent=2),encoding='utf-8')

def coverage_rows(rows):
    return ''.join(f'<div class="coverage-row"><span class="map-label">{pages}</span><span><a href="#{id}">{title}</a>：{desc}</span></div>' for pages,id,title,desc in rows)
cover='<h3>第四讲 · 57页</h3>'+coverage_rows([
('P1–10','t01','对称加密接口与分组原理','封面、目录、分节并入路线；定义2.3、分类、算法与模式、双射'),
('P11–14','t02','混淆、扩散与可逆结构','轮密钥加、密钥编排、Feistel与SPN'),
('P15–16','t02','完整SPN原例','S/P表、四轮加密、逆推与实验；P17分节并入路线'),
('P18–27','t03','DES算法','奇校验、IP与逆表、E/S/P及全部查表课堂题、解密'),
('P28–30','t04','DES密钥与安全限制','PC-1/PC-2、移位表、首轮子密钥及穷举'),
('P31–32','t04','多重DES','XEX3、双重DES、中间相遇、四种三重DES形式'),
('P33–39','t05','AES轮变换','参数、状态布局、SubBytes、ShiftRows、GF(2⁸)、MixColumns和解密'),
('P40–41','t06','AES密钥扩展','完整递推、Rcon与首轮原例'),
('P42–52','t07','SM4结构与密钥编排','沿革、四字窗口、完整S盒、T与T′、FK/CK和解密'),
('P53–54','t08','SM4国标向量与算法对照','首轮细算、完整32轮与DES/AES/SM4比较'),
('P55','t06','作业1：AES轮密钥','原矩阵保留、按列读取、完整首轮展开'),
('P56–57','t08','作业2：对称轮密钥变体','加密等于解密证明、一次请求攻击；结束页并入收束')])+ '<h3>第五讲 · 48页</h3>'+coverage_rows([
('P1–8','m01','ECB与CBC','封面目录分节并入路线，概述、分类、重复结构与CBC链'),
('P9–12','m02','CFB、OFB与CTR','完整公式、初始量约束、并行性、末段与五模式对照'),
('P13–14','m01','课堂题：ECB/CBC','原玩具算法、全部中间计算与反验；实验见m03'),
('P15–19','f01','流密码基础','分节导读、OTP、二元加法、流安全条件及算法比较'),
('P20–25','z01','ZUC基础','沿革、符号、三层结构、装入、反馈和快速模运算'),
('P26–28','z02','ZUC完整运行流程','BR、F、32拍初始化与舍弃拍、输出顺序'),
('P29–32','z02','EEA3与课堂题','全部参数、IV构造、100/1000比特取位及实验'),
('P33','f01','定向篡改','100改900的异或推导；模式错误传播另见m03'),
('P34–37','a01','MAC定义与边界','分节导读、工作原理、定义2.4、完整性与来源'),
('P38–41','a01','HMAC完整流程及小测','内外层、密钥规范化、全部长度题和异或题'),
('P42–44','a02','CBC-MAC与KDF/HKDF','固定长度边界、定义2.5、熵勘误与两阶段推导'),
('P45','m03','作业1：CBC/CTR','完整逐块计算、玩具与实际密码错误传播比较'),
('P46','f01','作业2：hello','全部ASCII、二进制、十六进制密文及攻击分析'),
('P47–48','a02','作业3：HKDF','600位三块递推与截取；结束页并入课程收束')])
corr=coverage_rows([
('第四讲P4–5','t01','正确性与保密性','可正确解密不代表必然安全；“结构简单”不是对称算法的定义'),
('第四讲P11、14','t02','扩散与结构比较','雪崩是统计趋势；SPN与Feistel速度须比较具体部件和轮数'),
('第四讲P31–32','t04','组合数不等于强度','中间相遇需时间与存储，且匹配后要排除假候选'),
('第四讲P42','t08','SM3与完整性','裸摘要不能独立认证来源；应使用MAC或签名等机制'),
('第五讲P4、7、14','m01','模式条件','ECB单块也不自动保证一般方案安全；CBC不保证所有密文块互异；填充须无歧义'),
('第五讲P16、18','f01','OTP与LCG','OTP不是唯一信息论安全构造；4个LCG输出不保证普遍唯一恢复未知参数'),
('第五讲P21、24','z01','移位符号与表示','W≫1是逻辑右移；模2³¹−1乘法按同余及非零表示理解'),
('第五讲P29','z02','初始化不创造独立新密钥','CK已有，初始化和流生成是确定性展开；IV唯一性依赖协议状态管理'),
('第五讲P35–36','a01','MAC的身份与删除边界','认证持钥者而非无条件身份绑定；不能发现整条丢弃，也不自动防重放'),
('第五讲P43–44','a02','KDF熵与HKDF单位','确定性派生不增加秘密熵；salt/r并非所有KDF都有，HKDF计数器单字节，RFC长度单位为字节')])
newsrc='''<h3>第四、第五讲原课件</h3><p><a href="./第四讲——对称加密.pptx">第四讲——对称加密.pptx</a>（57页）<br><a href="./第五讲——工作模式、流密码与消息认证码.pptx">第五讲——工作模式、流密码与消息认证码.pptx</a>（48页）</p>
<h3>新增章节核验来源</h3><p><a href="https://csrc.nist.gov/pubs/fips/197/final" target="_blank" rel="noopener noreferrer">NIST FIPS 197：AES</a>：核对状态布局、轮变换与密钥编排。<br><a href="https://csrc.nist.gov/pubs/sp/800/38/a/final" target="_blank" rel="noopener noreferrer">NIST SP 800-38A：五种保密模式</a>：核对IV约束、并行性、错误传播。<br><a href="https://csrc.nist.gov/pubs/sp/800/38/b/upd1/final" target="_blank" rel="noopener noreferrer">NIST SP 800-38B：CMAC</a>：核对CBC-MAC改进的认证用途。<br><a href="https://www.rfc-editor.org/rfc/rfc2104.html" target="_blank" rel="noopener noreferrer">RFC 2104：HMAC</a>：核对内外层与密钥长度处理。<br><a href="https://www.rfc-editor.org/rfc/rfc5869.html" target="_blank" rel="noopener noreferrer">RFC 5869：HKDF</a>：核对Extract/Expand、单字节计数器、长度单位与输入熵条件。<br><a href="https://www.etsi.org/deliver/etsi_ts/135200_135299/135222/16.00.00_60/ts_135222v160000p.pdf" target="_blank" rel="noopener noreferrer">3GPP/ETSI TS 35.222：ZUC规范入口</a>，及<a href="https://www.gsma.com/solutions-and-impact/technologies/security/wp-content/uploads/2019/05/eea3eia3zucv16.pdf" target="_blank" rel="noopener noreferrer">ETSI/SAGE ZUC v1.6算法规范</a>。新增数值以原课件为依据独立复算，拓展反例与评分点为教学补充。</p>'''

def assemble():
    rendered=B/'additions.rendered.html'
    additions=rendered.read_text(encoding='utf-8') if rendered.exists() else raw
    final=original
    final=final.replace('<section class="appendix" id="coverage">',additions+'\n<section class="appendix" id="coverage">',1)
    nav=[]
    for idx,m in enumerate(modules,19):
        if idx in [19,27,31,33]:nav.append('<p class="nav-group">'+{19:'对称加密与分组算法 · 第四讲',27:'工作模式与流密码 · 第五讲',31:'ZUC深入演算 · 第五讲',33:'认证与密钥派生 · 第五讲'}[idx]+'</p>')
        nav.append(f'<a class="nav-link" data-nav="{m["id"]}" href="#{m["id"]}"><span class="nav-number">{idx:02d}</span><span class="nav-label">{m["title"]}</span><span class="nav-dot" aria-hidden="true"></span></a>')
    final=final.replace('<p class="no-results"', ''.join(nav)+'<p class="no-results"',1)
    final=final.replace('</section>\n<section class="appendix" id="corrections">',cover+'</section>\n<section class="appendix" id="corrections">',1) if False else final
    # Insert appendices by parsing only those small original sections; preserve the existing chapters byte for byte.
    for id,extra in [('coverage',cover),('corrections',corr),('sources',newsrc)]:
        pat=rf'(<section class="appendix" id="{id}">[\s\S]*?)(</section>)'
        final,n=re.subn(pat,lambda m:m[1]+extra+m[2],final,count=1)
        assert n==1,id
    DATA['modules'].extend(modules);DATA['quizzes'].extend(quizzes)
    final=re.sub(r'const DATA=\{[\s\S]*?\};\s*</script>',lambda m:'const DATA='+json.dumps(DATA,ensure_ascii=False,separators=(',',':'))+';</script>',final,count=1)
    final=final.replace('第二讲·第三讲深度学习指南','第二讲至第五讲深度学习指南').replace('第二讲与第三讲','第二讲至第五讲')
    final=final.replace('两讲共 108 页，重组为 18 个知识模块。','四讲共 213 页，重组为 34 个知识模块。')
    final=final.replace('18章 · 54题','34章 · 102题').replace('108页覆盖对照','213页覆盖对照')
    final=final.replace('18章、54题','34章、102题').replace('主要依据你提供的两份课件','主要依据你提供的四份课件')
    final=final.replace('完整古典原例与SM3推导，单文件可离线学习。','古典密码、SM3、DES/AES/SM4、工作模式、ZUC、HMAC/HKDF完整课件讲解与作业，单文件可离线学习。')
    final=final.replace('乘2ᵏ可按31位循环左移实现。','乘2ᵏ可按31位循环左移实现；基本移位量取0≤k<31，更大的k先对31取模。')
    final=re.sub(r'(学习模块\s*\d+\s*/\s*)18',r'\g<1>34',final)
    hero='''<div class="callout intuition"><strong>已续入第四讲与第五讲 · 新增16章、48道选择题</strong><p>从共享密钥出发，走完DES、AES、SM4的轮变换，再学习五种工作模式、ZUC、HMAC与HKDF。保留全部课堂与作业原例，补全中间计算，并在对应章节标明条件与勘误。</p><div class="button-row"><a class="btn btn-primary" href="#t01">开始第四讲：对称加密</a><a class="btn" href="#m01">开始第五讲：模式与认证</a><a class="btn" href="#t06">AES原作业</a><a class="btn" href="#m03">模式演算实验</a></div></div>'''
    final=final.replace('<div class="intro-grid">',hero+'<div class="intro-grid">',1)
    final=final.replace('最后，亲手走完 SM3 从三个字母到 256 位摘要的每一步。','再亲手走完SM3、AES、SM4的计算，理解工作模式、流密码、认证与密钥派生的使用条件。')
    final=final.replace('最后贯通完整算法','贯通完整算法与应用')
    final=final.replace('第 15–18 章 · 从字节算到可核验的摘要','第15–18章：SM3；第19–34章：对称加密、流密码与认证')
    root=LH.fromstring(final);ct=len(root.cssselect('.cloze')) if False else len(root.xpath('//*[contains(concat(" ",normalize-space(@class)," ")," cloze ")]'))
    thinks=len(root.xpath('//details[contains(concat(" ",normalize-space(@class)," ")," think ")]'))
    final=final.replace('共 54 道针对性选择题、74 处主动回忆、35 道折叠思考题，以及 3 个演算实验。',f'共102道针对性选择题、{ct}处主动回忆、{thinks}道折叠思考题，以及6个交互演算实验。')
    final=final.replace('三个实验台的输入不计入学习进度。','六个实验台的输入不计入学习进度。')
    final=final.replace('0 / 54 题已答','0 / 102 题已答').replace('>/ 74<',f'>/ {ct}<').replace('aria-valuemax="74"',f'aria-valuemax="{ct}"')
    final=final.replace('</head>','<style>'+ (B/'new-style.css').read_text(encoding='utf-8')+'</style></head>')
    final=final.replace('</body>','<script>'+ (B/'new-labs.js').read_text(encoding='utf-8')+'</script></body>')
    from mark_homework import annotate
    final=annotate(final)
    TARGET.write_text(final,encoding='utf-8')
    print(json.dumps(dict(chapters=len(DATA['modules']),questions=len(DATA['quizzes']),concepts=ct,thinking=thinks,bytes=TARGET.stat().st_size,prerendered=rendered.exists()),ensure_ascii=False))

if __name__=='__main__' and (B/'new-labs.js').exists():assemble()

