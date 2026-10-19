from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r'C:\University Courses\计算机网络lab')
WORK = ROOT / '_lab2_工作'
WORK.mkdir(exist_ok=True)
REPORT = ROOT / '10244602411_刘小凡_Lab2_Ethernet_待补实测版.docx'
GUIDE = ROOT / 'Lab2_刘小凡_抓包与截图操作指南.docx'

def font(run, size=11.5, bold=False, color='000000', face='宋体'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), face)

def new_doc(title, subtitle):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.9)
    s.left_margin, s.right_margin = Cm(2.3), Cm(2.3)
    s.header_distance, s.footer_distance = Cm(.8), Cm(.85)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st = d.styles[name]
        st.font.name = 'Times New Roman'
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        st.font.color.rgb = RGBColor(0,0,0)
        st.paragraph_format.widow_control = True
    normal = d.styles['Normal']
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(.8)
    for nm,sz in [('Heading 1',14),('Heading 2',12),('Heading 3',11.5)]:
        st = d.styles[nm]
        st.font.size = Pt(sz)
        st.font.bold = True
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    d.styles['Caption'].font.size = Pt(10)
    d.styles['Caption'].paragraph_format.first_line_indent = Cm(0)
    d.styles['Caption'].paragraph_format.space_after = Pt(8)
    for el in d.styles._element.xpath('.//w:pBdr'):
        el.getparent().remove(el)
    p = d.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    font(p.add_run(title),18,True)
    p = d.add_paragraph(subtitle,style='Subtitle')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs: font(r,12)
    p = s.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('华东师范大学软件学院  计算机网络实验'),9)
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t in ['第 ', ' 页']:
        if t == ' 页':
            x = OxmlElement('w:fldSimple'); x.set(qn('w:instr'),'PAGE')
            r = OxmlElement('w:r'); tx = OxmlElement('w:t'); tx.text='1'; r.append(tx); x.append(r); p._p.append(x)
        font(p.add_run(t),9)
    d.core_properties.author='刘小凡'
    d.core_properties.title=title+' '+subtitle
    return d

def p(d,text,bold=False,indent=True):
    x=d.add_paragraph()
    if not indent: x.paragraph_format.first_line_indent=Cm(0)
    font(x.add_run(text),bold=bold)
    return x

def h(d,text,level=1): return d.add_paragraph(text,style=f'Heading {level}')

def page(d): d.add_page_break()

def table(d,heads,rows,widths,header=True):
    t=d.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for i,txt in enumerate(heads): t.rows[0].cells[i].text=txt
    if not header: t._tbl.remove(t.rows[0]._tr)
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row): cells[i].text=str(txt)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0 and header: trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
            tcpr.append(margins)
            if ri==0 and header:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF4'); tcpr.append(sh)
            for pp in cell.paragraphs:
                pp.paragraph_format.first_line_indent=Cm(0); pp.paragraph_format.space_after=Pt(1); pp.paragraph_format.line_spacing=1.15
                if ci==0 and len(cell.text)<24: pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs: font(rr,10.5,ri==0 and header)
    gap=d.add_paragraph()
    gap.paragraph_format.space_after=Pt(0)
    gap.paragraph_format.line_spacing=Pt(3)
    font(gap.add_run(''),1)
    return t

def shot(d,n,title,evidence):
    x=p(d,f'图 {n}  {title}  ［待插入本人截图］',True,False)
    x.paragraph_format.keep_with_next=True
    for r in x.runs: font(r,10.5,True,'666666')
    x=p(d,evidence,indent=False)
    for r in x.runs: font(r,10.5,color='666666')

def image(d,path,caption):
    x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    x.paragraph_format.first_line_indent=Cm(0); x.paragraph_format.keep_with_next=True
    r=x.add_run(); shape=r.add_picture(str(path),width=Cm(16.0)); shape._inline.docPr.set('descr',caption)
    x=d.add_paragraph(caption,style='Caption'); x.alignment=WD_ALIGN_PARAGRAPH.CENTER

def cmd(d,text):
    x=p(d,text,indent=False)
    x.paragraph_format.left_indent=Cm(.3)
    x.paragraph_format.space_after=Pt(6)
    for r in x.runs:
        font(r,10.5,face='等线'); r.font.name='Consolas'

FREG = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc',28)
FBIG = ImageFont.truetype(r'C:\Windows\Fonts\msyhbd.ttc',32)
FSM = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc',25)

def centered(draw,box,text,ft=FREG):
    x0,y0,x1,y1=box
    lines=text.split('\n'); heights=[]
    for line in lines:
        bb=draw.textbbox((0,0),line,font=ft); heights.append(bb[3]-bb[1])
    total=sum(heights)+12*(len(lines)-1); y=(y0+y1-total)/2
    for line,height in zip(lines,heights):
        bb=draw.textbbox((0,0),line,font=ft)
        draw.text(((x0+x1-(bb[2]-bb[0]))/2,y-bb[1]),line,font=ft,fill='#142233'); y+=height+12

def diagrams():
    im=Image.new('RGB',(1700,625),'white'); dr=ImageDraw.Draw(im)
    dr.text((36,20),'Ethernet II 帧的字节位置',font=FBIG,fill='black')
    dr.text((36,72),'示例条件：IPv4 无选项，ICMP Echo 数据 32 B，未带 VLAN 标签',font=FSM,fill='#444444')
    fields=[(35,300,'目的 MAC\n6 B\n偏移 0–5','#E8EEF4'),(300,565,'源 MAC\n6 B\n偏移 6–11','#E8EEF4'),(565,795,'Type\n2 B\n偏移 12–13','#E8EEF4'),(795,1035,'IPv4 首部\n20 B\n偏移 14–33','#EDF3E8'),(1035,1280,'ICMP 首部\n8 B\n偏移 34–41','#F7F0E1'),(1280,1510,'Echo 数据\n32 B\n偏移 42–73','#F7F0E1')]
    for x0,x1,tx,col in fields:
        dr.rectangle((x0,135,x1,335),fill=col,outline='#666666',width=2); centered(dr,(x0,135,x1,335),tx,FSM)
    for ya,yb in [(135,335)]:
        for x in range(1520,1661,15):
            dr.line((x,ya,min(x+8,1660),ya),fill='#666666',width=2); dr.line((x,yb,min(x+8,1660),yb),fill='#666666',width=2)
        for y in range(135,336,15):
            dr.line((1520,y,1520,min(y+8,335)),fill='#666666',width=2); dr.line((1660,y,1660,min(y+8,335)),fill='#666666',width=2)
    centered(dr,(1520,135,1660,335),'FCS\n4 B',FSM)
    for x0,x1,y,tx in [(35,795,365,'以太网首部 14 B'),(795,1510,365,'以太网有效载荷 60 B'),(35,1510,440,'典型抓包记录 74 B  不含 FCS')]:
        dr.line((x0,y,x1,y),fill='#333333',width=2); dr.line((x0,y-10,x0,y+10),fill='#333333',width=2); dr.line((x1,y-10,x1,y+10),fill='#333333',width=2)
        centered(dr,(x0,y+12,x1,y+58),tx,FSM)
    dr.text((36,535),'FCS 用虚线表示；前导码和 SFD 不在本图内。图中各框宽度不按字节数缩放。',font=FSM,fill='#444444')
    im.save(WORK/'frame_structure.png')
    im=Image.new('RGB',(1700,560),'white'); dr=ImageDraw.Draw(im)
    dr.text((35,20),'本地链路与跨网通信的地址范围',font=FBIG,fill='black')
    boxes=[(35,160,470,385,'本机\nIPv4  I_C\nMAC  M_C'),(630,160,1080,385,'默认网关\n本地接口 IPv4  I_G\n本地接口 MAC  M_G'),(1250,160,1665,385,'远程主机\nIPv4  I_S\n远端 MAC 无法从此抓包获知')]
    for b in boxes:
        dr.rectangle(b[:4],fill='#F4F6F8',outline='#657383',width=2); centered(dr,b[:4],b[4],FSM)
    for x0,x1,tx in [(480,620,'本地链路'),(1090,1240,'Internet')]:
        dr.line((x0,270,x1,270),fill='#304C70',width=4); dr.polygon([(x1,270),(x1-14,262),(x1-14,278)],fill='#304C70')
        centered(dr,(x0-25,90,x1+25,140),tx,FSM)
    dr.text((35,425),'本机抓到的请求：Ethernet  M_C → M_G     IPv4  I_C → I_S',font=FREG,fill='#142233')
    dr.text((35,480),'MAC 标识当前链路的收发接口；IPv4 标识该抓包位置看到的源与目的主机。',font=FSM,fill='#444444')
    im.save(WORK/'address_scope.png')

def report():
    d=new_doc('华东师范大学软件学院实验报告','Lab2 Ethernet 以太网帧分析')
    table(d,['实验信息','个人信息','记录信息'],[
      ['课程  计算机网络','姓名  刘小凡','实验日期  ［待填］'],
      ['实验编号  Lab2','学号  10244602411','实验时间  2 学时'],
      ['实验名称  Ethernet','年级  ［待填］','实验成绩  __________']],[5.5,5.4,5.5],header=False)
    h(d,'一 实验目的')
    p(d,'本实验通过 ICMP Echo 流量分析 Ethernet II 的字节结构与地址范围，并用广播和组播验证 I/G 判断位，区分抓包显示与实际链路字段。')
    for text in ['1. 正确选择网络接口，设置捕获过滤器和 MAC 地址解析，保存可复查的抓包文件。','2. 通过协议详情与原始字节确认目的地址、源地址和 EtherType 的位置及长度。','3. 对照本机 IP、默认网关及 ARP 缓存，解释远程 ping 的目的 MAC 为什么指向下一跳。','4. 识别广播地址与普通组播地址，并回答 IEEE 802.3 与 Ethernet II 的格式区别。']: p(d,text,indent=False)
    h(d,'二 实验内容与实验步骤')
    p(d,'先在实际联网接口上以 icmp 为捕获过滤器，运行四次 IPv4 ping 并保存单播迹线；再用 ether multicast 捕获约 30 秒背景流量，保存广播与组播迹线。分别检查 Ethernet、IPv4 和 ICMP 字段，绘制帧结构图与地址范围图，最后完成实验手册的 Turn-in 与拓展问题。')
    p(d,'捕获过滤器在开始抓包前决定记录哪些帧；显示过滤器在抓包后筛选已经记录的帧。两者使用不同语法。例如 ether multicast 用于捕获，eth.dst.ig == 1 用于显示。')
    h(d,'三 实验环境')
    table(d,['项目','配置或实测记录'],[['操作系统','Windows  ［版本待填］'],['抓包工具','Wireshark 4.6.8  本机安装版本'],['命令行工具','Windows CMD 的 ping、ipconfig、arp'],['捕获接口和链路类型','［待填接口］  Ethernet'],['抓包文件','Lab2_刘小凡_单播.pcapng；Lab2_刘小凡_广播组播.pcapng'],['网络状态','［待记录是否启用 VPN 及所在网络］']],[4.0,12.4])
    page(d)
    h(d,'四 实验过程与分析')
    h(d,'4 1 单播迹线的获取与定位',2)
    p(d,'使用 ping -4 -n 4 www.baidu.com 产生四个 ICMP Echo Request。参数 -4 固定为 IPv4，避免域名优先解析为 IPv6 时与 icmp 捕获过滤器不匹配。应先验证目标可以回复，再在开始捕获后重新执行命令。')
    shot(d,1,'单播捕获接口与过滤器','证据应包含选中接口、Ethernet 链路类型、icmp 捕获过滤器及关闭混杂模式的设置。')
    shot(d,2,'启用 MAC 地址解析','证据应显示 Resolve MAC addresses 已勾选。厂家名称仅辅助识别，地址记录使用完整的六字节 MAC。')
    shot(d,3,'IPv4 ping 结果','记录：目标 IPv4 ［待填］；发送 ［待填］；接收 ［待填］；丢失 ［待填］；平均往返时间 ［待填］ ms。')
    p(d,'抓包完成后选择 Info 为 Echo (ping) request 的帧。利用 ICMP 的 Identifier 和 Sequence Number 对应请求与回复，确认流量来自本次操作。若全部成功且没有其他 ICMP 流量，通常可见四个请求和四个回复；实际数量以保存的迹线为准。')
    shot(d,4,'请求帧的协议层与原始字节','选中请求帧并展开 Frame、Ethernet II、IPv4 和 ICMP；点击 Ethernet II 整行，使底部高亮其 14 字节首部。')
    table(d,['观测项目','实测值'],[['选取的请求帧编号','［待填］'],['对应回复帧编号','［待填］'],['请求与回复数量','请求 ［待填］  回复 ［待填］'],['Frame Length 与 Captured Length','［待填］ B / ［待填］ B'],['EtherType 与 IPv4 Protocol','［待填］ / ［待填］'],['ICMP Type 与 Code','请求 ［待填］ / ［待填］；回复 ［待填］ / ［待填］']],[7.0,9.4])
    p(d,'Ethernet II 的 Type=0x0800 表示其载荷为 IPv4；IPv4 的 Protocol=1 表示载荷为 ICMP。ICMP 是网络层控制协议，不能因它位于 IPv4 首部后就将其当作 TCP 或 UDP 传输层协议。Echo Request 的 Type=8、Echo Reply 的 Type=0，两者的 Code 通常为 0。')
    page(d)
    h(d,'4 2 Ethernet II 的字段与长度',2)
    p(d,'以下字节偏移从抓包记录首字节开始，首字节偏移为 0。对于未带 VLAN 标签的 Ethernet II 帧，前 14 字节依次为目的 MAC、源 MAC 和 Type，随后紧接上层负载。IPv4 和 ICMP 的起始位置还需根据实测首部长度确认。')
    table(d,['字段','偏移与长度','含义或验证方法'],[
      ['目的 MAC','0–5  共 6 B','请求通常指向本地下一跳的 MAC'],
      ['源 MAC','6–11  共 6 B','请求应对应选中捕获接口的 MAC'],
      ['EtherType','12–13  共 2 B','IPv4 为 08 00，即 0x0800'],
      ['IPv4 首部','从 14 开始  通常 20 B','以 Header Length 的实测值为准'],
      ['ICMP Echo 首部','IPv4 首部之后  8 B','Type、Code、Checksum、Identifier、Sequence'],
      ['Echo 数据','ICMP 首部之后','Windows 默认 ping 通常为 32 B'],
      ['FCS','标准帧末尾  4 B','多数主机抓包不显示，不能从缺失推断未发送']],[3.6,4.5,8.3])
    image(d,WORK/'frame_structure.png','图 5  默认 32 字节 Echo 数据下的 Ethernet II 字节结构示例')
    p(d,'在图 5 的条件下，以太网有效载荷为 20 + 8 + 32 = 60 B，典型抓包记录为 14 + 60 = 74 B。标准以太网帧加上 4 B FCS 后为 78 B，不包含前导码和帧起始定界符 SFD。这里的 74 B 是条件示例，实测值需由图 4 核对。')
    p(d,'此示例的负载 60 B 已超过未标记以太网的最小负载 46 B，因此不需要填充。短 ARP 帧的抓包长度可能只有 42 B，原因可能是抓取位置或驱动没有呈现填充与 FCS；标准以太网最小帧长 64 B 的规定仍然适用。')
    p(d,'若从 WLAN 抓包却显示 Ethernet II，通常是系统提供了拟以太网表示。该表示可以用于检查 MAC 地址和协议封装，但不能直接把上述标准以太网帧长当成无线链路实际发送的 802.11 帧长。VLAN 标签和 FCS 是否可见也取决于网卡、驱动与抓取位置。')
    page(d)
    h(d,'4 3 MAC 地址的作用范围',2)
    p(d,'使用 ipconfig /all 确认实际捕获接口的 IPv4、物理地址和默认网关，再用 arp -a 查询同一接口的网关 IPv4 对应 MAC。将这些信息与请求帧的 Ethernet 和 IPv4 字段逐项对照，避免把虚拟网卡或其他接口的地址混入分析。')
    shot(d,6,'捕获接口的 IP 配置','显示实际捕获接口的物理地址、IPv4、子网掩码和默认网关。')
    shot(d,7,'默认网关的 ARP 映射','显示对应接口下默认网关 IPv4 及其完整 MAC。若网关不在缓存中，先 ping 网关再查询。')
    table(d,['对象','IPv4 地址','本地可观测 MAC'],[['本机捕获接口 I_C / M_C','［待填］','［待填］'],['默认网关本地接口 I_G / M_G','［待填］','［待填］'],['远程主机 I_S','［待填］','该本地抓包无法确定']],[5.5,4.5,6.4])
    image(d,WORK/'address_scope.png','图 8  本机经默认网关访问远程主机的地址范围')
    p(d,'在普通的本地联网、未受 VPN 改写路由的场景中，访问异地目标时，请求帧的源 MAC 为本机接口 M_C，目的 MAC 为下一跳网关 M_G；IPv4 源和目的地址分别为 I_C 与 I_S。回复在本地链路上通常以网关为源 MAC、本机为目的 MAC。下一跳应由路由选择及 ARP 证据确认。')
    p(d,'每经过一条新的链路，路由器通常重新封装二层帧，因此 MAC 地址的作用范围是当前链路。IP 支持跨网络寻址，但 TTL、校验和以及可能发生的 NAT 转换都会使报文部分字段改变。远程主机可能有其网卡 MAC，只是本机的这份迹线不能给出该地址。')
    page(d)
    h(d,'4 4 广播与组播迹线',2)
    p(d,'第二次捕获使用 ether multicast 过滤器，在相同物理联网接口上记录约 30 秒背景流量。该过滤器依据目的 MAC 的组地址位筛选，包含广播和普通组播，不包含目的 MAC 为单播的普通帧。流量种类与数量随网络环境变化，需按实际捕获结果统计。')
    shot(d,9,'广播组播捕获设置','显示选中接口、ether multicast 捕获过滤器及关闭混杂模式。')
    shot(d,10,'广播目的地址与 I G 位','显示过滤器 eth.dst == ff:ff:ff:ff:ff:ff；展开 Ethernet II 的 Destination，显示 IG bit 为 Group。若为 ARP，还应展开其操作码和目标 IP。')
    shot(d,11,'普通组播目的地址与 I G 位','显示过滤器 eth.dst.ig == 1 && eth.dst != ff:ff:ff:ff:ff:ff；展开 Destination 与 IG bit，并记录上层协议。没有匹配结果时如实记录。')
    table(d,['记录项','实测结果'],[['抓包起止时间或持续时间','［待填］'],['迹线总帧数','［待填］'],['广播帧数','［待填］'],['普通组播帧数','［待填］'],['选取的广播帧编号和上层协议','［待填］'],['选取的普通组播帧编号和目的 MAC','［待填；未捕获则记未观察到］']],[8.1,8.3])
    p(d,'以太网广播地址为 ff:ff:ff:ff:ff:ff，48 位全部为 1。判断单播或组播的 I/G 位是目的 MAC 第一字节的最低有效位，即 bit 0：0 表示 Individual，1 表示 Group。广播也是组地址的一个特殊取值；组地址不一定是广播。［2］')
    p(d,'例如 ff 的最低位为 1；普通组播地址常以 01 或 33 开头，它们的最低位也为 1。以太网在每个地址字节内先发送最低位，所以手册中的“线上先发送的第一位”与这里的最低有效位一致。旁边的 U/L 位即 bit 1 判断全局或本地管理，与 I/G 的收件对象判断不同。')
    p(d,'若样本为 ARP 请求，应结合 Opcode、Sender IP 和 Target IP 解释作用。一般地址解析请求询问某个 IP 对应的 MAC；只有满足相应特征时，才能称为 Gratuitous ARP 或冲突检测流量。不能因“目的 MAC 是广播”就把所有 ARP 帧都写成免费 ARP。')
    p(d,'观测结论：［待依据图 10 与图 11 的协议、地址和统计数量填写，说明是否观察到广播和普通组播。］',indent=False)
    page(d)
    h(d,'4 5 抓包现象与问题分析',2)
    p(d,'抓不到 ping 时，首先区分目标不回复与选错接口：命令行成功而迹线为空，应核对实际路由、捕获接口和过滤器；域名使用 IPv6 时，应使用 -4 重新生成 IPv4 流量。VPN 会改变路由，物理 WLAN 上可能只看到隧道外层流量，虚拟接口中的二层信息也不能直接代表本地物理网关。')
    p(d,'未观察到广播或普通组播，可能与网络中设备较少、无线网络抑制广播、抓包接口选择或观测时长有关。可先延长自然背景观测并记录时长；如需主动生成流量，应另行保存迹线，标注生成方式，不能把主动流量计入原先 30 秒的背景统计。')
    p(d,'本次实际遇到的问题与处理：［待记录现象、排查步骤、处理结果；若无异常则据实填写。］',indent=False)
    h(d,'五 实验结果总结')
    h(d,'5 1 Turn in 题回答',2)
    p(d,'帧结构图见图 5。目的 MAC、源 MAC 与 Type 的长度分别为 6 B、6 B 和 2 B，以太网首部共 14 B；IP 数据报是它的有效载荷。FCS 在标准以太网帧末尾占 4 B，本图以虚线表示，具体抓包是否呈现需单独核对。')
    p(d,'地址范围图见图 8。本机和网关位于本地链路，远程主机经 Internet 到达。请求的以太网目的地址通常属于本地下一跳，IPv4 目的地址属于远程主机；图中 I_C、M_C、I_G、M_G、I_S 对应值见地址表。')
    p(d,'广播地址是 ff:ff:ff:ff:ff:ff。单播与组播判断位为目的 MAC 第一字节的最低有效位 I/G：0 为单播，1 为组播或广播；第一字节的最高位不是该判断位。')
    page(d)
    h(d,'5 2 IEEE 802 3 拓展问题',2)
    p(d,'1. 首部长度。Ethernet II 与 IEEE 802.3 的 MAC 首部都为 14 B，区别在末尾两字节被解释为 Type 或 Length。IEEE 802.3 后通常还接 LLC：当 Control 为单字节 UI 格式时，DSAP、SSAP、Control 共 3 B，合计 17 B；若再接 5 B SNAP，则合计 22 B。其他 LLC 控制格式可能用两字节 Control，因此需逐帧测量。Padding 和 FCS 在尾部，不应计入前部首部长度。［3］')
    p(d,'2. 帧类型判别。查看源 MAC 后的 16 位 Length/Type 字段：值不大于 1500 时按长度解释，值不小于 1536（0x0600）时按 EtherType 解释。1501–1535 不属于这两种常规取值范围。识别 LLC 的迹线还需确认底层链路确为 IEEE 802.3，因为无线 802.11 帧也可带 LLC。［4］')
    p(d,'3. 上层协议识别。LLC 使用服务访问点字段区分协议，接收端主要看 DSAP。若为 SNAP，常见 DSAP=0xaa、SSAP=0xaa、Control=0x03，随后五字节为 OUI 和协议标识；OUI=0x000000 时，两字节协议标识沿用 EtherType，例如 IPv4 为 0x0800。［3］')
    p(d,'拓展实测记录：［在本人迹线中用 llc 查找；有样本则记录帧号、Length、LLC 字段和长度，无样本则据实说明未观察到。］',indent=False)
    h(d,'5 3 实测结果与实验体会',2)
    p(d,'结果核对：［待填请求帧的首部长度、帧长、地址对应关系，以及广播和普通组播的数量。将实测值与字段定义对照，说明有无不一致及其原因。］',indent=False)
    p(d,'本实验的关键是把二层字段、上层协议与观测位置联系起来：14 B 首部能够从高亮的原始字节直接验证；目的 MAC 与网关 ARP 映射的对照能够说明二层寻址范围；I/G 位能够解释广播与普通组播的共同特征。抓包中不显示 FCS、填充或实际无线首部，属于观测方式的限制，应与协议格式本身分开分析。')
    p(d,'实验体会：［完成实际操作后，补充本人在选择接口、对应请求回复或展开字段时的具体发现与改进建议。］',indent=False)
    page(d)
    h(d,'六 附录')
    h(d,'6 1 关键命令与过滤器',2)
    table(d,['用途','表达式'],[['生成 IPv4 Echo 流量','ping -4 -n 4 www.baidu.com'],['查看接口与网关','ipconfig /all'],['查看 IPv4 到 MAC 映射','arp -a'],['单播捕获过滤器','icmp'],['广播组播捕获过滤器','ether multicast'],['仅显示 Echo 请求','icmp.type == 8'],['仅显示以太网广播','eth.dst == ff:ff:ff:ff:ff:ff'],['仅显示普通以太网组播','eth.dst.ig == 1 && eth.dst != ff:ff:ff:ff:ff:ff'],['寻找 LLC 样本','llc']],[6.0,10.4])
    h(d,'6 2 数据与资料来源',2)
    p(d,'实测数据文件：［待补本人两份 pcapng 文件的实际文件名和抓包日期］。报告中的图 5 为明确条件下的结构示例；实测截图和地址应以对应迹线为准。')
    refs=[
      '［1］课程资料 计算机网络实验手册  Ethernet 实验；计算机网络实验报告模板。',
      '［2］Wireshark Foundation  Ethernet 与 Ethernet Display Filter Reference。https://wiki.wireshark.org/Ethernet  https://www.wireshark.org/docs/dfref/e/eth.html',
      '［3］RFC 1042  A Standard for the Transmission of IP Datagrams over IEEE 802 Networks。https://www.rfc-editor.org/rfc/rfc1042.html',
      '［4］IANA  IEEE 802 Numbers。https://www.iana.org/assignments/ieee-802-numbers',
      '［5］张荣昊  Lab2 Ethernet 实验报告。用于参考章节组织；其地址、统计与截图不作为本人实测结果。']
    for t in refs:
        x=p(d,t,indent=False)
        for r in x.runs: font(r,9.5)
    d.save(REPORT)

def step(d,num,text):
    x=p(d,f'{num}. {text}',indent=False)
    x.paragraph_format.left_indent=Cm(.5); x.paragraph_format.first_line_indent=Cm(-.5)

def guide():
    d=new_doc('Lab2 抓包与截图操作指南','刘小凡  Ethernet 实验')
    p(d,'按顺序完成单播和广播组播两次抓包。保存两份 pcapng 后再截取证据；只要文件保存成功，缺少的截图以后可以重新打开补拍，不需要重复实验。截图使用 Win + Shift + S，框住与当前步骤有关的界面，并按下文编号发送。',indent=False)
    h(d,'开始前先确认网络')
    p(d,'本次检查发现电脑同时连接 WLAN 与 Cisco AnyConnect 虚拟接口“以太网 2”，且 VPN 默认路由优先。建议临时在 AnyConnect 中点击 Disconnect，保持 Wi-Fi 连接，再使用 WLAN 完成实验。若必须保留 VPN，先记录该状态；WLAN 上的 icmp 过滤器可能只看到零条记录，因为 ping 被封装进了隧道。')
    p(d,'打开命令行的方法：按 Win + R，输入 cmd，按回车。运行 ipconfig，找到实际联网接口。Wi-Fi 对应 WLAN，有线网络对应连接中的以太网；不要选择 VMware、WSL、蓝牙、Loopback 或“本地连接*”来代替真实上网接口。')
    h(d,'第一部分 单播抓包')
    h(d,'步骤 1 先检查 ping 能否回复',2)
    step(d,1,'在 CMD 输入下列命令，回车后等待四次测试结束。')
    cmd(d,'ping -4 -n 4 www.baidu.com')
    step(d,2,'看到“来自 某个 IP 的回复”说明可继续。若全是超时，可换另一个确实能回复的远程目标，例如 www.bing.com，仍保留 -4 -n 4；记住成功目标，后面使用同一个目标。')
    h(d,'步骤 2 设置捕获接口并拍第 1 张',2)
    step(d,1,'打开 Wireshark。如果正在抓包，先点左上角红色方块停止。按 Ctrl + K，或点击顶部“捕获”→“选项”。')
    step(d,2,'在“输入 / Input”标签的接口列表中单击 WLAN。若列表有多行被选中，先单击目标接口，只保留它；确认旁边链路类型为 Ethernet。')
    step(d,3,'在窗口下方“Capture filter for selected interfaces / 所选接口的捕获过滤器”输入 icmp。它是捕获选项窗口里的输入框，不是主窗口顶部的绿色显示过滤器。绿色表示语法有效。')
    step(d,4,'取消“在所有接口上启用混杂模式 / Enable promiscuous mode on all interfaces”的勾选；如果接口行还有单独的混杂模式设置，也保持关闭。无线“监控模式”保持关闭。')
    step(d,5,'此时先别点开始。截取窗口，包含选中接口、Ethernet、icmp 和混杂模式设置，命名为 01_单播捕获设置.png。')
    h(d,'步骤 3 打开 MAC 解析并拍第 2 张',2)
    step(d,1,'仍在捕获选项窗口中，点击上方“选项 / Options”标签。找到右侧“Name Resolution / 名称解析”。')
    step(d,2,'勾选“Resolve MAC addresses / 解析 MAC 地址”，截取该窗口，命名为 02_MAC解析设置.png。')
    step(d,3,'如果该选项不在这里，可关闭窗口，在主窗口菜单“视图”→“名称解析”→“启用 MAC 地址解析”中开启。看到勾选后截图即可。')
    h(d,'步骤 4 捕获 ping 并拍第 3 张',2)
    step(d,1,'回到捕获选项，点击底部“开始”。此时 Wireshark 已进入抓包。')
    step(d,2,'切换到 CMD，再次运行 ping -4 -n 4 www.baidu.com，或刚才验证成功的目标。等待四次结束后，将命令、目标 IP、四次结果和底部统计一起截图，命名为 03_ping结果.png。')
    step(d,3,'返回 Wireshark，点左上角红色方块停止。通常成功的四次 ping 会有四个 request 和四个 reply，不要为了凑数修改结果。')
    step(d,4,r'按 Ctrl + Shift + S，保存为 C:\University Courses\计算机网络lab\Lab2_刘小凡_单播.pcapng。若快捷键无响应，点“文件”→“另存为”，格式选 pcapng。')
    h(d,'步骤 5 展开单播帧并拍第 4 张',2)
    step(d,1,'在主窗口顶部显示过滤器中输入 icmp，按回车。它可以列出请求和回复。单击第一条 Info 中写着 Echo (ping) request 的记录。')
    step(d,2,'中间协议详情区左边的小三角表示展开。依次展开 Frame、Ethernet II、Internet Protocol Version 4、Internet Control Message Protocol。')
    step(d,3,'保留以下字段可见：Frame 的长度；Ethernet 的 Destination、Source、Type；IPv4 的 Header Length、Total Length、Protocol、Source Address、Destination Address；ICMP 的 Type 和 Code。')
    step(d,4,'点击 Ethernet II 文字整行，不是小三角。底部十六进制会高亮前 14 个字节。截图要同时带上方请求回复列表、中间字段和下方字节，命名为 04_单播帧分层.png。')
    p(d,'界面空间不够时：最大化 Wireshark，把上方列表与中间详情之间的分隔线向上拖，把中间与底部字节之间的分隔线向下拖，让中间区更高。无需把 IP 所有子字段展开。仍放不下时可以截两张补充图，我会择优放入报告。',indent=False)
    h(d,'步骤 6 查询地址并拍第 5 和第 6 张',2)
    step(d,1,'在 CMD 运行 ipconfig /all。找到与抓包相同的 WLAN 适配器，截图包括标题、物理地址、IPv4、子网掩码和默认网关，命名为 05_地址证据_ipconfig.png。')
    cmd(d,'ipconfig /all')
    step(d,2,'在 CMD 运行 arp -a。找到“接口”地址与上述 WLAN IPv4 相同的那一组，找到默认网关 IP 对应的“物理地址”，截取接口标题和这一行，命名为 06_地址证据_arp.png。')
    cmd(d,'arp -a')
    step(d,3,'如果没有网关这一行，先 ping 网关一次再运行 arp -a。例如当前 WLAN 网关是 192.168.31.1，可运行 ping -4 -n 1 192.168.31.1。若实际网关已变，替换成 ipconfig 显示的值。')
    p(d,'发送的请求帧中，Source MAC 应对应本机捕获接口；Destination MAC 通常对应该网关。MAC 显示厂家名时，括号中的六组十六进制才是完整地址。如果对应不上，先保留证据，不要手动改成相同值。',indent=False)
    h(d,'第二部分 广播与组播抓包')
    h(d,'步骤 7 修改捕获设置并拍第 7 张',2)
    step(d,1,'确认刚才单播文件已保存，再按 Ctrl + K 打开捕获选项。仍选择同一物理联网接口，Ethernet 类型及混杂模式设置保持一致。')
    step(d,2,'将捕获过滤器 icmp 替换成下列内容。这次也写在捕获选项窗口的过滤器中。')
    cmd(d,'ether multicast')
    step(d,3,'截取接口与过滤器设置，命名为 07_广播捕获设置.png。点击开始，在未主动生成测试流量的情况下等待约 30 秒，然后点红色方块停止。若弹出保存前一份的提示，确认单播已保存后继续。')
    step(d,4,'按 Ctrl + Shift + S，另存为 Lab2_刘小凡_广播组播.pcapng。不要覆盖单播文件。记下大约捕获了多久；下方状态栏的总分组数也保留。')
    h(d,'步骤 8 找广播帧并拍第 8 张',2)
    step(d,1,'在主窗口顶部的显示过滤器输入下列内容，按回车。')
    cmd(d,'eth.dst == ff:ff:ff:ff:ff:ff')
    step(d,2,'单击一条显示出来的记录，优先选择 ARP。中间区展开 Ethernet II，再展开它下面的 Destination。应看到完整地址 ff:ff:ff:ff:ff:ff，以及 IG bit 的 Group / broadcast/multicast 描述。')
    step(d,3,'如果是 ARP，展开 Address Resolution Protocol，保留 Opcode、Sender IP、Target IP 可见。截图同时包括顶部过滤器、列表、Destination 与 IG bit，以及底部总数 / Displayed 数，命名为 08_广播帧与IG位.png。')
    p(d,'注意选择的是 Ethernet II → Destination → IG bit。LG bit 是另外一个地址管理位；第一字节最低位才是 I/G。只需展示真实界面，不需要手工圈写字段，我会在图注中说明。',indent=False)
    h(d,'步骤 9 找普通组播帧并拍第 9 张',2)
    step(d,1,'在主窗口顶部显示过滤器替换成下列内容，按回车。')
    cmd(d,'eth.dst.ig == 1 && eth.dst != ff:ff:ff:ff:ff:ff')
    step(d,2,'如果有记录，单击一条，展开 Ethernet II → Destination，使目的地址和 IG bit 为 Group 可见；保留列表中的上层协议与底部 Displayed 数。命名为 09_组播帧与IG位.png。')
    step(d,3,'如果没有记录，截取空结果或直接告诉我未观察到。广播和组播不需要硬凑固定数量，也无需与参考报告相同。总数与广播数、普通组播数应能核对；需要时我会从 pcapng 重新统计。')
    h(d,'没有捕获到背景流量时怎样处理')
    p(d,'先确认过滤器写在正确位置，且抓包接口仍是实际联网的物理接口。可重新捕获，延长到约 60 秒；保存为新文件并记录新时长。如果仍没有自然广播，报告可以说明未观察到，或者按手册使用教师提供的迹线并明确其来源。')
    p(d,'如需补一份主动广播样本，可在另一份新抓包中运行 ping -4 -n 1 加同网段一个经确认未使用的 IPv4 地址，触发本机为该地址发送 ARP 查询。不需要也不要删除 ARP 缓存。不能随便把参考报告的 IP 填入命令；如果未确认可用目标，先把两份迹线发来，我会再判断是否需要补抓。主动样本应另存并标注，不能充当原先自然背景统计。')
    h(d,'IEEE 802 3 拓展观察')
    p(d,'可打开广播组播文件，在主窗口显示过滤器输入 llc，按回车。有结果时，确认详情显示 IEEE 802.3 和 Logical-Link Control，展开 Length、DSAP、SSAP、Control，截图并保留底部字节。没有结果时据实记录未观察到；报告已准备格式区别和三道拓展题的原理回答。')
    p(d,'若从其他已有迹线或教师提供的样例补充 LLC，应保留该文件的来源和日期。发现 LLC 不等于已经找到 IEEE 802.3，需同时检查底层是否为 802.11。')
    materials_heading=h(d,'发给我的材料清单')
    materials_heading.paragraph_format.page_break_before=True
    table(d,['材料','文件名或内容'],[
      ['单播迹线','Lab2_刘小凡_单播.pcapng'],['背景迹线','Lab2_刘小凡_广播组播.pcapng'],
      ['设置截图','01 单播设置；02 MAC 解析；07 广播设置'],
      ['单播证据','03 ping 结果；04 请求帧分层'],
      ['地址证据','05 ipconfig；06 arp'],
      ['广播组播证据','08 广播与 IG 位；09 普通组播与 IG 位'],
      ['补充文字','实际实验日期、年级；是否保留 VPN；遇到的异常']],[4.8,11.6])
    p(d,'截图直接粘贴到对话即可，分批发送也可以。两份 pcapng 放在当前实验文件夹或附上文件。收到后可以核验字节长度、统计数量和地址关系，将截图裁去无关空白，替换报告中的待填字段，更新图号与结构图。',indent=False)
    h(d,'界面与过滤器依据')
    for t in ['Wireshark User’s Guide  Capture Options。https://www.wireshark.org/docs/wsug_html_chunked/ChCapCaptureOptions.html','Wireshark Ethernet Display Filter Reference。https://www.wireshark.org/docs/dfref/e/eth.html']:
        x=p(d,t,indent=False)
        for r in x.runs: font(r,9)
    d.save(GUIDE)

diagrams()
report()
guide()
print(REPORT)
print(GUIDE)
