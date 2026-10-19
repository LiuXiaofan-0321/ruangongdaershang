# -*- coding: utf-8 -*-
"""check_font_coverage.py - which installed CJK font covers every character used
in the report? Any character a font lacks gets silently substituted by LibreOffice
with a *different* typeface, which is what makes strokes look uneven.
"""
import sys
from docx import Document
from fontTools.ttLib import TTFont

DOC = sys.argv[1] if len(sys.argv) > 1 else \
    r"C:\University Courses\计算机网络lab\10244602411_刘小凡_Lab2_Ethernet_实测版_v2.docx"

d = Document(DOC)
chars = set()
for p in d.paragraphs:
    chars |= set(p.text)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            chars |= set(c.text)
for s in d.sections:
    for p in s.header.paragraphs:
        chars |= set(p.text)
    for p in s.footer.paragraphs:
        chars |= set(p.text)

CJK = {c for c in chars if "\u3000" <= c <= "\u9fff" or c in "、（）【】「」；：，。①-⑳"}
print(f"distinct CJK/punct chars used in document: {len(CJK)}")

FONTS = {
    "DengXian":        (r"C:\Windows\Fonts\Deng.ttf",   (0,)),
    "DengXian Bold":   (r"C:\Windows\Fonts\Dengb.ttf",  (0,)),
    "SimSun":          (r"C:\Windows\Fonts\simsun.ttc", (0,)),
    "SimHei":          (r"C:\Windows\Fonts\simhei.ttf", (0,)),
    "Microsoft YaHei": (r"C:\Windows\Fonts\msyh.ttc",   (0,)),
    "KaiTi":           (r"C:\Windows\Fonts\simkai.ttf", (0,)),
    "FangSong":        (r"C:\Windows\Fonts\simfang.ttf", (0,)),
}

results = {}
for name, (path, idxs) in FONTS.items():
    cov = set()
    for i in idxs:
        try:
            f = TTFont(path, fontNumber=i, lazy=True)
            for table in f["cmap"].tables:
                cov |= set(table.cmap.keys())
        except Exception as exc:  # noqa: BLE001
            print(f"  !! {name}: {exc}")
    missing = sorted(c for c in CJK if ord(c) not in cov)
    results[name] = missing
    preview = "".join(missing[:45])
    print(f"{name:16s} missing {len(missing):3d}/{len(CJK)}  {preview}")

print()
best = sorted(results.items(), key=lambda kv: len(kv[1]))
print("best coverage:", ", ".join(f"{n}({len(m)})" for n, m in best[:3]))
