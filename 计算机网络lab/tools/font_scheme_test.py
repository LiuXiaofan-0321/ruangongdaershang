# -*- coding: utf-8 -*-
"""font_scheme_test.py - build the same sample report with different font schemes,
so we can check which one LibreOffice renders WITHOUT substituting another CJK face.

Usage: python font_scheme_test.py <out.docx> <latin> <cjk_regular> <cjk_bold_family>
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT, LATIN, CJK, CJKB = sys.argv[1:5]
HEAD_FONT = sys.argv[5] if len(sys.argv) > 5 else CJKB   # explicit heading family

d = Document()


def add(text, size=10.5, bold=False, latin=LATIN, ea=CJK, align=None, indent=True):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = latin
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), ea)
    return p


def heading(text, size):
    p = d.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = HEAD_FONT
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    return p


add("华东师范大学软件学院实验报告", 17, True, latin=HEAD_FONT, ea=HEAD_FONT,
    align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
add("Lab2  Ethernet 以太网帧的捕获与分析", 11.5, True, latin=HEAD_FONT, ea=HEAD_FONT,
    align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

heading("一、实验目的", 14)
add("本实验以 ICMP Echo 流量为切入点，用 Wireshark 捕获并逐字段拆解以太网帧，"
    "掌握 DIX Ethernet（Ethernet II）的帧格式与地址作用范围。")
heading("二、实验内容与实验步骤", 14)
heading("4.2  Ethernet II 的字段与长度（逐字节验证）", 12)
add("长度核对：以太网首部 14 B + IPv4 首部 20 B + ICMP 首部 8 B + 数据 32 B = 74 B，"
    "与 Wireshark 报告的 Frame Length 完全一致。")

t = d.add_table(rows=3, cols=3)
data = [["字段", "偏移", "实测值"],
        ["目的 MAC", "0–5", "28:d1:27:df:d4:95"],
        ["源 MAC", "6–11", "14:ac:60:56:df:85"]]
for ri, row in enumerate(data):
    for ci, v in enumerate(row):
        cell = t.rows[ri].cells[ci]
        cell.text = v
        for pp in cell.paragraphs:
            pp.paragraph_format.first_line_indent = Cm(0)
            pp.paragraph_format.space_after = Pt(1)
            for rr in pp.runs:
                rr.font.size = Pt(9.5)
                rr.font.bold = (ri == 0)
                rr.font.name = HEAD_FONT if ri == 0 else LATIN
                ea = HEAD_FONT if ri == 0 else CJK
                rr._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), ea)
b = OxmlElement("w:tblBorders")
for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "4")
    el.set(qn("w:color"), "9DB2C6")
    b.append(el)
t._tbl.tblPr.append(b)

d.add_paragraph()
for line in ["0000  28 d1 27 df d4 95 14 ac 60 56 df 85 08 00 45 00",
             "0010  00 3c 21 e6 00 00 80 01 00 00 c0 a8 1f 87 b4 65"]:
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(line)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK)

d.save(OUT)
print("wrote", OUT)
