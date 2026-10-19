from pathlib import Path
import json, sys, zipfile, hashlib
import fitz
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parent.parent
WORK=Path(__file__).resolve().parent
ref=Path(r'C:\Users\18963\.zcode\workspace\default\大二上A\大二上A\02_计算机网络\实验\实验报告及源码\Lab3\10245101445-张荣昊-Lab3.pdf')
pdf=fitz.open(ref)
(WORK/'reference.txt').write_text('\n'.join(f'PAGE {i+1}\n{p.get_text()}' for i,p in enumerate(pdf)),encoding='utf-8')
print('REFERENCE_PAGES',len(pdf))
for i,p in enumerate(pdf):
    p.get_pixmap(matrix=fitz.Matrix(1,1)).save(WORK/f'reference-{i+1}.png')
doc=Document(ROOT/'10244602411_刘小凡_Ethernet.docx')
texts=[]
for p in doc.paragraphs:
    if p.text: texts.append(p.text)
for i,t in enumerate(doc.tables):
    texts.append(f'TABLE {i}\n'+'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows))
(WORK/'lab2.txt').write_text('\n'.join(texts),encoding='utf-8')
print('LAB2',len(doc.paragraphs),len(doc.tables),len(doc.inline_shapes))
for s in doc.sections:
    print('LAB2_SECTION',s.page_width.cm,s.page_height.cm,s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm)
template=fitz.open(ROOT/'_lab2_工作'/'模板.pdf')
for i,p in enumerate(template): p.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(WORK/f'template-{i+1}.png')
print('TEMPLATE_PAGES',len(template))
print('REFERENCE_TEXT\n'+(WORK/'reference.txt').read_text(encoding='utf-8'))
