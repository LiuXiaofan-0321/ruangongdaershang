from pathlib import Path
import subprocess, sys, json, socket, datetime, os
sys.stdout.reconfigure(encoding='utf-8')
WORK=Path(__file__).resolve().parent
ROOT=WORK.parent
WS=Path(r'C:\Program Files\Wireshark')
IFACE=r'\Device\NPF_{C59C96C9-17AD-4F4E-ADD7-2E166380D039}'
TARGET='182.61.200.108'
LOCAL='172.31.165.86'
for d in ['captures','evidence','logs','qa']: (WORK/d).mkdir(exist_ok=True)
def run(args,name,timeout=60):
    p=subprocess.run(args,capture_output=True,timeout=timeout)
    txt=p.stdout.decode('utf-8',errors='replace')+'\n'+p.stderr.decode('utf-8',errors='replace')
    (WORK/'logs'/name).write_text(txt,encoding='utf-8')
    print(json.dumps({'file':name,'returncode':p.returncode,'output':txt},ensure_ascii=False))
    return p
mode=sys.argv[1]
if mode=='launch':
    args=[str(WS/'Wireshark.exe'),'-i',IFACE,'-p','-f','tcp port 80','-N','n','-o','ip.check_checksum:TRUE','-w',str(WORK/'captures'/'Lab3_HTTP.pcapng')]
    p=subprocess.Popen(args)
    print(json.dumps({'pid':p.pid,'args':args}))
elif mode=='fetch':
    args=[r'C:\Windows\System32\curl.exe','--noproxy','*','-4','--resolve',f'www.baidu.com:80:{TARGET}','--interface',LOCAL,'--connect-timeout','10','--max-time','25','-sS','-D',str(WORK/'logs'/'http_headers.txt'),'-o',str(WORK/'logs'/'baidu.html'),'-w','HTTP_CODE=%{http_code}\nREMOTE_IP=%{remote_ip}\nLOCAL_IP=%{local_ip}\nSIZE_DOWNLOAD=%{size_download}\nTIME_TOTAL=%{time_total}\n','http://www.baidu.com/']
    run(args,'curl.txt',35)
    print((WORK/'logs'/'http_headers.txt').read_text(encoding='utf-8'))
elif mode=='trace':
    cap=subprocess.Popen([str(WS/'tshark.exe'),'-i',IFACE,'-p','-f','icmp','-a','duration:95','-w',str(WORK/'captures'/'Lab3_Traceroute.pcapng')],stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    import time
    time.sleep(2)
    run([r'C:\Windows\System32\tracert.exe','-4','-h','30','-w','700',TARGET],'tracert.txt',100)
    out,err=cap.communicate(timeout=100)
    (WORK/'logs'/'trace_capture.txt').write_text((out+err).decode('utf-8',errors='replace'),encoding='utf-8')
    print('TRACE_CAPTURE',cap.returncode)
elif mode=='environment':
    run([r'C:\Windows\System32\ipconfig.exe','/all'],'ipconfig.txt')
    run([str(WS/'tshark.exe'),'--version'],'wireshark_version.txt')
    run([r'C:\Windows\System32\route.exe','print','-4'],'routes.txt')
    print('DATE',datetime.datetime.now().astimezone().isoformat())
elif mode=='inspect-template':
    from docx import Document
    from lxml import etree
    doc=Document(WORK/'template.docx')
    for i,p in enumerate(doc.paragraphs):
        print(i,repr(p.text),p.style.name,etree.tostring(p._p,encoding='unicode')[:1800])
    for i,t in enumerate(doc.tables):
        print('TABLE',i)
        for r in t.rows: print([c.text for c in r.cells])
        print(etree.tostring(t._tbl,encoding='unicode')[:14000])
    for s in doc.sections: print('SECTION',s.page_width.cm,s.page_height.cm,s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm)
