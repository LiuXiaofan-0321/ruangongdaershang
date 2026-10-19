from pathlib import Path
from copy import deepcopy
import json,re,sys,hashlib,zipfile
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor

sys.stdout.reconfigure(encoding='utf-8')
W=Path(__file__).resolve().parent;R=W.parent;DATA=R/'_lab3_work'
SRC=R/'10244602411_刘小凡_IPv4.docx'
OUT=R/'10244602411_刘小凡_IPv4_优化版.docx'
IMG=W/'images';IMG.mkdir(exist_ok=True)
F=json.loads((DATA/'facts.json').read_text(encoding='utf-8'))
HOPS=json.loads((DATA/'logs'/'source_trace.json').read_text(encoding='utf-8'))
assert F['checksum']=='cb73' and F['sum']=='4fffb' and F['http_target_count']==19
assert len(HOPS)==17 and HOPS[8]['probes'][0]['address'] is None
source_hash=hashlib.sha256(SRC.read_bytes()).hexdigest()

def font(n=38,bold=False,mono=False):
    return ImageFont.truetype(r'C:\Windows\Fonts\consola.ttf' if mono else (r'C:\Windows\Fonts\msyhbd.ttc' if bold else r'C:\Windows\Fonts\msyh.ttc'),n)
def centered(draw,box,text,n=38,bold=False,mono=False):
    f=font(n,bold,mono);x0,y0,x1,y1=box
    bb=draw.multiline_textbbox((0,0),text,font=f,spacing=7,align='center')
    draw.multiline_text(((x0+x1-bb[2])/2,(y0+y1-(bb[3]-bb[1]))/2-bb[1]),text,font=f,fill='#17212b',spacing=7,align='center')

def ipv4_diagram():
    im=Image.new('RGB',(2100,990),'white');g=ImageDraw.Draw(im)
    g.text((80,22),'第 4 帧 IPv4 首部与数据报字节布局',font=font(44,True),fill='black')
    g.text((80,85),'首部从 IP 字节 0 起算 每行 4 字节 共享字节的字段合并为一格',font=font(31),fill='#454545')
    x=150;ww=1880;y=155;hh=102
    rows=[[(1,'Version + IHL\n0x45'),(1,'DS\n0x00'),(2,'Total Length\n0x0075 = 117 B')],
          [(2,'Identification\n0xedc0'),(2,'Flags + Fragment Offset\n0x4000  DF=1  MF=0  Offset=0')],
          [(1,'TTL\n128'),(1,'Protocol\n6 (TCP)'),(2,'Header Checksum\n0x0000（捕获值）')],
          [(4,'Source Address   172.31.165.86')],[(4,'Destination Address   182.61.200.108')]]
    for ri,row in enumerate(rows):
        g.text((18,y+ri*hh+35),f'{ri*4:02d}-{ri*4+3:02d}',font=font(27,mono=True),fill='#444444')
        offset=0
        for count,label in row:
            a=x+offset*ww/4;b=x+(offset+count)*ww/4
            g.rectangle((a,y+ri*hh,b,y+(ri+1)*hh),fill='#edf2f6' if ri<3 else '#f4f6f8',outline='#66717e',width=2)
            centered(g,(a,y+ri*hh,b,y+(ri+1)*hh),label,34 if ri==1 else 38)
            offset+=count
    g.text((80,704),'完整捕获帧 131 B（下条按字节比例绘制）',font=font(36,True),fill='black')
    labels=[(14,'Ethernet\n14 B'),(20,'IPv4 首部\n20 B'),(20,'TCP 首部\n20 B'),(77,'HTTP 请求\n77 B')]
    offset=0;y=767
    for n,label in labels:
        a=80+offset*1950/131;b=80+(offset+n)*1950/131
        g.rectangle((a,y,b,y+100),fill='#f1f3f5' if offset==0 else '#e7eef5',outline='#66717e',width=2)
        centered(g,(a,y,b,y+100),label,32)
        offset+=n
    for n in [0,14,34,54,131]:
        xx=80+n*1950/131
        g.text((xx-14,877),str(n),font=font(27,mono=True),fill='#333333')
    g.text((80,931),'IP Total Length = 117 B = IP 首部 20 B + IP 载荷 97 B；IP 载荷 = TCP 首部 20 B + HTTP 请求 77 B',font=font(28),fill='#333333')
    path=IMG/'ipv4_structure.png';im.save(path);return path

def path_diagram():
    im=Image.new('RGB',(2100,880),'white');g=ImageDraw.Draw(im)
    g.text((60,20),'本机 WLAN 到目标服务器的 17 跳 ICMP 测量路径',font=font(43,True),fill='black')
    nodes=[{'hop':0,'ip':'172.31.165.86','org':'本机 WLAN'}]
    for h in HOPS:
        ip=next((p['address'] for p in h['probes'] if p['address']),None)
        org='未知' if ip is None or ip.startswith('10.') else ('百度 / 终点' if h['hop']==17 else '百度' if ip.startswith('182.61.') else 'CERNET')
        nodes.append({'hop':h['hop'],'ip':ip or '* * *','org':org})
    positions=[]
    for i,node in enumerate(nodes):
        rr=i//6;cc=i%6 if rr%2==0 else 5-i%6
        positions.append((55+cc*337,110+rr*225))
    for i in range(17):
        a,b=positions[i],positions[i+1]
        if a[1]==b[1]:
            direction=1 if b[0]>a[0] else -1
            start=(a[0]+(290 if direction==1 else 0),a[1]+75);end=(b[0]+(0 if direction==1 else 290),b[1]+75)
            g.line((start,end),fill='#576574',width=4)
            g.polygon([end,(end[0]-direction*13,end[1]-8),(end[0]-direction*13,end[1]+8)],fill='#576574')
        else:
            start=(a[0]+145,a[1]+150);end=(b[0]+145,b[1])
            g.line((start,end),fill='#576574',width=4);g.polygon([end,(end[0]-8,end[1]-13),(end[0]+8,end[1]-13)],fill='#576574')
    for node,(x,y) in zip(nodes,positions):
        fill='#e7eef5' if node['hop'] in [0,17] else '#f4f4f4' if node['ip']=='* * *' else '#ffffff'
        g.rectangle((x,y,x+290,y+150),fill=fill,outline='#66717e',width=2)
        centered(g,(x,y+5,x+290,y+48),'起点' if node['hop']==0 else f"第 {node['hop']} 跳",32,True)
        centered(g,(x,y+50,x+290,y+101),node['ip'],32,mono=True)
        centered(g,(x,y+102,x+290,y+147),node['org'],29)
    g.text((60,766),'* 表示 900 ms 内无回复，未知跳保留；箭头表示测得的跳序，不保证真实路由器之间直接相连。',font=font(30),fill='#333333')
    g.text((60,812),'组织名称依据 APNIC 地址登记，仅表示登记归属，不推断设备物理位置；完整 RTT 输出见附录 C。',font=font(30),fill='#333333')
    path=IMG/'internet_path.png';im.save(path);return path

def crop_evidence(name,box,out):
    image=Image.open(DATA/'evidence'/name)
    image.crop(box).save(IMG/out)
    return IMG/out

def setfont(r,size=11,bold=False,mono=False):
    r.font.name='Consolas' if mono else 'Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.italic=False;r.font.color.rgb=RGBColor(0,0,0)
    rf=r._r.get_or_add_rPr().rFonts
    for k in ['ascii','hAnsi','cs']:rf.set(qn('w:'+k),'Consolas' if mono else 'Times New Roman')
    rf.set(qn('w:eastAsia'),'Microsoft YaHei' if bold else 'SimSun')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+k),None)

d=Document(SRC);old=list(d.element.body);sect=deepcopy(d.sections[0]._sectPr)
for sty in ['Title','Heading 1','Heading 2','Heading 3','Caption']:
    if sty not in d.styles:
        created=d.styles.add_style(sty,WD_STYLE_TYPE.PARAGRAPH,builtin=True)
        created.base_style=d.styles['Normal']
    if sty.startswith('Heading '):
        lev=OxmlElement('w:outlineLvl');lev.set(qn('w:val'),str(int(sty[-1])-1))
        d.styles[sty].element.get_or_add_pPr().append(lev)
roles={};ordered=[]
def formatp(p,role='body'):
    f=p.paragraph_format;f.keep_with_next=False;f.keep_together=False;f.widow_control=True;f.page_break_before=False
    f.space_before=Pt(0);f.space_after=Pt(3);f.line_spacing=1.25;f.left_indent=Cm(0);f.right_indent=Cm(0);f.first_line_indent=Pt(22)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    pp=p._p.get_or_add_pPr()
    for t in ['pBdr','shd','framePr','tabs']:
        for e in list(pp.findall(qn('w:'+t))):pp.remove(e)
    snap=pp.find(qn('w:snapToGrid'))
    if snap is None:snap=OxmlElement('w:snapToGrid');pp.append(snap)
    snap.set(qn('w:val'),'0')
    if role=='body':
        p.style=d.styles['Normal']
        for r in p.runs:setfont(r)
    elif role.startswith('heading'):
        level=int(role[-1]);p.style=d.styles['Heading '+str(level)];f.first_line_indent=Pt(0);f.keep_with_next=True;f.keep_together=True
        f.space_before=Pt(10 if level<3 else 7);f.space_after=Pt(5 if level<3 else 3);f.line_spacing=1.2;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:setfont(r,12.5 if level==1 else 11.5 if level==2 else 11,True)
    elif role=='caption':
        p.style=d.styles['Caption'];f.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.line_spacing=1.15;f.keep_together=True;f.space_after=Pt(7)
        for r in p.runs:setfont(r,9.5)
    elif role=='mono':
        p.style=d.styles['Normal'];f.first_line_indent=Pt(0);f.left_indent=Cm(.2);f.line_spacing=1.15;f.space_after=Pt(1);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:setfont(r,9.5,mono=True)
    elif role=='image':
        p.style=d.styles['Normal'];f.first_line_indent=Pt(0);f.line_spacing=1;f.space_before=Pt(4);f.space_after=Pt(2);f.keep_with_next=True;f.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    roles[p._p]=role
    return p._p
def P(text,role='body'):
    p=d.add_paragraph();p.add_run(text);return formatp(p,role)
def use(i,text=None,role='body'):
    e=old[i];p=Paragraph(e,d)
    if text is not None:p.clear();p.add_run(text)
    ordered.append(formatp(p,role));return e
def add(text,role='body'):ordered.append(P(text,role))
def heading(text,level=2):add(text,'heading'+str(level))
def fig(path,caption,width=16.5,extra=None,extra_width=9):
    p=d.add_paragraph();p.add_run().add_picture(str(path),width=Cm(width));ordered.append(formatp(p,'image'))
    if extra:
        p=d.add_paragraph();p.add_run().add_picture(str(extra),width=Cm(extra_width));ordered.append(formatp(p,'image'))
    add(caption,'caption')
def table(i,caption,widths):
    add(caption,'caption');ordered[-1].get_or_add_pPr().append(OxmlElement('w:keepNext'))
    t=Table(old[i],d);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr
    for name in ['tblpPr','tblBorders','tblW']:
        for e in list(pr.findall(qn('w:'+name))):pr.remove(e)
    tw=OxmlElement('w:tblW');tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)/2.54*1440)));pr.append(tw)
    border=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
    pr.append(border)
    for j,w in enumerate(widths):t.columns[j].width=Cm(w)
    for ri,row in enumerate(t.rows):
        row.height=None;tr=row._tr.get_or_add_trPr()
        for el in list(tr):
            if el.tag in [qn('w:trHeight'),qn('w:tblHeader'),qn('w:cantSplit')]:tr.remove(el)
        tr.append(OxmlElement('w:cantSplit'))
        if ri==0:tr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr()
            for name in ['tcMar','shd']:
                for el in list(cp.findall(qn('w:'+name))):cp.remove(el)
            mar=OxmlElement('w:tcMar')
            for edge,value in [('top',75),('bottom',75),('left',90),('right',90)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4' if ri==0 else 'FFFFFF');cp.append(sh)
            for p in c.paragraphs:
                formatp(p);f=p.paragraph_format;f.first_line_indent=Pt(0);f.line_spacing=1.12;f.space_after=Pt(0);f.keep_together=True;f.keep_with_next=ri<len(t.rows)-1
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 and widths[ci]<2 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:setfont(r,10,ri==0)
    ordered.append(t._tbl);return t

structure=ipv4_diagram();pathdraw=path_diagram()
namepic=crop_evidence('02_name_resolution.png',(655,120,1320,340),'name_resolution_crop.png')
capturemain=crop_evidence('01_capture_options.png',(52,130,1880,426),'capture_interface_crop.png')
capturefilter=crop_evidence('01_capture_options.png',(40,770,1890,905),'capture_filter_crop.png')
outpic=crop_evidence('04_outgoing_ipv4.png',(0,0,1700,790),'outgoing_header_crop.png')
listpic=crop_evidence('03_target_connection.png',(5,0,1810,455),'http_list_crop.png')
liststatus=crop_evidence('03_target_connection.png',(1395,1150,1860,1180),'http_list_status.png')

# Title and inline metadata replace the overlapping floating template furniture.
title=Paragraph(old[0],d);title.style=d.styles['Title'];title.paragraph_format.first_line_indent=Pt(0);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after=Pt(8);title.paragraph_format.keep_with_next=True;title.paragraph_format.line_spacing=1.2
for r in title.runs:setfont(r,16,True)
ordered.append(title._p)
sub=P('Lab3 IPv4 数据报与网络路径','caption');Paragraph(sub,d).paragraph_format.keep_with_next=True;ordered.append(sub)
meta=Table(old[2],d)
for x in list(meta._tbl.tblPr.findall(qn('w:tblpPr'))):meta._tbl.tblPr.remove(x)
meta.alignment=WD_TABLE_ALIGNMENT.LEFT;meta.autofit=False
for ci,w in enumerate([7.3,4.8,5.0]):meta.columns[ci].width=Cm(w)
for row in meta.rows:
    for ci,c in enumerate(row.cells):
        c.width=Cm([7.3,4.8,5.0][ci])
        for p in c.paragraphs:
            formatp(p);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:setfont(r,10.5)
ordered.append(meta._tbl)

use(4,role='heading1')
use(11,'本实验通过真实 HTTP 流量和逐 TTL 的 ICMP 探测，分析 IPv4 首部结构、分片判据与网络路径，并独立验算接收报文的首部校验和。具体目标如下。')
for i in range(12,17):
    if i==15:use(i,'4. 使用逐 TTL 探测发现本机到远端服务器的测量路径，保留未知跳，并依据可核实资料标注公网地址的登记组织；')
    else:use(i)
use(6,role='heading1')
use(17,'实验按手册的五个步骤进行：捕获 HTTP 迹线；展开 IPv4 首部并核对字节；绘制字节布局并回答六道问题；根据同一目标服务器的逐跳探测绘制路径；选取接收方向、首部 20 B 且校验和非零的报文完成验算。')
add('提交内容分别为 IPv4 结构图和六题答案、网络路径图及实际探测输出、校验和的十个 16 位字与求和过程，均在第四节及附录中给出。')
use(8,role='heading1')
table(25,'表 1  实验环境与通信端点',[3.6,13.0])
use(10,role='heading1')
use(28,'4.1 HTTP 流量的获取与定位','heading2')
use(18,'本次使用 curl 获取网页。参数 -4 强制使用 IPv4，--resolve 固定 www.baidu.com 的服务器地址，--interface 绑定 WLAN 源地址，--noproxy 禁用 curl 的显式代理；这些参数并不关闭系统 TUN，路径探测的影响另见 4.5 节。命令按 Windows PowerShell 的续行语法书写如下。')
bt=chr(96)
for line in [f'curl.exe -4 --noproxy "*" --interface 172.31.165.86 {bt}',f'  --resolve www.baidu.com:80:182.61.200.108 {bt}',f'  -o baidu.html -w "%{{http_code}} %{{size_download}} %{{time_total}}\\n" {bt}','  http://www.baidu.com/']:
    add(line,'mono')
use(23,'捕获前在 Wireshark 选择 WLAN、Ethernet 链路层和 tcp port 80 捕获过滤器，关闭混杂模式并启用网络名称解析。tcp port 80 本身不限定 IP 版本，本次由 curl -4 产生 IPv4 流量。执行后 curl 返回 HTTP 200，实际下载 2381 B，用时约 0.855 s；响应首部的 Content-Length: 2381 与下载文件大小一致。')
fig(capturemain,'图 1  捕获选项的两处局部：WLAN 行使用 Ethernet；tcp port 80 已设置，混杂模式关闭。',16.5,capturefilter,16.5)
fig(namepic,'图 2  同一次捕获设置中 Resolve network names 已勾选；此图为原始窗口的局部截图。',12.5)
use(34,'开始捕获后执行上述 GET，结束后停止并保存迹线。原始 Lab3_HTTP.pcapng 共 108 帧；按 tcp.stream == 0 导出目标连接，得到 Lab3_HTTP_target.pcapng 的 19 帧，其中发送 10 帧、接收 9 帧。原始前 19 帧就是该连接，因此导出后的帧号保持一致。交互阶段见表 2。')
table(35,'表 2  目标连接的报文阶段',[2.0,2.5,2.8,4.0,5.3])
tt=Table(old[35],d);tt.cell(3,4).text='ACK 与数据段；第 9、14 帧被标记为重传'
fig(DATA/'evidence'/'07_id_ttl_fragmentation.png','图 3  导出的目标连接全部 19 帧：Source、Destination、IP ID、TTL、IP Length、DF、MF 与 Offset。')
fig(DATA/'evidence'/'06_http_200_stream.png','图 4  同一 TCP 流的实际内容：GET / HTTP/1.1、HTTP/1.1 200 OK 和 Content-Length: 2381。')
add('列表中的 Previous segment not captured、Dup ACK 和 Retransmission 提示反映 TCP 序号、丢失或重传分析，不能等同于 IPv4 分片。此次传输最终成功，HTTP 响应的完整正文大小以 curl 保存的文件和响应首部核对。')
use(40,'4.2 IPv4 首部的逐字段检查','heading2')
use(41,'选取第 4 帧 HTTP GET，IPv4 首部长度为 20 B，不含选项。图 5 显示实测字段及原始字节，表 3 将字段与首部内偏移逐项对应。偏移从 IPv4 首部第一个字节起算。')
fig(outpic,'图 5  第 4 帧 IPv4 字段和覆盖首部的原始字节局部：Total Length=117 B、TTL=128、Checksum=0x0000。')
table(44,'表 3  第 4 帧 IPv4 首部字段与字节偏移',[4.4,1.25,1.25,3.8,5.9])
ft=Table(old[44],d);ft.cell(6,4).text='本次 HTTP 发送方向的初值';ft.cell(8,4).text='捕获值为零，与发送端校验和卸载现象相符'
use(45,'本次 HTTP 发送方向的首部校验和在抓包处均为 0x0000，Wireshark 提示可能由 IP checksum offload 引起。这与发送端抓包发生在网卡填写校验和之前的现象相符[6]，不能据此认定线上报文错误，也不能仅凭零值证明每一帧的硬件处理过程。接收方向校验和非零且验证通过，4.6 节选取第 5 帧做独立验算。')
heading('4.3 IPv4 数据报的字节布局')
add('手册要求标出首部各字段的字节位置与长度，并合并共享字节的字段。图 6 将 20 B 首部按每行 4 B 展开；下方完整帧条带按字节数等比绘制，区分以太网首部、IPv4 首部、TCP 首部和 HTTP 请求。')
fig(structure,'图 6  依据第 4 帧绘制的字节布局：帧长 131 B，IP 总长度 117 B，IP 载荷 97 B；本图为示意图。')
use(53,'Version 与 IHL 是两个独立的 4 位字段，共享第 0 字节；Flags 的 3 位与 Fragment Offset 的 13 位共享第 6、7 字节。本次无 Options，IPv4 首部占帧内偏移 14–33，TCP 从帧内偏移 34 开始。')
heading('4.4 六道问题的回答')
add('图 3 已将目标连接全部 19 帧的双向字段列在同一视图。以下结合该列表、原始首部字节与手册要求逐题回答。')
qa=[('本机与远端 IPv4 地址','本机 WLAN 地址为 172.31.165.86，目标服务器为 182.61.200.108；收发方向由源、目的地址的顺序区分。这是本机抓包点看到的地址，不由此断言服务器看到的本机地址，NAT 可能改写地址。'),
('Total Length 的统计范围','Total Length 包含 IPv4 首部及其载荷，不含以太网首部。第 4 帧：117 B = 20 B IP 首部 + 20 B TCP 首部 + 77 B HTTP 请求；帧长 131 B 与 IP 总长度之差为 14 B，恰好是本次 Ethernet II 首部。'),
('Identification 字段的变化','本机方向依次为 0xedbe 至 0xedc7，逐包加 1；接收方向为 0xedbe、0xb73c、0xb73e、0xb73f、0xb740、0xb741、0xb744、0xb745、0xb746，存在跳号。两方向在第 1、2 帧还出现了相同值。因此本次 ID 不在连接内恒定，也不保证双向不同；不能把观察到的递增趋势推广为协议要求。ID 用于分片重组，不是 TCP 连接号。RFC 6864 所称 atomic datagram 满足 DF=1、MF=0、Offset=0，允许其 ID 重复[2]；本次 19 帧满足这一条件。'),
('发送方向 TTL 的初值','本次 HTTP 连接的 10 个发送包在 WLAN 抓包处 TTL 均为 128，低于 8 位字段的最大值 255。路径探测则故意使用 TTL 1–17。接收 SYN/ACK 的 TTL 为 45，其余接收包为 40；未确认对端初始 TTL，不能用 64−接收 TTL 直接推断路由跳数。'),
('未分片报文的判据','MF=0 且 Fragment Offset=0 表示本包未分片。19 帧均满足该判据。DF=1 是禁止后续分片的指示，不是判断未分片的必要条件；DF=0 的短包也可能不分片。MF=0 但 Offset>0 是最后一片。Fragment Offset 以 8 B 为单位；图 3 的 Offset=0 在两种单位下数值相同。'),
('IHL 的长度编码','本次目标连接首部均为 20 B。首字节 0x45 的高 4 位 0100 表示 Version=4，低 4 位 0101 表示 IHL=5。IHL 以 32 位字即 4 B 为单位，5×4=20 B；合法的最小值为 5，4 位最大值 15 对应 60 B，选项与填充最多共 40 B。')]
for i,(q,a) in enumerate(qa,1):heading(f'{i} {q}',3);add(a)
heading('4.5 Internet 路径与隧道影响的排查')
use(67,'首先对与 HTTP 相同的服务器运行 tracert -4 -h 30 -w 700 182.61.200.108，原始输出只显示两跳：')
use(68,role='mono');use(69,role='mono')
use(70,'两跳结果本身不能证明异常，因为目标也可能离源很近；关键证据是对应 WLAN 迹线中 Echo Request 的 TTL 全为 128，未体现 tracert 的逐 TTL 探测。结合 singbox_tun 的默认路由指向 172.18.0.2，可判断该结果受到本机隧道影响，不能据此描述 WLAN 的逐跳转发路径。排查未修改 VPN 或系统路由。')
use(71,'随后通过 Windows IcmpSendEcho2Ex API 绑定 SourceAddress=172.31.165.86，每跳发送 3 次探测，单次超时 900 ms。保存 API 返回结果，并在 WLAN 抓包核对实际 TTL。该实现为绑定 WLAN 源地址的 traceroute 脚本，附录 C 保存其完整输出[4]。')
use(64,'逐 TTL 探测使不同距离的路由器返回 ICMP Time Exceeded。第 5 帧的外层源地址为 10.100.4.1、目的为本机，ICMP Type=11、Code=0；内部引用原始请求的 IPv4 首部与 ICMP 首部，供发送方匹配探测，见图 7[3]。')
fig(DATA/'evidence'/'09_icmp_time_exceeded.png','图 7  路径迹线第 5 帧：路由器 10.100.4.1 返回 Type 11 / Code 0，并引用本机发往服务器的原始请求。')
use(72,'实际探测在第 17 跳收到目标回复。76 帧中有 51 个 Echo Request、23 个 Time Exceeded 和 2 个 Echo Reply。TTL 1–17 各对应 3 个发出的探测。仅使用 icmp.type == 8 会同时匹配超时报文引用的内层 Echo 请求；图 8 的过滤条件同时排除了外层 Type 11，因此显示 51 帧，而不是包含引用的 74 帧。')
fig(DATA/'evidence'/'08_path_ttl_probes.png','图 8  TTL 1、2、3 各三次 Echo Request 的列表局部，以及同一视图的统计栏；过滤后显示 51 个请求。',16.5,DATA/'evidence'/'08_path_status.png')
fig(pathdraw,'图 9  依据 source_trace.json 绘制的 17 跳测量路径，含本机、目标服务器与未知跳；本图为示意图。')
use(79,'本次 PTR 反向解析未成功，组织标注使用 APNIC RDAP 的地址登记资料。202.120.95.246 所在登记范围为 202.112.0.0–202.121.255.255；101.4.116.61 属于 101.4.0.0/14，219.224.103.65 属于 219.224.0.0/15，182.61.255.32 属于 182.61.0.0/16。登记组织不代表设备物理位置，私网地址 10.100.4.1 的组织保持未知。查询结果与完整 RTT 输出分别见附录 A、C。')
use(80)
use(81,'可确认的是：所响应的公网地址先后落在 CERNET 与百度的登记范围，目标 182.61.200.108 在第 17 跳回复。现有 PTR 和登记资料不能确认具体校园设备、上海交通大学节点或路由器地理位置，因此不作这类推断。')
heading('4.6 接收包 IPv4 首部校验和验算')
use(83,'按手册要求选择目标连接第 5 帧：接收方向、首部 20 B、校验和非零。其源为 182.61.200.108、目的为 172.31.165.86，实收首部校验和 0xcb73，Wireshark 标记 [correct]。图 10 高亮了这 20 B；首部字节如下。')
use(84,role='mono')
fig(DATA/'evidence'/'05_incoming_checksum.png','图 10  目标连接第 5 帧的 20 B IPv4 首部高亮，实收校验和 0xcb73，经 Wireshark 验证为 correct。')
use(85,'按网络字节序将首部分成十个 16 位字，包含实收校验和字段进行普通加法，再将进位回卷，见表 4。')
table(86,'表 4  接收首部的十个 16 位字与普通累计和',[1.3,2.7,8.6,4.0])
ct=Table(old[86],d)
for row,text in [(7,'源地址前半：182.61'),(8,'源地址后半：200.108'),(9,'目的地址前半：172.31'),(10,'目的地址后半：165.86')]:ct.cell(row,2).text=text
use(87,'普通和为 0x4fffb。高 16 位 0x0004 加回低 16 位 0xfffb，得到 0xffff；再取反为 0x0000，完成接收验证。此求和只覆盖 IPv4 首部，不覆盖 TCP 或 HTTP 载荷[1]。')
add('接收验证：0x4fffb → 0xfffb + 0x0004 = 0xffff → 取反 = 0x0000。','mono')
use(88,'另将首部校验和字段临时置零进行生成复算：普通和 0x43488，回卷得到 0x348c，取反为 0xcb73。该值与原始报文及 Wireshark 结果一致。置零仅用于计算，未改写抓包文件。')
add('生成复算：0x43488 → 0x3488 + 0x0004 = 0x348c → 取反 = 0xcb73。','mono')
use(89,'上述字节和求和结果已通过 analyze.py 从 pcapng 重新读取并独立核验，与 GUI 显示相互印证。第 5 帧 Ethernet 帧长为 60 B，而 IP 总长度为 40 B；14 B 链路首部之外的 6 B 填充不参与 IP 首部校验和计算。')

use(97,'五、实验结果总结','heading1')
heading('5.1 实验结论')
for i in range(99,104):
    if i==99:use(i,'关于 IPv4 首部：本次目标 HTTP 连接的所有报文首部长度均为 20 B，无选项字段。首字节 0x45 编码版本 4 与首部长度 5，IHL 以 4 B 为单位，最大值 15 对应 60 B。Total Length 统计 IPv4 首部与载荷的总和，不含链路层首部。')
    elif i==102:use(i,'关于校验和：接收包首部以实收值求和回卷得到 0xffff，置零生成复算得到 0xcb73；发送包捕获值为零与校验和卸载现象相符，不能据此判定线上错误。')
    else:use(i)
heading('5.2 实验体会')
use(105,'本次实验中，发送包的零校验和和 tracert 的两跳结果都需要结合抓包位置分析。前者应与校验和卸载区分，后者需检查实际发出的 TTL。把工具提示与独立字段验证相结合，才能避免把本机处理现象误写为网络结论。')
use(106);use(107)
use(108,'六、附录','heading1')
heading('附录 A 迹线与验证文件')
table(110,'表 5  可复核的实验数据文件',[6.3,10.3])
Table(old[110],d).cell(7,1).text='17 跳路径结构，图 9 的数据来源'
heading('附录 B 复核命令')
add('以下命令以 PowerShell 为例，反引号表示续行。抓包文件名按附录 A；绑定源地址的路径探测实现为 source_trace.py。')
for line in [f'tshark.exe -r Lab3_HTTP_target.pcapng -Y "frame.number == 5" {bt}','  -o ip.check_checksum:TRUE -V -x',f'tshark.exe -r Lab3_Path_WLAN.pcapng {bt}',f'  -Y "icmp.type == 8 && !(icmp.type == 11)" {bt}','  -T fields -E occurrence=f -e frame.number -e ip.src -e ip.dst -e ip.ttl']:
    add(line,'mono')
heading('附录 C 完整逐跳测量输出')
add('以下为 source_trace.txt 的完整实测输出，星号表示探测超时。')
Paragraph(ordered[-1],d).paragraph_format.keep_with_next=True
trace_lines=(DATA/'logs'/'source_trace.txt').read_text(encoding='utf-8').splitlines()
for i,line in enumerate(trace_lines):
    add(line,'mono')
    if i<len(trace_lines)-1:Paragraph(ordered[-1],d).paragraph_format.keep_with_next=True
heading('附录 D 参考资料')
refs=[('[1] RFC 791 Internet Protocol','https://www.rfc-editor.org/rfc/rfc791.html'),('[2] RFC 6864 Updated Specification of the IPv4 ID Field','https://www.rfc-editor.org/rfc/rfc6864.html'),('[3] RFC 792 Internet Control Message Protocol','https://www.rfc-editor.org/rfc/rfc792.html'),('[4] Microsoft Learn IcmpSendEcho2Ex function','https://learn.microsoft.com/en-us/windows/win32/api/icmpapi/nf-icmpapi-icmpsendecho2ex'),('[5] Wireshark IPv4 display filter reference','https://www.wireshark.org/docs/dfref/i/ip.html'),('[6] Wireshark User Guide Checksums','https://www.wireshark.org/docs/wsug_html_chunked/ChAdvChecksums.html')]
for title,url in refs:
    add(title);p=Paragraph(ordered[-1],d);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
    add(url,'mono')

# Replace only the body content, preserving the source header/footer and section geometry.
body=d.element.body
for e in list(body):body.remove(e)
for e in ordered:body.append(e)
body.append(sect)
for stname in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
    st=d.styles[stname];st.font.color.rgb=RGBColor(0,0,0)
    pp=st.element.find(qn('w:pPr'))
    if pp is not None:
        for e in list(pp.findall(qn('w:pBdr'))):pp.remove(e)
# Restore formatting after replacing text in some table cells.
for t in d.tables[1:]:
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=ri<len(t.rows)-1
                p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(0)
                for r in p.runs:setfont(r,10,ri==0)
settings=d.settings.element
uf=settings.find(qn('w:updateFields'))
if uf is None:uf=OxmlElement('w:updateFields');settings.append(uf)
uf.set(qn('w:val'),'true')
d.core_properties.author='刘小凡';d.core_properties.title='Lab3 IPv4 数据报与网络路径实验报告'
d.core_properties.comments=''
d.save(OUT)
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==source_hash
check=Document(OUT)
major=[p.text for p in check.paragraphs if p.style.name=='Heading 1']
assert major==['一、实验目的','二、实验内容与实验步骤','三、实验环境','四、实验过程与分析','五、实验结果总结','六、附录'],major
captions=[p.text for p in check.paragraphs if p.style.name=='Caption' and p.text.startswith('图 ')]
assert [int(re.match(r'图 (\d+)',x)[1]) for x in captions]==list(range(1,11))
with zipfile.ZipFile(SRC) as a,zipfile.ZipFile(OUT) as b:
    for name in a.namelist():
        if re.match(r'word/(header|footer)\d+\.xml$',name):assert a.read(name)==b.read(name),name
(W/'checks.json').write_text(json.dumps({'source_sha256':source_hash,'output':str(OUT),'sections':major,'figures':len(captions),'inline_images':len(check.inline_shapes),'tables':len(check.tables),'source_unchanged':True,'header_footer_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
print('SAVED',OUT)
