from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

root=Path(r'C:\University Courses\计算机网络lab')
source=root/'10244602411_刘小凡_Protocol_Layer_优化版.docx'
out=root/'10244602411_刘小凡_Protocol_Layer_补图优化版.docx'
temp=Path(r'C:\Users\18963\AppData\Local\Temp')
doc=Document(source)
ps=list(doc.paragraphs)
shots={
 'ip':'codex-clipboard-0190e312-89f6-4d6e-a7cf-daad4686240d.png',
 'tcp':'codex-clipboard-b923275d-635e-4ac2-8740-5a972f9cef66.png',
 'reassembly':'codex-clipboard-9e0e3234-ad0e-49b3-b2ee-3b8e1e17879b.png',
 'download':'codex-clipboard-e40aebe3-45a0-4ded-a4a3-f13d0e2e1788.png',
 'demux':'codex-clipboard-6137536c-7150-4642-90e5-e93087682e1b.png'}

def rewrite(p,text):
    # Preserve paragraph formatting and the first run's typography.
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r=p.add_run(text)
    if props is not None:r._r.insert(0,props)

def picture(p,key,box=None):
    path=temp/shots[key]
    w,h=Image.open(path).size
    if box is None:box=(0,0,w,h)
    x0,y0,x1,y1=box
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format
    pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0)
    pf.space_before=Pt(4);pf.space_after=Pt(2)
    pf.keep_with_next=True;pf.keep_together=True
    shape=p.add_run().add_picture(str(path),width=Cm(16.4))
    # Native Word cropping keeps the complete original screenshot embedded.
    shape.height=int(shape.width*(y1-y0)/(x1-x0))
    fill=shape._inline.xpath('.//pic:blipFill')[0]
    rect=OxmlElement('a:srcRect')
    for side,v in [('l',x0/w),('t',y0/h),('r',(w-x1)/w),('b',(h-y1)/h)]:
        rect.set(side,str(round(v*100000)))
    fill.insert(1,rect)
    shape._inline.docPr.set('descr',key+' Wireshark screenshot from lab1.pcapng')
    return p

def insert_after(anchor,text='',style='Normal'):
    p=doc.add_paragraph(text,style)
    anchor._p.addnext(p._p)
    return p

def add_evidence(anchor,lead,key,caption,box):
    intro=insert_after(anchor,lead)
    intro.paragraph_format.keep_with_next=True
    pic=insert_after(intro)
    picture(pic,key,box)
    cap=insert_after(pic,caption,'Caption')
    cap.paragraph_format.keep_with_next=False
    cap.paragraph_format.keep_together=True
    return cap

# 4.2: preserve the GET analysis, then add its TCP evidence and a separately
# identified response IPv4 example. The latter has Total Length 885, not 117.
anchor=add_evidence(ps[32],
 '在 No. 99227 中展开 TCP 首部，可直接核对源端口 57891、目的端口 80、TCP Segment Len=77 和 Header Length=20 bytes。选中整个 TCP 协议项后，字节区高亮其 20 B 首部，与上表及 GET 字节结构图一致。',
 'tcp','图 0  No. 99227 TCP 首部字段及 20 B 字节高亮',
 (0,580,1422,1280))
add_evidence(anchor,
 '另选取下载响应 No. 101068 核对 IPv4 首部。其 Header Length=20 bytes，Total Length=885，Protocol=TCP (6)；加上 14 B Ethernet 首部，帧长为 899 B。该图展示的是响应报文，GET 请求的 IPv4 Total Length 仍为 117 B。',
 'ip','图 0  No. 101068 响应报文的 IPv4 首部及原始字节',
 (0,590,1422,1305))

# Replace the old, collapsed response screenshot with the detailed new one.
ps[35].clear()
picture(ps[35],'reassembly',(0,580,1422,1510))
rewrite(ps[36],'图 0  No. 101068 HTTP 响应重组及重复序列范围')
rewrite(ps[34],ps[34].text+' 展开的重组详情显示 6 条分段记录，其中 No. 102270 与 No. 101067 的 payload 范围同为 789–1812，No. 101119、No. 102271 与 No. 101068 的范围同为 1813–2657。重复范围不增加唯一字节数，因此重组长度仍为 2658 B。')

# 4.4: actual filtered download list supplements the existing calculation.
add_evidence(ps[42],
 '使用 tcp.stream == 95 && tcp.srcport == 80 筛选下载方向，共显示 12 条捕获记录，其中包含重复数据段和 FIN 重传。下表所采用的 7 个首轮报文只是其中的子集，不能把 12 条记录全部当作首次有效传输。',
 'download','图 0  下载方向全部 12 条报文及 HTTP 200 OK',
 (0,0,1422,1020))

# Explicitly state the scope of the existing result without silently replacing it.
rewrite(ps[47],
 '因此，按上述首次有效传输口径计算，协议首部开销约为 12.8%，HTTP 数据约占 87.2%；把所选短帧 padding 计入非应用数据时约为 13.3%。这两个比例对应表中所选的 7 个报文，不代表包含重传的全部 12 条下载记录。')

# 4.5: the shorter screenshot gives direct evidence for the IPv4 key.
add_evidence(ps[49],
 'No. 101068 的 IPv4 首部明确显示 Protocol: TCP (6)，验证 IPv4 将载荷交给 TCP 处理；前面的响应 IPv4 截图还显示 Ethernet Type: IPv4 (0x0800)，验证 Ethernet 到 IPv4 的分用。',
 'demux','图 0  No. 101068 IPv4 Protocol 分用字段及首部高亮',
 (0,435,1422,1055))

# Update appendix with filters actually visible in the supplied screenshots.
p=insert_after(ps[85],'tcp.stream == 95 && tcp.srcport == 80')
insert_after(p,'frame.number == 101068')

# Number all figure captions in their final document order.
n=0
for p in doc.paragraphs:
    if re.match(r'^图\s*\d+',p.text):
        n+=1
        rewrite(p,re.sub(r'^图\s*\d+',f'图 {n}',p.text))
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0)

# Keep the appendix together and remove a duplicate inherited list marker.
ps[80].paragraph_format.page_break_before=True
pr=ps[63]._p.get_or_add_pPr()
numpr=pr.find(qn('w:numPr'))
if numpr is not None:pr.remove(numpr)

# Keep table headings repeated after page breaks; avoid rows splitting.
for table in doc.tables[1:]:
    pr=table.rows[0]._tr.get_or_add_trPr()
    if pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
    for row in table.rows:
        pr=row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))

# Remove the inherited decorative line directly beneath the title.
ppr=ps[0]._p.pPr
if ppr is not None:
    border=ppr.find(qn('w:pBdr'))
    if border is not None:ppr.remove(border)
for style_name in ['Title']:
    ppr=doc.styles[style_name]._element.find(qn('w:pPr'))
    if ppr is not None:
        border=ppr.find(qn('w:pBdr'))
        if border is not None:ppr.remove(border)

doc.save(out)
print(str(out));print('figures',n,'images',len(doc.inline_shapes))
